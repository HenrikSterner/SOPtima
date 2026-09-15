from __future__ import annotations

import io
import html
import hashlib
import json
import re
import unicodedata
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime, time as dt_time, timedelta
from pathlib import Path
from typing import Any

import pandas as pd
import streamlit as st


ROOT = Path(__file__).resolve().parent
HTML_SOURCE = ROOT / "vejlederfordelingBETA.html"
DEFAULT_TEACHERS_FILE = ROOT / "lærere.csv"
DEFAULT_TEMPLATE = ROOT / "SOP - SKABELON.docx"
TEST_STUDENTS_FILE = ROOT / "testdata" / "sop_test_elever.csv"
TEST_TEACHERS_FILE = ROOT / "testdata" / "sop_test_lærere.csv"
NO_WISHES_LABEL = "Ingen ønsker"
UNASSIGNED_LABEL = "Ingen vejleder"
COPYRIGHT = "SOPtima (Copyright Henrik Sterner 2026)"


def repair_text(value: Any) -> str:
    """Gør også ældre filer med dobbeltkodet UTF-8 læsbare."""
    text = "" if value is None else str(value).strip()
    if text.casefold() in {"nan", "nat", "none", "<na>"}:
        return ""
    if any(marker in text for marker in ("Ã", "Â", "�")):
        try:
            fixed = text.encode("latin1").decode("utf-8")
            before_markers = sum(text.count(marker) for marker in ("Ã", "Â", "�"))
            after_markers = sum(fixed.count(marker) for marker in ("Ã", "Â", "�"))
            if after_markers < before_markers:
                return fixed
        except (UnicodeEncodeError, UnicodeDecodeError):
            pass
    return text


def normal_key(value: Any) -> str:
    return re.sub(r"\s+", " ", repair_text(value)).strip().casefold()


def canonical_subject(value: Any) -> str:
    subject = normal_key(value)
    subject = re.sub(r"\s+[abc]\*?$", "", subject)
    if subject in {"komm/info", "komm/it", "kommunikation og it", "kommunikation og it"}:
        return "kommunikation og it"
    repaired = {
        "idï¿½historie": "idehistorie",
        "idéhistorie": "idehistorie",
        "idrï¿½t": "idraet",
        "idræt": "idraet",
        "erhvervsï¿½konomi": "erhvervsoekonomi",
        "erhvervsøkonomi": "erhvervsoekonomi",
    }
    if subject in repaired:
        return repaired[subject]
    subject = "".join(char for char in unicodedata.normalize("NFD", subject) if unicodedata.category(char) != "Mn")
    return (
        subject.replace("æ", "ae")
        .replace("ø", "oe")
        .replace("å", "aa")
        .replace("Ã¦", "ae")
        .replace("Ã¸", "oe")
        .replace("Ã¥", "aa")
    )


def parse_capacity(value: Any, *, default: int = 0, maximum: int = 50) -> int:
    """Læs et max-tal uden at acceptere tvetydige eller negative værdier."""
    text = repair_text(value)
    if not text:
        return default
    match = re.fullmatch(r"\s*(\d+)(?:\.0+)?\s*(?:elever?|opgaver?)?\s*", text, flags=re.I)
    if not match:
        raise ValueError(f"Max-værdien '{text}' skal være et helt tal mellem 0 og {maximum}.")
    capacity = int(match.group(1))
    if capacity > maximum:
        raise ValueError(f"Max-værdien {capacity} er større end den tilladte grænse på {maximum}.")
    return capacity


def unique_wishes(student: dict[str, Any]) -> list[str]:
    """Returnér højst to forskellige, normaliserede lærerønsker."""
    return list(dict.fromkeys(normal_key(wish) for wish in student.get("wishes", []) if normal_key(wish)))[:2]


def wish_status(student: dict[str, Any], assigned: list[str | None]) -> tuple[int, int]:
    """Returnér antal opfyldte og antal afgivne, forskellige ønsker."""
    wishes = unique_wishes(student)
    fulfilled = len(set(wishes) & {normal_key(item) for item in assigned if item})
    return fulfilled, len(wishes)


def wish_status_label(student: dict[str, Any], assigned: list[str | None]) -> str:
    fulfilled, total = wish_status(student, assigned)
    return "Ingen angivet" if total == 0 else f"{fulfilled}/{total}"


def safe_spreadsheet_value(value: Any) -> Any:
    """Neutralisér tekst, som ellers kan blive fortolket som en regnearksformel."""
    if isinstance(value, str) and re.match(r"^\s*[=+\-@]", value):
        return "'" + value
    return value


def safe_excel_frame(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.map(safe_spreadsheet_value)


def teacher_identity(value: Any) -> tuple[str, str]:
    """Læs både initialer og etiketten 'Navn (initialer)' fra en tidsplanfil."""
    label = repair_text(value)
    match = re.search(r"\(([^()]+)\)\s*$", label)
    if match:
        teacher_id = normal_key(match.group(1))
        name = repair_text(label[: match.start()]) or teacher_id
        return teacher_id, name
    return normal_key(label), label


def extract_embedded(name: str, terminator: str) -> Any:
    if not HTML_SOURCE.exists():
        return None
    source = HTML_SOURCE.read_text(encoding="utf-8")
    match = re.search(rf"const {name} = (.*?);\s*{terminator}", source, re.S)
    if not match:
        return None
    try:
        return json.loads(match.group(1))
    except json.JSONDecodeError:
        return None


@st.cache_data(show_spinner=False)
def embedded_data() -> tuple[list[dict[str, Any]], list[dict[str, Any]], dict[str, int]]:
    students = extract_embedded("STUDENTS", r"const TEACHERS") or []
    teachers = extract_embedded("TEACHERS", r"const DOCX_TEMPLATE_PARTS") or []
    capacities = extract_embedded("TEACHER_CAPACITIES", r"const teacherCapacities") or {}
    # Cloud-demoen behøver ikke den gamle HTML-kilde; brug de fiktive CSV-data som fallback.
    if not students and TEST_STUDENTS_FILE.exists():
        students = parse_students(read_path_table(TEST_STUDENTS_FILE))
    if not teachers and TEST_TEACHERS_FILE.exists():
        teachers, capacities = parse_teachers(read_path_table(TEST_TEACHERS_FILE))
    for student in students:
        student["name"] = repair_text(student.get("name"))
        student["className"] = repair_text(student.get("className"))
        student["subjects"] = [repair_text(item) for item in student.get("subjects", [])]
        student["wishes"] = [normal_key(item) for item in student.get("wishes", []) if item]
    for teacher in teachers:
        teacher["id"] = normal_key(teacher.get("id"))
        teacher["name"] = repair_text(teacher.get("name") or teacher["id"])
        teacher["subjects"] = [repair_text(item) for item in teacher.get("subjects", [])]
        teacher["holds"] = [repair_text(item) for item in teacher.get("holds", []) if repair_text(item)]
    return students, teachers, {normal_key(k): int(v) for k, v in capacities.items()}


def read_uploaded_table(uploaded: Any) -> pd.DataFrame:
    data = uploaded.getvalue()
    filename = uploaded.name.casefold()
    if filename.endswith((".xlsx", ".xlsm", ".xls")):
        return pd.read_excel(io.BytesIO(data), sheet_name=0, dtype=object)
    try:
        return pd.read_csv(io.BytesIO(data), sep=None, engine="python", dtype=object, encoding="utf-8-sig")
    except UnicodeDecodeError:
        return pd.read_csv(io.BytesIO(data), sep=None, engine="python", dtype=object, encoding="cp1252")


def read_path_table(path: Path) -> pd.DataFrame:
    if path.suffix.casefold() in {".xlsx", ".xlsm", ".xls"}:
        return pd.read_excel(path, sheet_name=0, dtype=object)
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return pd.read_csv(path, sep=None, engine="python", dtype=object, encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"Kunne ikke læse {path.name}.")


def column_for(columns: list[Any], *patterns: str) -> Any | None:
    normalised = {column: normal_key(column) for column in columns}
    for column, key in normalised.items():
        if all(pattern in key for pattern in patterns):
            return column
    return None


def value_at(row: pd.Series, column: Any | None) -> str:
    return repair_text(row.get(column, "")) if column is not None else ""


def parse_schedule_upload(uploaded: Any) -> dict[str, Any]:
    """Læs en færdig vejlederliste til direkte generering af en tidsplan.

    Formatet er én række pr. elev med kolonnerne Elev, Klasse, Lærer 1,
    Fag 1, Lærer 2, Fag 2 og eventuelt bemærkninger fra lærerne. Lærerne
    oprettes automatisk ud fra de navne/initialer, der står i arket.
    """
    table = read_uploaded_table(uploaded).dropna(how="all").copy()

    def has_schedule_header(columns: list[Any]) -> bool:
        keys = [normal_key(column) for column in columns]
        has_student = any("elev" in key for key in keys)
        has_class = any("klasse" in key or "hold" in key for key in keys)
        teacher_and_subject_columns = sum(
            "lærer" in key or "laerer" in key or "vejleder" in key or "fag" in key or "subject" in key
            for key in keys
        )
        return has_student and has_class and teacher_and_subject_columns >= 4

    def find_header_row(raw: pd.DataFrame) -> int | None:
        for row_index, row in raw.head(25).iterrows():
            if has_schedule_header([value for value in row.tolist() if repair_text(value)]):
                return int(row_index)
        return None

    # Accepter også Excel-filer med titel/instruktioner før tabelhovedet eller
    # et andet første ark, så længe et ark indeholder den angivne tabel.
    if not has_schedule_header(list(table.columns)) and uploaded.name.casefold().endswith((".xlsx", ".xlsm")):
        sheets = pd.read_excel(io.BytesIO(uploaded.getvalue()), sheet_name=None, header=None, dtype=object)
        for raw in sheets.values():
            header_row = find_header_row(raw)
            if header_row is None:
                continue
            table = raw.iloc[header_row + 1 :].copy()
            table.columns = [repair_text(value) or f"Kolonne {index + 1}" for index, value in enumerate(raw.iloc[header_row].tolist())]
            table = table.dropna(how="all")
            break
    columns = list(table.columns)

    def find_column(*alternatives: tuple[str, ...]) -> Any | None:
        for patterns in alternatives:
            found = column_for(columns, *patterns)
            if found is not None:
                return found
        return None

    name_col = find_column(("elev",), ("elevnavn",), ("navn",))
    class_col = find_column(("klasse",), ("hold",))
    teacher_columns = [
        find_column(("lærer", "1"), ("laerer", "1"), ("vejleder", "1"), ("teacher", "1")),
        find_column(("lærer", "2"), ("laerer", "2"), ("vejleder", "2"), ("teacher", "2")),
    ]
    subject_columns = [
        find_column(("fag", "1"), ("subject", "1")),
        find_column(("fag", "2"), ("subject", "2")),
    ]
    notes_col = (
        find_column(("bemærk",), ("bemerk",), ("note",), ("kommentar",))
        or find_column(("ændring",), ("aendring",))
    )

    if name_col is None or class_col is None or any(column is None for column in teacher_columns + subject_columns):
        found_columns = ", ".join(repair_text(column) or "(tom)" for column in columns)
        raise ValueError(
            "Tidsplanarket skal indeholde kolonnerne Elev, Klasse, Lærer 1, Fag 1, "
            "Lærer 2 og Fag 2. Kolonnen Bemærkninger/ændringer fra lærerne er valgfri. "
            f"Fundne kolonner: {found_columns}."
        )

    students: list[dict[str, Any]] = []
    teachers_by_id: dict[str, dict[str, Any]] = {}
    assignments: list[list[str]] = []
    invalid_rows: list[str] = []
    for row_number, (_, row) in enumerate(table.iterrows(), 2):
        name = value_at(row, name_col)
        if not name:
            continue
        teacher_names = [value_at(row, column) for column in teacher_columns]
        subjects = [value_at(row, column) for column in subject_columns]
        if any(not value for value in teacher_names + subjects):
            invalid_rows.append(str(row_number))
            continue

        teacher_identities = [teacher_identity(value) for value in teacher_names]
        teacher_ids = [item[0] for item in teacher_identities]
        for (teacher_id, teacher_name), subject in zip(teacher_identities, subjects):
            teacher = teachers_by_id.setdefault(
                teacher_id,
                {"id": teacher_id, "name": teacher_name, "subjects": [], "holds": []},
            )
            if normal_key(teacher.get("name")) == teacher_id and normal_key(teacher_name) != teacher_id:
                teacher["name"] = teacher_name
            if canonical_subject(subject) not in {canonical_subject(item) for item in teacher["subjects"]}:
                teacher["subjects"].append(subject)

        students.append(
            {
                "id": f"E{len(students) + 1:03d}",
                "name": name,
                "className": value_at(row, class_col),
                "subjects": subjects,
                "subjectsWithLevel": subjects[:],
                "wishes": [],
                "projectTitle": "",
                "projectDescription": "",
                "scheduleNotes": value_at(row, notes_col),
            }
        )
        assignments.append(teacher_ids)

    if invalid_rows:
        rows = ", ".join(invalid_rows[:10])
        suffix = " …" if len(invalid_rows) > 10 else ""
        raise ValueError(f"Række(r) {rows}{suffix} mangler elevens to lærere eller fag.")
    if not students:
        raise ValueError("Der blev ikke fundet nogen komplette elever i tidsplanarket.")

    teachers = sorted(teachers_by_id.values(), key=lambda item: item["name"].casefold())
    loads = {teacher["id"]: 0 for teacher in teachers}
    for assigned in assignments:
        for teacher_id in set(assigned):
            loads[teacher_id] += 1
    return {
        "students": students,
        "teachers": teachers,
        "solution": {"assignments": assignments, "loads": loads},
        "source_name": getattr(uploaded, "name", "upload.xlsx"),
    }


def make_schedule_input_template() -> bytes:
    """Lav et tomt Excel-ark med det forventede tidsplanformat."""
    columns = [
        "Elev",
        "Klasse",
        "Lærer 1",
        "Fag 1",
        "Lærer 2",
        "Fag 2",
        "",
        "Bemærkninger/ændringer fra lærerne",
    ]
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        safe_excel_frame(pd.DataFrame(columns=columns)).to_excel(writer, index=False, sheet_name="Tidsplan-input")
        sheet = writer.book["Tidsplan-input"]
        sheet.freeze_panes = "A2"
        for index, width in enumerate((24, 12, 22, 22, 22, 22, 4, 45), 1):
            sheet.column_dimensions[chr(64 + index)].width = width
    return output.getvalue()


def parse_students(table: pd.DataFrame) -> list[dict[str, Any]]:
    table = table.dropna(how="all").copy()
    columns = list(table.columns)
    name_col = column_for(columns, "elevnavn") or column_for(columns, "navn")
    class_col = column_for(columns, "klasse") or column_for(columns, "hold")
    subject_cols = []
    wish_cols = []
    for index in (1, 2):
        subject_cols.append(
            column_for(columns, "fag", str(index))
            or column_for(columns, "subject", str(index))
            or (columns[index + 1] if len(columns) > index + 1 else None)
        )
        wish_cols.append(
            column_for(columns, "ønske", str(index))
            or column_for(columns, "ønsk", str(index))
            or column_for(columns, "onske", str(index))
            or column_for(columns, "wish", str(index))
        )
    if name_col is None or any(item is None for item in subject_cols):
        raise ValueError("Elevarket skal indeholde kolonnerne Elevnavn/Navn, Fag 1 og Fag 2. Klasse/Hold anbefales, men er valgfri.")

    title_col = column_for(columns, "projekttitel") or column_for(columns, "projekt", "titel") or column_for(columns, "projecttitle") or column_for(columns, "projektnavn")
    description_col = column_for(columns, "projektbeskrivelse") or column_for(columns, "projekt", "beskrivelse") or column_for(columns, "projectdescription") or column_for(columns, "beskrivelse") or column_for(columns, "description")
    students = []
    invalid_rows = []
    for row_number, (_, row) in enumerate(table.iterrows(), 1):
        name = value_at(row, name_col)
        subjects = [value_at(row, column) for column in subject_cols]
        row_has_data = any(repair_text(value) for value in row.tolist())
        if row_has_data and not name:
            invalid_rows.append(f"række {row_number} mangler elevnavn")
            continue
        if name and not any(subjects):
            invalid_rows.append(f"{name} mangler begge fag")
            continue
        wishes = [normal_key(value_at(row, column)) for column in wish_cols]
        students.append(
            {
                "id": f"E{len(students) + 1:03d}",
                "name": name,
                "className": value_at(row, class_col) or "Ukendt klasse/hold",
                "subjects": subjects,
                "wishes": [wish for wish in wishes if wish],
                "subjectsWithLevel": subjects[:],
                "projectTitle": value_at(row, title_col),
                "projectDescription": value_at(row, description_col),
            }
        )
    if invalid_rows:
        suffix = " …" if len(invalid_rows) > 10 else ""
        raise ValueError("Elevarket indeholder ugyldige rækker: " + "; ".join(invalid_rows[:10]) + suffix + ".")
    if not students:
        raise ValueError("Der blev ikke fundet nogen elever i elevarket.")
    return students


def subject_from_course(value: Any, technical_subjects: dict[str, str]) -> str | None:
    subject = repair_text(value)
    if not subject or normal_key(subject).startswith("bro"):
        return None
    if normal_key(subject) in technical_subjects:
        return technical_subjects[normal_key(subject)]
    if "teknikfag" in normal_key(subject):
        # Bevar den konkrete retning, så fx "Digitalt design ... teknikfag"
        # kan matches direkte mod elevens fag uden en separat oversigt.
        subject = re.sub(r"\s+[ABC]\*?$", "", subject, flags=re.I)
        return subject.strip() or "Teknikfag"
    subject = re.sub(r"^\d+(?:vf)?\s*\.\s*", "", subject, flags=re.I)
    subject = re.sub(r"^\d+(?:vf)?\s+", "", subject, flags=re.I)
    subject = re.sub(r"\s+[ABC]\*?$", "", subject, flags=re.I)
    return subject.strip() or None


def parse_teachers(table: pd.DataFrame, technical: dict[str, str] | None = None) -> tuple[list[dict[str, Any]], dict[str, int]]:
    technical = technical or {}
    columns = list(table.columns)
    id_col = column_for(columns, "initial") or column_for(columns, "lærer")
    name_col = column_for(columns, "lærernavn") or column_for(columns, "navn")
    course_col = column_for(columns, "fag")
    capacity_col = column_for(columns, "max") or column_for(columns, "kapac")
    if id_col is None or course_col is None:
        raise ValueError("Lærerarket skal indeholde kolonnerne Initialer og Fag.")
    class_col = column_for(columns, "klasse")
    hold_columns = [column for column in columns if normal_key(column).startswith("hold")][:4]
    if not hold_columns and class_col is not None:
        hold_columns = [class_col]
    teachers_by_id: dict[str, dict[str, Any]] = {}
    capacities: dict[str, int] = {}
    subject_columns = [column for column in columns if normal_key(column).startswith("fag")]
    invalid_rows = []
    for row_number, (_, row) in enumerate(table.iterrows(), 1):
        teacher_id = normal_key(row.get(id_col, ""))
        if not teacher_id:
            if any(repair_text(value) for value in row.tolist()):
                invalid_rows.append(f"række {row_number} mangler initialer")
            continue
        if capacity_col is not None:
            try:
                parsed_capacity = parse_capacity(value_at(row, capacity_col))
                capacities[teacher_id] = min(capacities.get(teacher_id, parsed_capacity), parsed_capacity)
            except ValueError as error:
                invalid_rows.append(f"{teacher_id}: {error}")
        subjects = []
        course_values = [row.get(course_col)] + [row.get(column) for column in subject_columns if column != course_col]
        for course in course_values:
            subject = subject_from_course(course, technical)
            if class_col is not None and "teknikfag" in normal_key(course):
                detailed = subject_from_course(row.get(class_col), technical)
                if detailed and detailed != "Teknikfag":
                    subject = detailed
            if subject and canonical_subject(subject) not in {canonical_subject(item) for item in subjects}:
                subjects.append(subject)
        teachers_by_id.setdefault(teacher_id, {"id": teacher_id, "name": value_at(row, name_col) or teacher_id, "subjects": [], "holds": []})
        for subject in subjects:
            if canonical_subject(subject) not in {canonical_subject(item) for item in teachers_by_id[teacher_id]["subjects"]}:
                teachers_by_id[teacher_id]["subjects"].append(subject)
        for column in hold_columns:
            hold = value_at(row, column)
            if hold and normal_key(hold) not in {normal_key(item) for item in teachers_by_id[teacher_id]["holds"]}:
                teachers_by_id[teacher_id]["holds"].append(hold)
    if invalid_rows:
        suffix = " …" if len(invalid_rows) > 10 else ""
        raise ValueError("Lærerarket indeholder ugyldige rækker: " + "; ".join(invalid_rows[:10]) + suffix + ".")
    teachers = sorted(teachers_by_id.values(), key=lambda item: item["id"])
    return teachers, capacities


def parse_teacher_upload(uploaded: Any, technical: dict[str, str] | None = None) -> tuple[list[dict[str, Any]], dict[str, int]]:
    """Læs både lærerlisten på ark 1 og max-tal på ark 2 i timefagfordeling.xlsx."""
    technical = technical or {}
    filename = uploaded.name.casefold()
    if not filename.endswith((".xlsx", ".xlsm", ".xls")):
        return parse_teachers(read_uploaded_table(uploaded), technical)
    sheets = pd.read_excel(io.BytesIO(uploaded.getvalue()), sheet_name=None, dtype=object)
    first = next(iter(sheets.values()))
    teachers, capacities = parse_teachers(first, technical)
    if len(sheets) > 1:
        second = list(sheets.values())[1]
        for _, row in second.iterrows():
            if len(row) < 2:
                continue
            teacher_id = normal_key(row.iloc[0])
            raw_capacity = repair_text(row.iloc[1])
            if teacher_id and raw_capacity:
                value = parse_capacity(raw_capacity)
                capacities[teacher_id] = min(capacities.get(teacher_id, value), value)
    return teachers, capacities


def teacher_label(teacher_id: str | None, teachers_by_id: dict[str, dict[str, Any]]) -> str:
    if not teacher_id:
        return "Ikke tildelt"
    teacher = teachers_by_id.get(normal_key(teacher_id))
    if not teacher:
        return teacher_id
    return f"{teacher['name']} ({teacher['id']})" if normal_key(teacher["name"]) != teacher["id"] else teacher["id"]


def build_candidates(teachers: list[dict[str, Any]]) -> dict[str, list[str]]:
    candidates: dict[str, list[str]] = defaultdict(list)
    for teacher in teachers:
        for subject in teacher["subjects"]:
            if teacher["id"] not in candidates[canonical_subject(subject)]:
                candidates[canonical_subject(subject)].append(teacher["id"])
    return dict(candidates)


def wished_count(student: dict[str, Any], assigned: list[str | None]) -> int:
    """Antal forskellige lærerønsker, som er repræsenteret i tildelingen."""
    return wish_status(student, assigned)[0]


def optimize(
    students: list[dict[str, Any]],
    teachers: list[dict[str, Any]],
    capacities: dict[str, int],
    K: int,
    double_limit: int,
    use_global_k: bool,
    allow_over_capacity: bool,
    lock_teacher_max: bool,
    prioritize_pairs: bool,
    prioritize_classes: bool,
    attempts: int = 120,
    progress_callback: Any = None,
) -> dict[str, Any]:
    candidates = build_candidates(teachers)
    teacher_ids = [teacher["id"] for teacher in teachers]

    def candidate_for(subject: Any) -> list[str]:
        return candidates.get(canonical_subject(subject), [])

    def capacity_of(teacher_id: str) -> int:
        value = int(capacities.get(teacher_id, 0))
        return min(K, value) if use_global_k else value

    permit_over = use_global_k and allow_over_capacity and not lock_teacher_max
    best = None
    seed_base = K * 100003 + len(students) * 97 + double_limit * 7919
    demand = Counter(canonical_subject(subject) for student in students for subject in student["subjects"])

    def make_attempt(seed: int) -> dict[str, Any]:
        rng = __import__("random").Random(seed)
        assignments: list[list[str | None]] = [[None, None] for _ in students]
        loads = {teacher_id: 0 for teacher_id in teacher_ids}
        double_loads = {teacher_id: 0 for teacher_id in teacher_ids}

        def capacity_available(teacher_id: str, index: int) -> bool:
            return (
                teacher_id in assignments[index]
                or loads[teacher_id] < capacity_of(teacher_id)
                or (permit_over and loads[teacher_id] < K)
            )

        def pressure(subject: str) -> float:
            ids = candidate_for(subject)
            available = sum(K if permit_over else capacity_of(teacher_id) for teacher_id in ids)
            return demand[canonical_subject(subject)] / max(1, available)

        def can_share(index: int, slot: int, teacher_id: str) -> bool:
            other = assignments[index][1 - slot]
            if teacher_id != other:
                return True
            return assignments[index][0] == teacher_id and assignments[index][1] == teacher_id or double_loads[teacher_id] < double_limit

        def assign(index: int, slot: int, teacher_id: str | None) -> None:
            before = assignments[index][:]
            before_set = {item for item in before if item}
            before_double = before[0] if before[0] and before[0] == before[1] else None
            assignments[index][slot] = teacher_id
            after = assignments[index]
            after_set = {item for item in after if item}
            after_double = after[0] if after[0] and after[0] == after[1] else None
            for teacher_id_ in before_set | after_set:
                if teacher_id_ in before_set and teacher_id_ not in after_set:
                    loads[teacher_id_] -= 1
                if teacher_id_ not in before_set and teacher_id_ in after_set:
                    loads[teacher_id_] += 1
            if before_double and before_double != after_double:
                double_loads[before_double] -= 1
            if after_double and before_double != after_double:
                double_loads[after_double] += 1

        coverage = []
        for index, student in enumerate(students):
            options = []
            for slot in (0, 1):
                allowed = set(candidate_for(student["subjects"][slot]))
                options.extend({"slot": slot, "teacher": wish} for wish in dict.fromkeys(student["wishes"]) if wish in allowed)
            coverage.append({"index": index, "options": options, "pressure": max(pressure(item) for item in student["subjects"]), "scarcity": min(len(candidate_for(item)) for item in student["subjects"]), "tie": rng.random()})
        coverage.sort(key=lambda item: (-item["pressure"], len(item["options"]), item["scarcity"], item["tie"]))
        for item in coverage:
            index = item["index"]
            student = students[index]
            choices = [option for option in item["options"] if can_share(index, option["slot"], option["teacher"]) and capacity_available(option["teacher"], index)]
            choices.sort(key=lambda option: (
                int(not candidate_for(student["subjects"][1 - option["slot"]]) or len([item for item in candidate_for(student["subjects"][1 - option["slot"]]) if item != option["teacher"]]) == 0),
                loads[option["teacher"]],
                -len([item for item in candidate_for(student["subjects"][1 - option["slot"]]) if item != option["teacher"]]),
                rng.random(),
            ))
            if choices:
                assign(index, choices[0]["slot"], choices[0]["teacher"])

        remaining = []
        for index, student in enumerate(students):
            for slot in (0, 1):
                if not assignments[index][slot]:
                    subject = student["subjects"][slot]
                    remaining.append({"index": index, "slot": slot, "count": len(candidate_for(subject)), "pressure": pressure(subject), "tie": rng.random()})
        remaining.sort(key=lambda item: (-item["pressure"], item["count"], item["tie"]))
        for item in remaining:
            index, slot = item["index"], item["slot"]
            if assignments[index][slot]:
                continue
            student = students[index]
            choices = [teacher_id for teacher_id in candidate_for(student["subjects"][slot]) if can_share(index, slot, teacher_id) and capacity_available(teacher_id, index)]
            other_slot = 1 - slot
            def choice_key(teacher_id: str) -> tuple[Any, ...]:
                other = candidate_for(student["subjects"][other_slot])
                future_bad = 0 if any(item != teacher_id and capacity_available(item, index) for item in other) else 1
                return (max(0, loads[teacher_id] + 1 - capacity_of(teacher_id)), future_bad, int(teacher_id not in student["wishes"]), loads[teacher_id], rng.random())
            choices.sort(key=choice_key)
            if choices:
                assign(index, slot, choices[0])

        for _ in range(8):
            moved = False
            overloaded = sorted((item for item in loads if loads[item] > capacity_of(item)), key=lambda item: -loads[item])
            for teacher_id in overloaded:
                held = []
                for index, student in enumerate(students):
                    for slot in (0, 1):
                        if assignments[index][slot] == teacher_id:
                            before = wished_count(student, assignments[index])
                            held.append((2 if teacher_id in student["wishes"] and before == 1 else 1 if teacher_id in student["wishes"] else 0, rng.random(), index, slot))
                held.sort()
                for _, _, index, slot in held:
                    if loads[teacher_id] <= capacity_of(teacher_id):
                        break
                    student = students[index]
                    alternatives = [item for item in candidate_for(student["subjects"][slot]) if item != teacher_id and can_share(index, slot, item) and capacity_available(item, index)]
                    alternatives.sort(key=lambda item: (int(item not in student["wishes"]), loads[item], rng.random()))
                    if alternatives:
                        assign(index, slot, alternatives[0])
                        moved = True
            if not moved:
                break

        for _ in range(4):
            improved = False
            order = sorted(range(len(students)), key=lambda index: (wished_count(students[index], assignments[index]), rng.random()))
            for index in order:
                student = students[index]
                for slot in (0, 1):
                    current = assignments[index][slot]
                    if not current or current in student["wishes"]:
                        continue
                    preferred = [item for item in dict.fromkeys(student["wishes"]) if item in candidate_for(student["subjects"][slot]) and can_share(index, slot, item) and capacity_available(item, index)]
                    preferred.sort(key=lambda item: loads[item])
                    if preferred:
                        assign(index, slot, preferred[0])
                        improved = True
            if not improved:
                break

        for _ in range(3):
            improved = False
            for index, student in enumerate(students):
                wishes = list(dict.fromkeys(student["wishes"]))
                if len(wishes) < 2:
                    continue
                current_count = wished_count(student, assignments[index])
                pairs = [(wishes[0], wishes[1]), (wishes[1], wishes[0])]
                for pair in pairs:
                    if pair[0] not in candidate_for(student["subjects"][0]) or pair[1] not in candidate_for(student["subjects"][1]) or wished_count(student, list(pair)) <= current_count:
                        continue
                    changes = Counter(item for item in pair)
                    changes.subtract(item for item in assignments[index] if item)
                    if all(loads[item] + delta <= capacity_of(item) or (permit_over and loads[item] + delta <= K) for item, delta in changes.items()):
                        assign(index, 0, None)
                        assign(index, 1, None)
                        assign(index, 0, pair[0])
                        assign(index, 1, pair[1])
                        improved = True
                        break
            if not improved:
                break

        stats = score_solution(students, teachers, capacities, assignments, loads, K, use_global_k, prioritize_pairs, prioritize_classes)
        return {"assignments": assignments, "loads": loads, "double_loads": double_loads, "stats": stats, "K": K, "double_limit": double_limit, "use_global_k": use_global_k, "allow_over_capacity": permit_over, "lock_teacher_max": lock_teacher_max}

    attempts = max(1, int(attempts))
    if progress_callback:
        progress_callback(0.0)
    for attempt in range(attempts):
        candidate = make_attempt(seed_base + attempt * 7919)
        if best is None or candidate["stats"]["score"] < best["stats"]["score"]:
            best = candidate
        if progress_callback:
            progress_callback((attempt + 1) / attempts)
    return best


def update_algorithm_progress(progress_bar: Any, label: str, value: float) -> None:
    """Opdater en Streamlit-progressbar med en ensartet procenttekst."""
    bounded = max(0.0, min(1.0, float(value)))
    progress_bar.progress(bounded, text=f"{label}: {round(bounded * 100)} %")


def score_solution(students: list[dict[str, Any]], teachers: list[dict[str, Any]], capacities: dict[str, int], assignments: list[list[str | None]], loads: dict[str, int], K: int, use_global_k: bool, prioritize_pairs: bool, prioritize_classes: bool) -> dict[str, Any]:
    candidates = build_candidates(teachers)
    def cap(teacher_id: str) -> int:
        value = capacities.get(teacher_id, 0)
        return min(K, value) if use_global_k else value
    unassigned = sum(len([item for item in assigned if item is None]) > 0 for assigned in assignments)
    none = one = both = desired_total = requested_total = no_wishes = all_wishes = capacity_blocked_slots = 0
    blocked_students = set()
    pair_counts = Counter()
    class_counts = Counter()
    for index, assigned in enumerate(assignments):
        student = students[index]
        count = wished_count(student, assigned)
        _, requested = wish_status(student, assigned)
        desired_total += count
        requested_total += requested
        if requested == 0:
            no_wishes += 1
        elif count == 0:
            none += 1
        elif count >= 2:
            both += 1
        else:
            one += 1
        if requested and count == requested:
            all_wishes += 1
        for slot, teacher_id in enumerate(assigned):
            if not teacher_id:
                possible = candidates.get(canonical_subject(student["subjects"][slot]), [])
                if possible and all(loads[item] >= cap(item) for item in possible):
                    capacity_blocked_slots += 1
                    blocked_students.add(student["id"])
        for teacher_id in set(item for item in assigned if item):
            class_counts[f"{student.get('className', 'Ukendt')}|{teacher_id}"] += 1
        if assigned[0] and assigned[1] and assigned[0] != assigned[1]:
            pair_counts[f"{'|'.join(sorted(canonical_subject(item) for item in student['subjects']))}::{ '|'.join(sorted(assigned)) }"] += 1
    overload = sum(max(0, load - cap(teacher_id)) for teacher_id, load in loads.items())
    over_teachers = sum(load > cap(teacher_id) for teacher_id, load in loads.items())
    k_overload = sum(max(0, load - K) for load in loads.values()) if use_global_k else 0
    k_over_teachers = sum(load > K for load in loads.values()) if use_global_k else 0
    max_load = max(loads.values(), default=0)
    load_squares = sum(load * load for load in loads.values())
    pair_score = sum(value * (value - 1) // 2 for value in pair_counts.values())
    class_score = sum(value * (value - 1) // 2 for value in class_counts.values())
    score = k_overload * 1e15 + unassigned * 1e12 + overload * 1e9 + none * 1e6 + (requested_total - desired_total) * 1e4 - (pair_score * .05 if prioritize_pairs else 0) - (class_score * .05 if prioritize_classes else 0) + max_load * 100 + load_squares
    return {"score": score, "unassigned": int(unassigned), "none": none, "one": one, "both": both, "no_wishes": no_wishes, "all_wishes": all_wishes, "wish_students": len(students) - no_wishes, "requested_total": requested_total, "desired_total": desired_total, "overload": overload, "over_teachers": over_teachers, "k_overload": k_overload, "k_over_teachers": k_over_teachers, "max_load": max_load, "same_teacher_pair_score": pair_score, "same_class_teacher_score": class_score, "capacity_blocked_slots": capacity_blocked_slots, "capacity_blocked_students": len(blocked_students)}


def subject_stats(
    students: list[dict[str, Any]],
    teachers: list[dict[str, Any]],
    capacities: dict[str, int],
    K: int,
    use_global_k: bool,
    allow_over_capacity: bool = False,
    lock_teacher_max: bool = False,
) -> pd.DataFrame:
    candidates = build_candidates(teachers)
    grouped: dict[str, dict[str, Any]] = {}
    for student in students:
        for subject in student["subjects"]:
            key = canonical_subject(subject)
            grouped.setdefault(key, {"Fag": subject, "Elever": set()})["Elever"].add(student["id"])
    rows = []
    permit_over = use_global_k and allow_over_capacity and not lock_teacher_max
    for key, item in grouped.items():
        ids = candidates.get(key, [])
        capacity = sum(K if permit_over else min(K, capacities.get(teacher_id, 0)) if use_global_k else capacities.get(teacher_id, 0) for teacher_id in ids)
        rows.append({"Fag": item["Fag"], "Elever": len(item["Elever"]), "Lærere": len(ids), "Kapacitet": capacity, "Elever pr. lærer": round(len(item["Elever"]) / len(ids), 1) if ids else None, "Lærerinitialer": ", ".join(ids), "Mangler kapacitet": max(0, len(item["Elever"]) - capacity)})
    return pd.DataFrame(rows).sort_values(["Mangler kapacitet", "Elever pr. lærer", "Elever"], ascending=[False, False, False]) if rows else pd.DataFrame()


def make_export(
    students: list[dict[str, Any]],
    teachers: list[dict[str, Any]],
    solution: dict[str, Any],
    capacities: dict[str, int],
    use_global_k: bool,
    selected_sheets: list[str] | None = None,
) -> bytes:
    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    rows = []
    for index, student in enumerate(students):
        assigned = solution["assignments"][index]
        fulfilled_wishes, requested_wishes = wish_status(student, assigned)
        rows.append({"Elev-ID": student["id"], "Elev": student["name"], "Klasse": student.get("className", ""), "Fag 1": student["subjects"][0], "Vejleder fag 1": teacher_label(assigned[0], teacher_map), "Vejleder-ID 1": assigned[0] or "", "Fag 2": student["subjects"][1], "Vejleder fag 2": teacher_label(assigned[1], teacher_map), "Vejleder-ID 2": assigned[1] or "", "Ønsker angivet": requested_wishes, "Ønsker opfyldt": fulfilled_wishes, "Projekttitel": student.get("projectTitle", ""), "Projektbeskrivelse": student.get("projectDescription", "")})
    distribution = pd.DataFrame(rows)
    load_rows = []
    for teacher in sorted(teachers, key=lambda item: teacher_label(item["id"], teacher_map)):
        load = solution["loads"].get(teacher["id"], 0)
        individual_limit = int(capacities.get(teacher["id"], 0))
        normal_limit = min(solution["K"], individual_limit) if use_global_k else individual_limit
        permit_over = use_global_k and solution.get("allow_over_capacity", False) and not solution.get("lock_teacher_max", False)
        hard_limit = solution["K"] if permit_over else normal_limit
        load_rows.append({"Lærer": teacher_label(teacher["id"], teacher_map), "Initialer": teacher["id"], "Fag": " · ".join(teacher["subjects"]), "Antal elever": load, "Individuelt max": individual_limit, "Hård grænse": hard_limit, "Over individuelt max": max(0, load - normal_limit), "Over hård grænse": max(0, load - hard_limit), "Dobbeltvejledninger": solution["double_loads"].get(teacher["id"], 0)})
    loads = pd.DataFrame(load_rows)
    stats = subject_stats(students, teachers, capacities, solution["K"], use_global_k, solution.get("allow_over_capacity", False), solution.get("lock_teacher_max", False))
    unassigned = distribution[distribution[["Vejleder-ID 1", "Vejleder-ID 2"]].eq("").any(axis=1)]
    if selected_sheets is not None:
        sheet_frames = {
            "Fordeling": distribution,
            "Lærerbelastning": loads,
            "Fagstatistik": stats,
            "Ikke tildelte": unassigned,
            "Elevdata": students_to_frame(students, teachers),
            "Lærerdata": teachers_to_frame(teachers, capacities),
        }
        sheet_order = ["Fordeling", "Lærerbelastning", "Fagstatistik", "Ikke tildelte", "Elevdata", "Lærerdata"]
        selected = [name for name in sheet_order if name in selected_sheets]
        if not selected:
            raise ValueError("Vælg mindst ét ark til Excel-filen.")
        selected_output = io.BytesIO()
        with pd.ExcelWriter(selected_output, engine="openpyxl") as writer:
            for sheet_name in selected:
                safe_excel_frame(sheet_frames[sheet_name]).to_excel(writer, index=False, sheet_name=sheet_name)
            for sheet in writer.book.worksheets:
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
                for column in sheet.columns:
                    width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
                    sheet.column_dimensions[column[0].column_letter].width = width
        return selected_output.getvalue()
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        safe_excel_frame(distribution).to_excel(writer, index=False, sheet_name="Fordeling")
        safe_excel_frame(loads).to_excel(writer, index=False, sheet_name="Lærerbelastning")
        safe_excel_frame(stats).to_excel(writer, index=False, sheet_name="Fagstatistik")
        safe_excel_frame(unassigned).to_excel(writer, index=False, sheet_name="Ikke tildelte")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(45, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
    return output.getvalue()


def _clock_label(value: datetime) -> str:
    return value.strftime("%H:%M")


def _schedule_pair_label(teacher_ids: tuple[str, ...], teacher_map: dict[str, dict[str, Any]]) -> str:
    return " + ".join(teacher_label(teacher_id, teacher_map) for teacher_id in teacher_ids) or "Mangler vejleder"


def _round_order_metrics(
    round_order: list[int],
    round_teachers: list[set[str]],
    round_pairs: list[set[tuple[str, ...]]] | None = None,
) -> tuple[int, int]:
    """Beregn antal lærerhuller og sammenhængende lærerpar for en rækkefølge."""
    positions = {round_number: position for position, round_number in enumerate(round_order)}
    teacher_positions: dict[str, list[int]] = defaultdict(list)
    for round_number, teachers_in_round in enumerate(round_teachers):
        position = positions[round_number]
        for teacher_id in teachers_in_round:
            teacher_positions[teacher_id].append(position)
    gap_rounds = sum(
        max(positions_for_teacher) - min(positions_for_teacher) + 1 - len(positions_for_teacher)
        for positions_for_teacher in teacher_positions.values()
        if len(positions_for_teacher) > 1
    )
    pair_affinity = 0
    if round_pairs is not None:
        pair_affinity = sum(
            len(round_pairs[left] & round_pairs[right])
            for left, right in zip(round_order, round_order[1:])
        )
    return gap_rounds, pair_affinity


def _optimise_teacher_round_order(
    round_teachers: list[set[str]],
    initial_order: list[int],
    round_pairs: list[set[tuple[str, ...]]] | None = None,
) -> list[int]:
    """Forsøg at samle hver lærers runder ved at optimere rækkefølgen af runder.

    Runderne ændres ikke, og dermed ændres hverken konfliktfriheden eller antallet
    af runder. Det er en heuristik, fordi den eksakte rækkefølgeoptimering ellers
    bliver unødigt dyr for store tidsplaner.
    """
    if len(initial_order) < 2:
        return initial_order[:]

    def score(order: list[int]) -> tuple[int, int]:
        gaps, affinity = _round_order_metrics(order, round_teachers, round_pairs)
        return gaps, -affinity

    def greedy_order(start: int) -> list[int]:
        remaining = set(initial_order)
        remaining.remove(start)
        order = [start]
        seen_teachers = set(round_teachers[start])
        while remaining:
            previous = order[-1]

            def candidate_key(round_number: int) -> tuple[int, int, int, int]:
                current_teachers = round_teachers[round_number]
                # Fortsæt først med lærere fra den foregående runde. En lærer,
                # der allerede har været aktiv, men ikke er aktiv nu, tæller som
                # et sandsynligt hul, hvis den dukker op igen.
                overlap = len(round_teachers[previous] & current_teachers)
                stale = len((seen_teachers - round_teachers[previous]) & current_teachers)
                affinity = (
                    len(round_pairs[previous] & round_pairs[round_number])
                    if round_pairs is not None
                    else 0
                )
                return stale, -overlap, -affinity, round_number

            next_round = min(remaining, key=candidate_key)
            order.append(next_round)
            remaining.remove(next_round)
            seen_teachers.update(round_teachers[next_round])
        return order

    # Brug den eksisterende rækkefølge som kandidat og suppler med nogle
    # forskellige greedy-starter. Det giver flere muligheder uden at gøre
    # tidsplanberegningen uforholdsmæssigt langsom.
    starts = [initial_order[0], initial_order[-1]]
    starts.extend(initial_order[:: max(1, len(initial_order) // 8)])
    candidate_orders = [initial_order[:], list(reversed(initial_order))]
    candidate_orders.extend(greedy_order(start) for start in dict.fromkeys(starts))
    best_order = min(candidate_orders, key=score)
    best_score = score(best_order)

    # Flytning af én hel runde kan lukke et hul, som simple naboombytninger ikke
    # kan finde. Begræns søgningen til den almindelige størrelse for en tidsplan;
    # ved meget store planer er greedy-resultatet stadig en gyldig forbedring.
    if len(best_order) <= 120:
        improved = True
        while improved:
            improved = False
            for source in range(len(best_order)):
                for target in range(len(best_order)):
                    if source == target or abs(source - target) < 2:
                        continue
                    candidate = best_order[:]
                    moved = candidate.pop(source)
                    candidate.insert(target, moved)
                    candidate_score = score(candidate)
                    if candidate_score < best_score:
                        best_order, best_score = candidate, candidate_score
                        improved = True
                        break
                if improved:
                    break
    return best_order


def _schedule_time(value: Any, *, field_name: str = "tidspunkt") -> dt_time:
    """Læs et tidspunkt fra tidsplanens indstillinger uden at gætte på formatet."""
    if isinstance(value, dt_time):
        return value.replace(second=0, microsecond=0)
    if isinstance(value, datetime):
        return value.time().replace(second=0, microsecond=0)
    text = repair_text(value)
    try:
        return datetime.strptime(text, "%H:%M").time()
    except ValueError as error:
        raise ValueError(f"{field_name.capitalize()} skal skrives som TT:MM, fx 10:30.") from error


def normalise_teacher_blocks(
    teacher_blocks: dict[str, Any] | None,
    teachers: list[dict[str, Any]],
) -> dict[str, list[tuple[dt_time, dt_time]]]:
    """Validér og saml lærernes eventuelle spærringer pr. dag.

    Værdierne kan være par af ``datetime.time`` eller ``TT:MM``. Overlappende
    spærringer for samme lærer samles, så tidsplanlægningen kun skal forholde
    sig til adskilte tidsrum.
    """
    known_ids = {normal_key(teacher["id"]) for teacher in teachers}
    result: dict[str, list[tuple[dt_time, dt_time]]] = {teacher_id: [] for teacher_id in known_ids}
    for raw_teacher_id, raw_intervals in (teacher_blocks or {}).items():
        teacher_id = normal_key(raw_teacher_id)
        if not teacher_id:
            continue
        if teacher_id not in known_ids:
            raise ValueError(f"Spærringen henviser til den ukendte lærer '{repair_text(raw_teacher_id)}'.")
        if raw_intervals is None:
            continue
        if isinstance(raw_intervals, tuple) and len(raw_intervals) == 2 and not isinstance(raw_intervals[0], (tuple, list, dict)):
            raw_intervals = [raw_intervals]
        for interval in raw_intervals:
            if not isinstance(interval, (tuple, list)) or len(interval) != 2:
                raise ValueError(f"Spærringen for {teacher_id} skal have et start- og sluttidspunkt.")
            start = _schedule_time(interval[0], field_name="Starttid for spærring")
            end = _schedule_time(interval[1], field_name="Sluttid for spærring")
            if start >= end:
                raise ValueError(f"Spærringen for {teacher_id} skal slutte efter den starter.")
            result[teacher_id].append((start, end))

    for teacher_id, intervals in result.items():
        merged: list[tuple[dt_time, dt_time]] = []
        for start, end in sorted(intervals):
            if merged and start <= merged[-1][1]:
                merged[-1] = (merged[-1][0], max(merged[-1][1], end))
            else:
                merged.append((start, end))
        result[teacher_id] = merged
    return result


def teacher_blocks_from_table(table: pd.DataFrame) -> dict[str, list[tuple[str, str]]]:
    """Omsæt redigeringstabellen i brugerfladen til spærringer for tidsplanen."""
    blocks: dict[str, list[tuple[str, str]]] = defaultdict(list)
    for _, row in table.iterrows():
        teacher_id = normal_key(row.get("Initialer", ""))
        start = repair_text(row.get("Spærret fra", ""))
        end = repair_text(row.get("Spærret til", ""))
        if not start and not end:
            continue
        label = repair_text(row.get("Lærer", "")) or teacher_id
        if not teacher_id:
            raise ValueError("En lærerspærring mangler initialer.")
        if not start or not end:
            raise ValueError(f"Angiv både fra- og til-tid for spærringen hos {label}.")
        blocks[teacher_id].append((start, end))
    return dict(blocks)


def schedule_time_choices(start_time: dt_time, end_time: dt_time, interval_minutes: int = 5) -> list[dt_time]:
    """Returnér de klokkeslet, der kan vælges for en lærers spærring."""
    if start_time >= end_time:
        return []
    start = datetime.combine(date.today(), start_time).replace(second=0, microsecond=0)
    end = datetime.combine(date.today(), end_time).replace(second=0, microsecond=0)
    choices = [start.time()]
    current = start
    while current + timedelta(minutes=interval_minutes) < end:
        current += timedelta(minutes=interval_minutes)
        choices.append(current.time())
    if choices[-1] != end.time():
        choices.append(end.time())
    return choices


def make_schedule(
    students: list[dict[str, Any]],
    teachers: list[dict[str, Any]],
    solution: dict[str, Any],
    start_time: dt_time,
    end_time: dt_time,
    student_minutes: int,
    pause_count: int,
    pause_minutes: int,
    transition_minutes: int,
    group_pairs: bool,
    lunch_mode: str = "Fast tidspunkt for alle lærere",
    lunch_start_time: dt_time = dt_time(12, 0),
    lunch_minutes: int = 30,
    search_attempts: int = 80,
    progress_callback: Any = None,
    avoid_teacher_gaps: bool = False,
    teacher_blocks: dict[str, Any] | None = None,
    floating_pauses: bool = False,
) -> dict[str, Any]:
    """Lav en parallel vejledningsplan ud fra den aktuelle fordeling."""
    if start_time >= end_time:
        raise ValueError("Sluttidspunktet skal ligge efter starttidspunktet.")
    if student_minutes < 1:
        raise ValueError("Vejledningstiden pr. elev skal være mindst ét minut.")
    if pause_count < 0 or pause_minutes < 0 or transition_minutes < 0:
        raise ValueError("Pauser og skiftetid kan ikke være negative.")
    if lunch_minutes < 1:
        raise ValueError("Frokostpausen skal være mindst ét minut.")
    if lunch_mode not in {"Fast tidspunkt for alle lærere", "Flydende for alle lærere"}:
        raise ValueError("Vælg, om frokostpausen skal være fast eller flydende.")
    search_attempts = max(1, int(search_attempts))
    if progress_callback:
        progress_callback(0.0)

    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    blocked_intervals = normalise_teacher_blocks(teacher_blocks, teachers)
    sessions = []
    for index, student in enumerate(students):
        assigned = solution["assignments"][index]
        if any(not teacher_id for teacher_id in assigned):
            raise ValueError("Fordelingen indeholder elever uden to vejledere. Ret fordelingen, før tidsplanen laves.")
        pair = tuple(sorted(set(assigned)))
        sessions.append({"student": student, "assigned": assigned, "pair": pair, "index": index})
    if not sessions:
        raise ValueError("Der er ingen elever i fordelingen.")

    if group_pairs:
        grouped: dict[tuple[str, ...], list[dict[str, Any]]] = {}
        first_index = {}
        for session in sessions:
            grouped.setdefault(session["pair"], []).append(session)
            first_index.setdefault(session["pair"], session["index"])
        pair_order = sorted(grouped, key=lambda pair: first_index[pair])
        ordered_sessions = [session for pair in pair_order for session in grouped[pair]]
    else:
        ordered_sessions = sessions[:]

    # Tildel runder som en kantfarvning: to vejledninger kan være samtidige,
    # når de ikke deler en lærer. Den tidligere rækkefølge-baserede løsning
    # kunne lægge uafhængige lærerpar efter hinanden og dermed næsten fordoble
    # dagens længde. DSATUR finder i stedet parallelle runder med færrest
    # mulige konflikter og stopper straks, når lærerbelastningens nedre grænse
    # er nået.
    teacher_sessions: dict[str, list[int]] = defaultdict(list)
    for session_index, session in enumerate(ordered_sessions):
        for teacher_id in session["pair"]:
            teacher_sessions[teacher_id].append(session_index)
    conflicts: list[set[int]] = [set() for _ in ordered_sessions]
    for session_indexes in teacher_sessions.values():
        for session_index in session_indexes:
            conflicts[session_index].update(other for other in session_indexes if other != session_index)
    degrees = [len(conflict_set) for conflict_set in conflicts]
    lower_bound = max(len(session_indexes) for session_indexes in teacher_sessions.values())
    best_colors: list[int] | None = None
    best_colour_count = len(ordered_sessions) + 1
    best_gap_score: int | None = None
    best_round_order: list[int] | None = None
    coloring_attempts = max(1, search_attempts)
    for attempt_index in range(coloring_attempts):
        rng = __import__("random").Random(20260913 + attempt_index)
        colors = [-1] * len(ordered_sessions)
        saturation: list[set[int]] = [set() for _ in ordered_sessions]
        uncoloured = set(range(len(ordered_sessions)))
        while uncoloured:
            max_saturation = max(len(saturation[index]) for index in uncoloured)
            candidates = [index for index in uncoloured if len(saturation[index]) == max_saturation]
            max_degree = max(degrees[index] for index in candidates)
            candidates = [index for index in candidates if degrees[index] == max_degree]
            session_index = rng.choice(candidates)
            used_colours = {colors[other] for other in conflicts[session_index] if colors[other] >= 0}
            colour = 0
            while colour in used_colours:
                colour += 1
            colors[session_index] = colour
            uncoloured.remove(session_index)
            for other in conflicts[session_index]:
                if other in uncoloured:
                    saturation[other].add(colour)
        colour_count = max(colors) + 1
        candidate_round_order = list(range(colour_count))
        candidate_gap_score = None
        if avoid_teacher_gaps:
            candidate_round_teachers = [set() for _ in range(colour_count)]
            candidate_round_pairs = [set() for _ in range(colour_count)]
            for session, colour in zip(ordered_sessions, colors):
                candidate_round_teachers[colour].update(session["pair"])
                candidate_round_pairs[colour].add(session["pair"])
            candidate_round_order = _optimise_teacher_round_order(
                candidate_round_teachers,
                candidate_round_order,
                candidate_round_pairs if group_pairs else None,
            )
            candidate_gap_score, _ = _round_order_metrics(
                candidate_round_order,
                candidate_round_teachers,
                candidate_round_pairs if group_pairs else None,
            )

        is_better = colour_count < best_colour_count
        if avoid_teacher_gaps and colour_count == best_colour_count:
            is_better = best_gap_score is None or candidate_gap_score < best_gap_score
        if is_better:
            best_colors = colors
            best_colour_count = colour_count
            best_gap_score = candidate_gap_score
            best_round_order = candidate_round_order
        if progress_callback:
            progress_callback(0.1 + 0.8 * (attempt_index + 1) / coloring_attempts)
        if not avoid_teacher_gaps and best_colour_count <= lower_bound:
            break

    # En ombytning af hele runder ændrer ikke lærer-konflikterne. Når
    # lærerpar-optimering er valgt, bruges det derfor som et sekundært mål:
    # runder med de samme lærerpar placeres så vidt muligt ved siden af
    # hinanden, uden at antallet af runder eller paralleliteten ændres.
    round_order = best_round_order or list(range(best_colour_count))
    if not avoid_teacher_gaps and group_pairs and best_colors is not None and best_colour_count > 1:
        pairs_in_round = [set() for _ in range(best_colour_count)]
        for session, colour in zip(ordered_sessions, best_colors):
            pairs_in_round[colour].add(session["pair"])
        affinity = [
            [len(pairs_in_round[left] & pairs_in_round[right]) for right in range(best_colour_count)]
            for left in range(best_colour_count)
        ]

        def order_score(order: list[int]) -> int:
            return sum(affinity[left][right] for left, right in zip(order, order[1:]))

        best_order = round_order
        best_order_score = -1
        for start in round_order:
            candidate_order = [start]
            remaining = set(round_order)
            remaining.remove(start)
            while remaining:
                previous = candidate_order[-1]
                best_affinity = max(affinity[previous][candidate] for candidate in remaining)
                next_round = min(candidate for candidate in remaining if affinity[previous][candidate] == best_affinity)
                candidate_order.append(next_round)
                remaining.remove(next_round)
            candidate_score = order_score(candidate_order)
            if candidate_score > best_order_score:
                best_order = candidate_order
                best_order_score = candidate_score
        improved = True
        while improved:
            improved = False
            for left in range(len(best_order) - 1):
                for right in range(left + 1, len(best_order)):
                    candidate_order = best_order[:]
                    candidate_order[left], candidate_order[right] = candidate_order[right], candidate_order[left]
                    if order_score(candidate_order) > best_order_score:
                        best_order = candidate_order
                        best_order_score = order_score(candidate_order)
                        improved = True
        round_order = best_order

    # Rækkefølgen af de farvede runder kan ændres uden at skabe en
    # dobbeltbooking. Når der er lærerspærringer, vælger vi derfor først de
    # runder, der faktisk kan begynde nu. Det forhindrer, at en spærring for
    # én lærer unødigt forsinker helt uafhængige lærerpar.
    if any(blocked_intervals.values()) and best_colors is not None:
        teachers_in_colour = [set() for _ in range(best_colour_count)]
        for session, colour in zip(ordered_sessions, best_colors):
            teachers_in_colour[colour].update(session["pair"])

        def first_unblocked_start(candidate: datetime, teacher_ids: set[str]) -> datetime:
            while True:
                candidate_end = candidate + timedelta(minutes=student_minutes)
                ends = [
                    datetime.combine(candidate.date(), blocked_end)
                    for teacher_id in teacher_ids
                    for blocked_start, blocked_end in blocked_intervals.get(teacher_id, [])
                    if candidate < datetime.combine(candidate.date(), blocked_end)
                    and candidate_end > datetime.combine(candidate.date(), blocked_start)
                ]
                if not ends:
                    return candidate
                candidate = max(ends)

        ordering_cursor = datetime.combine(date.today(), start_time)
        remaining_colours = set(round_order)
        reordered_rounds = []
        while remaining_colours:
            candidates = [
                (first_unblocked_start(ordering_cursor, teachers_in_colour[colour]), round_order.index(colour), colour)
                for colour in remaining_colours
            ]
            chosen_start, _, chosen_colour = min(candidates)
            reordered_rounds.append(chosen_colour)
            remaining_colours.remove(chosen_colour)
            ordering_cursor = chosen_start + timedelta(minutes=student_minutes + transition_minutes)
        round_order = reordered_rounds

    colour_to_round = {colour: round_index for round_index, colour in enumerate(round_order)}
    for session, colour in zip(ordered_sessions, best_colors or []):
        session["round"] = colour_to_round[colour]
    last_round = max(session["round"] for session in ordered_sessions)
    final_round_teachers = [set() for _ in range(last_round + 1)]
    for session in ordered_sessions:
        final_round_teachers[session["round"]].update(session["pair"])
    teacher_gap_rounds, _ = _round_order_metrics(list(range(last_round + 1)), final_round_teachers)
    floating_lunch_round = max(1, (last_round + 1) // 2)
    if lunch_mode == "Fast tidspunkt for alle lærere":
        lunch_reference = datetime.combine(date.today(), lunch_start_time)
        day_reference = datetime.combine(date.today(), start_time)
        estimated_lunch_round = int((lunch_reference - day_reference).total_seconds() // 60) // max(1, student_minutes + transition_minutes)
        lunch_pause_round = max(1, min(last_round, estimated_lunch_round))
    else:
        lunch_pause_round = floating_lunch_round
    actual_pause_count = min(pause_count, last_round)
    global_pause_count = actual_pause_count

    def evenly_spaced_positions(first: int, last: int, count: int) -> list[int]:
        if count <= 0 or first > last:
            return []
        candidates = list(range(first, last + 1))
        positions = [
            candidates[max(0, min(len(candidates) - 1, round((index + 1) * (len(candidates) + 1) / (count + 1)) - 1))]
            for index in range(count)
        ]
        return list(dict.fromkeys(positions))

    if floating_pauses:
        # Med to pauser bliver der én i hver side af frokosten. Placeringerne
        # er altid efter en runde og før en efterfølgende runde.
        before_count = (global_pause_count + 1) // 2
        after_count = global_pause_count - before_count
        pause_after = set(evenly_spaced_positions(1, lunch_pause_round, before_count))
        pause_after.update(evenly_spaced_positions(lunch_pause_round + 1, last_round, after_count))
        remaining_positions = [position for position in range(1, last_round + 1) if position not in pause_after]
        while len(pause_after) < global_pause_count and remaining_positions:
            pause_after.add(remaining_positions.pop(0))
    else:
        pause_after = set(evenly_spaced_positions(1, last_round, global_pause_count))

    plan_start = datetime.combine(date.today(), start_time)
    day_end = datetime.combine(date.today(), end_time)
    lunch_start_dt = datetime.combine(date.today(), lunch_start_time)
    lunch_end_dt = lunch_start_dt + timedelta(minutes=lunch_minutes)
    if lunch_mode == "Fast tidspunkt for alle lærere" and (lunch_start_dt < plan_start or lunch_end_dt > day_end):
        raise ValueError("Den faste frokostpause skal ligge inden for tidsrummet og kunne rumme hele pausen.")
    floating_lunch_round = max(1, (last_round + 1) // 2)
    lunch_interval: tuple[datetime, datetime] | None = None
    lunch_added = False
    round_starts = []
    timeline_rows = []
    regular_pause_intervals: list[tuple[datetime, datetime]] = []
    cursor = plan_start

    def next_available_round_start(candidate: datetime, round_teachers: set[str]) -> datetime:
        """Flyt en runde til efter alle berørte læreres spærringer.

        Runderne bevarer deres indbyrdes rækkefølge. Derfor kan en spærring
        skabe et hul i den fælles plan, men den kan aldrig føre til, at en
        lærer bliver booket i sit spærrede tidsrum.
        """
        while True:
            session_end = candidate + timedelta(minutes=student_minutes)
            overlaps = []
            for teacher_id in round_teachers:
                for blocked_start, blocked_end in blocked_intervals.get(teacher_id, []):
                    start_dt = datetime.combine(candidate.date(), blocked_start)
                    end_dt = datetime.combine(candidate.date(), blocked_end)
                    if candidate < end_dt and session_end > start_dt:
                        overlaps.append(end_dt)
            if not overlaps:
                return candidate
            candidate = max(overlaps)

    for round_number in range(last_round + 1):
        round_teachers = final_round_teachers[round_number]
        # En fast frokostpause og en individuel spærring kan begge forskyde en
        # runde. Gentag derfor kontrollen, til tidspunktet er lovligt for dem
        # begge – også når en spærring netop flytter en runde hen over frokost.
        while True:
            if not lunch_added:
                if lunch_mode == "Flydende for alle lærere":
                    add_lunch = round_number == floating_lunch_round
                else:
                    round_end = cursor + timedelta(minutes=student_minutes)
                    add_lunch = cursor >= lunch_start_dt or round_end > lunch_start_dt
                if add_lunch:
                    if lunch_mode == "Fast tidspunkt for alle lærere":
                        pause_start, pause_end = lunch_start_dt, lunch_end_dt
                        cursor = max(cursor, pause_end)
                    else:
                        pause_start = cursor
                        pause_end = cursor + timedelta(minutes=lunch_minutes)
                        cursor = pause_end
                    lunch_interval = (pause_start, pause_end)
                    timeline_rows.append({
                        "_sort": (round_number, -1), "Type": "Frokostpause", "Start": _clock_label(pause_start),
                        "Slut": _clock_label(pause_end), "Varighed (min.)": lunch_minutes, "Elev": "", "Klasse": "",
                        "Lærer(e)": "Alle lærere", "Information": "Frokostpause", "Bemærkning": "",
                    })
                    lunch_added = True
                    continue
            available_start = next_available_round_start(cursor, round_teachers)
            if lunch_mode == "Fast tidspunkt for alle lærere" and not lunch_added:
                # Hvis spærringen flyttede starten frem over frokost, lægges den
                # faste pause ind, før runden planlægges.
                if available_start >= lunch_start_dt or available_start + timedelta(minutes=student_minutes) > lunch_start_dt:
                    cursor = available_start
                    continue
            cursor = available_start
            break
        round_starts.append(cursor)
        round_end = cursor + timedelta(minutes=student_minutes)
        if round_number < last_round:
            next_start = round_end + timedelta(minutes=transition_minutes)
            if round_number + 1 in pause_after:
                pause_end = next_start + timedelta(minutes=pause_minutes)
                regular_pause_intervals.append((next_start, pause_end))
                timeline_rows.append({
                    "_sort": (round_number, 1), "Type": "Pause", "Start": _clock_label(next_start),
                    "Slut": _clock_label(pause_end), "Varighed (min.)": pause_minutes, "Elev": "", "Klasse": "",
                    "Lærer(e)": "", "Information": "Pause", "Bemærkning": "",
                })
                next_start = pause_end
            cursor = next_start
        else:
            cursor = round_end
    if not lunch_added:
        if lunch_mode == "Fast tidspunkt for alle lærere":
            final_lunch_start = lunch_start_dt
            final_lunch_end = lunch_end_dt
        else:
            final_lunch_start = cursor
            final_lunch_end = cursor + timedelta(minutes=lunch_minutes)
        lunch_interval = (final_lunch_start, final_lunch_end)
        timeline_rows.append({
            "_sort": (last_round + 1, -1), "Type": "Frokostpause", "Start": _clock_label(final_lunch_start),
            "Slut": _clock_label(final_lunch_end), "Varighed (min.)": lunch_minutes, "Elev": "", "Klasse": "",
            "Lærer(e)": "Alle lærere", "Information": "Frokostpause", "Bemærkning": "",
        })
        cursor = max(cursor, final_lunch_end)
    if cursor > day_end:
        required = int((cursor - plan_start).total_seconds() // 60)
        available = int((day_end - plan_start).total_seconds() // 60)
        rounds = last_round + 1
        guidance_total = rounds * student_minutes
        transition_total = last_round * transition_minutes
        regular_pause_total = global_pause_count * pause_minutes
        lunch_total = lunch_minutes if lunch_interval is not None else 0
        accounted_total = guidance_total + transition_total + regular_pause_total + lunch_total
        extra_idle_minutes = max(0, required - accounted_total)
        teacher_loads = Counter(
            teacher_id
            for session in sessions
            for teacher_id in session["pair"]
        )
        busiest = sorted(teacher_loads.items(), key=lambda item: (-item[1], item[0]))[:3]
        busiest_text = ", ".join(
            f"{teacher_label(teacher_id, teacher_map)} ({load} elev.)"
            for teacher_id, load in busiest
        )
        minimum_end = _clock_label(plan_start + timedelta(minutes=required))
        explanation = [
            f"{len(sessions)} elever er fordelt på {rounds} vejledningsrunder, fordi den samme lærer ikke kan vejlede to elever samtidig.",
            f"Selve vejledningerne bruger {rounds} × {student_minutes} minutter = {guidance_total} minutter.",
            f"Skift mellem runder bruger {last_round} × {transition_minutes} minutter = {transition_total} minutter.",
            f"De almindelige pauser bruger {global_pause_count} × {pause_minutes} minutter = {regular_pause_total} minutter i den fælles tidsplan.",
            f"Frokostpausen bruger {lunch_total} minutter og er sat til {lunch_mode.casefold()}.",
            f"De mest belastede lærere er {busiest_text}.",
        ]
        if extra_idle_minutes:
            if lunch_mode == "Fast tidspunkt for alle lærere":
                explanation.append(
                    f"Den faste frokosttid giver desuden {extra_idle_minutes} minutters nødvendig tidsjustering, så ingen vejledning ligger i frokostpausen."
                )
            else:
                active_blocks = [
                    f"{teacher_label(teacher_id, teacher_map)}: {_clock_label(datetime.combine(date.today(), start))}–{_clock_label(datetime.combine(date.today(), end))}"
                    for teacher_id, intervals in blocked_intervals.items()
                    for start, end in intervals
                ]
                if active_blocks:
                    explanation.append(
                        f"Lærerspærringer giver desuden {extra_idle_minutes} minutters nødvendig ventetid. "
                        f"Aktive spærringer: {'; '.join(active_blocks[:8])}."
                    )
                else:
                    explanation.append(f"Planens rundeplacering giver desuden {extra_idle_minutes} minutters nødvendig ventetid.")
        sessions_after_day = [
            session for session in ordered_sessions
            if round_starts[session["round"]] + timedelta(minutes=student_minutes) > day_end
        ]
        if any(blocked_intervals.values()) and sessions_after_day:
            affected_students = list(dict.fromkeys(session["student"]["name"] for session in sessions_after_day))
            preview = ", ".join(affected_students[:8])
            suffix = " …" if len(affected_students) > 8 else ""
            explanation.append(
                f"{len(sessions_after_day)} vejledning(er) for {len(affected_students)} elev(er) falder efter dagens sluttid på grund af lærerspærringerne. "
                f"De resterende skal lægges en anden dag, hvis spærringerne og sluttiden fastholdes: {preview}{suffix}."
            )
        raise ValueError(
            f"Tidsplanen kan ikke afsluttes kl. {_clock_label(day_end)}. "
            f"Med de valgte indstillinger skal den mindst afsluttes kl. {minimum_end}; "
            f"tidsrummet fra kl. {_clock_label(plan_start)} til kl. {_clock_label(day_end)} er kun {available} minutter.\n\n"
            + "\n".join(f"• {line}" for line in explanation)
            + f"\n\nDer mangler derfor {required - available} minutter. Forlæng sluttiden, reducér elevtid, skiftetid eller almindelige pauser, eller justér lærerspærringerne."
        )

    session_rows = []
    teacher_rows = []
    pair_counts = Counter()
    pair_first_last: dict[tuple[str, ...], list[str]] = {}
    for session in ordered_sessions:
        session_start = round_starts[session["round"]]
        session_end = session_start + timedelta(minutes=student_minutes)
        student = session["student"]
        assigned = session["assigned"]
        pair = session["pair"]
        pair_label = _schedule_pair_label(pair, teacher_map)
        time_label = f"{_clock_label(session_start)}–{_clock_label(session_end)}"
        session_rows.append({
            "Start": _clock_label(session_start), "Slut": _clock_label(session_end), "Tid": time_label,
            "Elev": student["name"], "Klasse": student.get("className", ""), "Fag 1": student["subjects"][0],
            "Vejleder 1": teacher_label(assigned[0], teacher_map), "Fag 2": student["subjects"][1],
            "Vejleder 2": teacher_label(assigned[1], teacher_map), "Lærerpar": pair_label,
            "Projekttitel": student.get("projectTitle", ""),
            "Bemærkninger/ændringer fra lærerne": student.get("scheduleNotes", ""),
        })
        timeline_rows.append({
            "_sort": (session["round"], 0), "Type": "Vejledning", "Start": _clock_label(session_start),
            "Slut": _clock_label(session_end), "Varighed (min.)": student_minutes, "Elev": student["name"],
            "Klasse": student.get("className", ""), "Lærer(e)": pair_label,
            "Information": student.get("projectTitle", "") or "SOP-vejledning",
            "Bemærkning": student.get("scheduleNotes", ""),
        })
        pair_counts[pair] += 1
        pair_first_last.setdefault(pair, [time_label, time_label])
        pair_first_last[pair][1] = time_label
        for teacher_id in dict.fromkeys(assigned):
            subject_slots = [subject for slot, subject in enumerate(student["subjects"][:2]) if assigned[slot] == teacher_id]
            co_teachers = list(dict.fromkeys(other_id for other_id in assigned if other_id != teacher_id))
            teacher_rows.append({
                "Type": "Vejledning",
                "Start": _clock_label(session_start), "Slut": _clock_label(session_end), "Tid": time_label,
                "Lærer": teacher_label(teacher_id, teacher_map), "Initialer": teacher_id, "Elev": student["name"],
                "Klasse": student.get("className", ""), "Fag": " · ".join(dict.fromkeys(subject_slots)),
                "Medvejleder": " + ".join(teacher_label(other_id, teacher_map) for other_id in co_teachers) or "—",
                "Projekttitel": student.get("projectTitle", ""),
                "Bemærkning": student.get("scheduleNotes", ""),
            })
    if lunch_interval is not None:
        pause_start, pause_end = lunch_interval
        for teacher in teachers:
            teacher_rows.append({
                "Type": "Frokostpause", "Start": _clock_label(pause_start), "Slut": _clock_label(pause_end),
                "Tid": f"{_clock_label(pause_start)}–{_clock_label(pause_end)}", "Lærer": teacher_label(teacher["id"], teacher_map),
                "Initialer": teacher["id"], "Elev": "", "Klasse": "", "Fag": "Frokostpause",
                "Medvejleder": "—", "Projekttitel": "", "Bemærkning": "Fælles frokostpause",
            })
    for pause_start, pause_end in regular_pause_intervals:
        for teacher in teachers:
            teacher_rows.append({
                "Type": "Pause", "Start": _clock_label(pause_start), "Slut": _clock_label(pause_end),
                "Tid": f"{_clock_label(pause_start)}–{_clock_label(pause_end)}", "Lærer": teacher_label(teacher["id"], teacher_map),
                "Initialer": teacher["id"], "Elev": "", "Klasse": "", "Fag": "Pause", "Medvejleder": "—",
                "Projekttitel": "", "Bemærkning": "Fælles planlagt pause",
            })
    for teacher_id, intervals in blocked_intervals.items():
        for blocked_start, blocked_end in intervals:
            teacher_rows.append({
                "Type": "Spærret", "Start": _clock_label(datetime.combine(date.today(), blocked_start)),
                "Slut": _clock_label(datetime.combine(date.today(), blocked_end)),
                "Tid": f"{_clock_label(datetime.combine(date.today(), blocked_start))}–{_clock_label(datetime.combine(date.today(), blocked_end))}",
                "Lærer": teacher_label(teacher_id, teacher_map), "Initialer": teacher_id,
                "Elev": "", "Klasse": "", "Fag": "Ikke tilgængelig", "Medvejleder": "—",
                "Projekttitel": "", "Bemærkning": "Lærerens angivne spærring",
            })
    timeline_rows.sort(key=lambda row: row["_sort"])
    for row in timeline_rows:
        row.pop("_sort", None)
    session_rows.sort(key=lambda row: row["Start"])
    teacher_rows.sort(key=lambda row: (row["Start"], row["Lærer"]))

    student_columns = ["Start", "Slut", "Tid", "Elev", "Klasse", "Fag 1", "Vejleder 1", "Fag 2", "Vejleder 2", "Lærerpar", "Projekttitel", "Bemærkninger/ændringer fra lærerne"]
    teacher_columns = ["Type", "Start", "Slut", "Tid", "Lærer", "Initialer", "Elev", "Klasse", "Fag", "Medvejleder", "Projekttitel", "Bemærkning"]
    timeline_columns = ["Type", "Start", "Slut", "Varighed (min.)", "Elev", "Klasse", "Lærer(e)", "Information", "Bemærkning"]
    pair_rows = [
        {"Lærerpar": _schedule_pair_label(pair, teacher_map), "Antal elever": count, "Tidsblok": " → ".join(pair_first_last[pair])}
        for pair, count in sorted(pair_counts.items(), key=lambda item: pair_first_last[item[0]][0])
    ]
    total_minutes = int((cursor - plan_start).total_seconds() // 60)
    if progress_callback:
        progress_callback(1.0)
    return {
        "students": pd.DataFrame(session_rows, columns=student_columns),
        "teachers": pd.DataFrame(teacher_rows, columns=teacher_columns),
        "timeline": pd.DataFrame(timeline_rows, columns=timeline_columns),
        "pairs": pd.DataFrame(pair_rows, columns=["Lærerpar", "Antal elever", "Tidsblok"]),
        "settings": {
            "Start": _clock_label(datetime.combine(date.today(), start_time)),
            "Slut": _clock_label(datetime.combine(date.today(), end_time)),
            "Minutter pr. elev": student_minutes,
            "Antal pauser": actual_pause_count,
            "Minutter pr. pause": pause_minutes,
            "Pauser fordeles om frokost": "Ja" if floating_pauses else "Nej",
            "Minutter mellem elever": transition_minutes,
            "Frokostpause (min.)": lunch_minutes,
            "Frokostpause": lunch_mode,
            "Fast frokosttid": _clock_label(lunch_start_dt) if lunch_mode == "Fast tidspunkt for alle lærere" else "Flydende",
            "Lærerpar samlet": "Ja" if group_pairs else "Nej",
            "Lærerspærringer": sum(len(intervals) for intervals in blocked_intervals.values()),
            "Forsøg at undgå lærerhuller": "Ja" if avoid_teacher_gaps else "Nej",
            "Lærerhuller (runder)": teacher_gap_rounds,
            "Optimeringsdybde": search_attempts,
            "Planlagt tidsforbrug (min.)": total_minutes,
        },
    }


def make_teacher_gantt_html(schedule: dict[str, Any], selected_teacher: str | None = None) -> str:
    """Lav en kompakt, visuel Gantt-plan for én eller alle lærere."""
    frame = schedule["teachers"]
    if frame.empty:
        return '<div class="gantt-empty">Der er ingen vejledninger at vise.</div>'

    def clock_minutes(value: Any) -> int:
        parsed = datetime.strptime(str(value), "%H:%M")
        return parsed.hour * 60 + parsed.minute

    settings = schedule["settings"]
    plan_start = clock_minutes(settings["Start"])
    plan_end = clock_minutes(settings["Slut"])
    visible_end = max(plan_end, max(clock_minutes(value) for value in frame["Slut"]))
    total_minutes = max(1, visible_end - plan_start)
    names = sorted(frame["Lærer"].dropna().unique().tolist(), key=str.casefold)
    if selected_teacher and selected_teacher != "Alle lærere":
        names = [selected_teacher] if selected_teacher in names else []

    palette = ["#0f817a", "#df8b2d", "#b44438", "#5367a8", "#8d5aa7", "#4f8b69", "#bd5f83"]

    def position(value: int) -> float:
        return max(0.0, min(100.0, (value - plan_start) / total_minutes * 100))

    ticks = list(range(plan_start, visible_end + 1, 30))
    if ticks[-1] != visible_end:
        ticks.append(visible_end)
    tick_markup = "".join(
        f'<span class="gantt-tick" style="left:{position(tick):.3f}%">{tick // 60:02d}:{tick % 60:02d}</span>'
        for tick in ticks
    )
    rows = []
    for name in names:
        teacher_rows = frame[frame["Lærer"] == name].sort_values(["Start", "Slut"])
        guidance_count = int((teacher_rows["Type"] == "Vejledning").sum()) if "Type" in teacher_rows else len(teacher_rows)
        blocks = []
        for _, row in teacher_rows.iterrows():
            if str(row.get("Type", "")) == "Pause":
                continue
            start = clock_minutes(row["Start"])
            end = clock_minutes(row["Slut"])
            is_lunch = str(row.get("Type", "")) == "Frokostpause"
            is_blocked = str(row.get("Type", "")) == "Spærret"
            is_pause = str(row.get("Type", "")) == "Pause"
            subject = "Frokostpause" if is_lunch else "Pause" if is_pause else "Ikke tilgængelig" if is_blocked else str(row.get("Fag", "Vejledning"))
            student = "Frokostpause" if is_lunch else "Pause" if is_pause else "Spærret" if is_blocked else str(row.get("Elev", "Elev"))
            co_teacher = "" if is_lunch or is_pause or is_blocked else str(row.get("Medvejleder", "")).strip(" —")
            color = "#71858a" if is_lunch else "#a57b28" if is_pause else "#9a5b63" if is_blocked else palette[sum(ord(char) for char in subject) % len(palette)]
            detail = f" · Medvejleder: {co_teacher}" if co_teacher else ""
            tooltip = html.escape(f"{row['Start']}–{row['Slut']} · {student} · {subject}{detail}", quote=True)
            co_teacher_html = f"<small>Medvejleder: {html.escape(co_teacher)}</small>" if co_teacher else ""
            blocks.append(
                f'<div class="gantt-block" title="{tooltip}" style="left:{position(start):.3f}%;width:{max(0.8, position(end) - position(start)):.3f}%;background:{color}">'
                f'<span>{html.escape(row["Start"])}–{html.escape(row["Slut"])}</span>'
                f'<strong>{html.escape(student)}</strong><small>{html.escape(subject)}</small>{co_teacher_html}</div>'
            )
        teacher_initials = " ".join(teacher_rows["Initialer"].dropna().astype(str).unique().tolist())
        search_terms = normal_key(f"{name} {teacher_initials}")
        rows.append(
            f'<div class="gantt-row" data-teacher="{html.escape(search_terms, quote=True)}"><button type="button" class="gantt-name teacher-link" data-teacher-select="{html.escape(search_terms, quote=True)}">{html.escape(name)}<small>{guidance_count} elev(er)</small></button>'
            f'<div class="gantt-track">{"".join(blocks)}</div></div>'
        )

    if not rows:
        return '<div class="gantt-empty">Vælg en lærer med planlagte vejledninger.</div>'
    return f'''<div class="teacher-gantt">
<style>
.teacher-gantt{{border:1px solid #dce4e2;border-radius:14px;background:#fff;padding:16px;overflow:auto;box-shadow:0 8px 24px #1933300b}}
.gantt-intro{{display:flex;justify-content:space-between;gap:16px;align-items:flex-end;margin-bottom:14px;color:#66777f;font-size:13px}}
.gantt-intro strong{{display:block;color:#172126;font:700 21px Georgia,serif;margin-bottom:3px}}
.gantt-scale{{margin-left:178px;height:28px;min-width:760px;position:relative;border-bottom:1px solid #dce4e2}}
.gantt-tick{{position:absolute;bottom:5px;transform:translateX(-50%);font-size:11px;color:#66777f;white-space:nowrap}}
.gantt-row{{display:flex;min-width:938px;min-height:82px;border-bottom:1px solid #edf1f0}}
.gantt-row:last-child{{border-bottom:0}}
.gantt-name{{width:162px;flex:0 0 162px;padding:15px 12px 8px 0;font-weight:700;color:#172126;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}}
.gantt-name small{{display:block;color:#66777f;font-size:11px;font-weight:400;margin-top:3px}}
.gantt-track{{position:relative;flex:1;margin:9px 0;background:repeating-linear-gradient(to right,#f4f8f7 0,#f4f8f7 calc(8.333% - 1px),#dfe9e6 calc(8.333% - 1px),#dfe9e6 8.333%);border-radius:8px;min-height:64px}}
.gantt-block{{position:absolute;top:5px;height:54px;border-radius:7px;padding:4px 7px;box-sizing:border-box;color:white;overflow:hidden;box-shadow:0 2px 5px #1721262b;line-height:1.15;font-size:10px;min-width:25px}}
.gantt-block span,.gantt-block strong,.gantt-block small{{display:block;overflow:hidden;text-overflow:ellipsis;white-space:nowrap}}
.gantt-block strong{{font-size:11px;margin-top:2px}}
.gantt-block small{{opacity:.88;margin-top:2px}}
</style>
<div class="gantt-intro"><div><strong>Visuel lærerplan</strong><span>Blokkene viser elevens vejledningstid. Hold musen over en blok for detaljer.</span></div><span>{html.escape(str(settings["Start"]))}–{html.escape(str(settings["Slut"]))}</span></div>
<div class="gantt-scale">{tick_markup}</div>
{"".join(rows)}
</div>'''


def excel_sheet_name(value: Any, used_names: set[str]) -> str:
    """Lav et kort, gyldigt og unikt Excel-fanenavn."""
    base = repair_text(value)
    for invalid_character in ("[", "]", ":", "*", "?", "/", "\\"):
        base = base.replace(invalid_character, " ")
    base = " ".join(base.split()).strip(" '") or "Lærer"
    base = base[:31]
    candidate = base
    suffix_number = 2
    while candidate.casefold() in used_names:
        suffix = f" ({suffix_number})"
        candidate = base[: 31 - len(suffix)].rstrip() + suffix
        suffix_number += 1
    used_names.add(candidate.casefold())
    return candidate


def make_schedule_excel(schedule: dict[str, Any]) -> bytes:
    teacher_plan = schedule["teachers"].copy()
    if not teacher_plan.empty:
        teacher_plan["_sort_teacher"] = teacher_plan["Lærer"].map(normal_key)
        teacher_plan = teacher_plan.sort_values(["_sort_teacher", "Start", "Slut"], kind="stable").drop(columns="_sort_teacher")
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        safe_excel_frame(schedule["students"]).to_excel(writer, index=False, sheet_name="Elevplan")
        safe_excel_frame(teacher_plan).to_excel(writer, index=False, sheet_name="Lærerplan")
        used_sheet_names = {"elevplan", "lærerplan", "tidslinje", "lærerpar", "indstillinger"}
        for teacher_name in sorted(teacher_plan["Lærer"].dropna().unique().tolist(), key=normal_key):
            personal_plan = teacher_plan[teacher_plan["Lærer"] == teacher_name].sort_values(["Start", "Slut"], kind="stable")
            safe_excel_frame(personal_plan).to_excel(
                writer,
                index=False,
                sheet_name=excel_sheet_name(teacher_name, used_sheet_names),
            )
        safe_excel_frame(schedule["timeline"]).to_excel(writer, index=False, sheet_name="Tidslinje")
        safe_excel_frame(schedule["pairs"]).to_excel(writer, index=False, sheet_name="Lærerpar")
        safe_excel_frame(pd.DataFrame([schedule["settings"]])).to_excel(writer, index=False, sheet_name="Indstillinger")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                sheet.column_dimensions[column[0].column_letter].width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    return output.getvalue()


def make_schedule_html(schedule: dict[str, Any]) -> str:
    settings = schedule["settings"]
    setting_lines = "".join(f"<li><strong>{html.escape(str(key))}:</strong> {html.escape(str(value))}</li>" for key, value in settings.items())
    gantt_html = make_teacher_gantt_html(schedule, "Alle lærere")
    teacher_frame = schedule["teachers"].sort_values(["Lærer", "Start", "Slut"], kind="stable")
    teacher_frame = teacher_frame[teacher_frame["Type"] == "Vejledning"]
    teacher_columns = ["Tid", "Elev", "Klasse", "Fag", "Medvejleder", "Projekttitel", "Bemærkning"]
    teacher_headers = "".join(f"<th>{html.escape(str(column))}</th>" for column in teacher_columns)
    teacher_agenda_rows = []
    for row in teacher_frame.fillna("").to_dict("records"):
        search_terms = normal_key(f"{row.get('Lærer', '')} {row.get('Initialer', '')}")
        cells = "".join(f"<td>{html.escape(str(row.get(column, '')))}</td>" for column in teacher_columns)
        teacher_agenda_rows.append(f'<tr data-teacher="{html.escape(search_terms, quote=True)}">{cells}</tr>')
    return f"""<!doctype html>
<html lang="da">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Vejledningsplan</title>
<style>
body{{margin:0;background:#f7fbf9;color:#172126;font:15px/1.45 Arial,sans-serif}}
main{{max-width:1400px;margin:auto;padding:28px 20px 60px}}
h1,h2{{font-family:Georgia,serif}} h1{{font-size:38px;margin:0 0 8px}} h2{{margin-top:30px}}
.intro,.card{{background:white;border:1px solid #dce4e2;border-radius:14px;padding:18px;margin:14px 0;box-shadow:0 8px 24px #1933300b}}
.intro p{{margin:4px 0;color:#66777f}} ul{{display:flex;flex-wrap:wrap;gap:8px 28px;padding-left:20px}}
.table-wrap{{overflow:auto;background:white;border:1px solid #dce4e2;border-radius:12px}}
table{{border-collapse:collapse;width:100%;min-width:760px}} th,td{{padding:9px 11px;border-bottom:1px solid #e2e9e7;text-align:left;vertical-align:top}} th{{background:#e7f4f1;color:#075f5b;white-space:nowrap}} tr:last-child td{{border-bottom:0}}
.teacher-search{{box-sizing:border-box;width:min(520px,100%);padding:11px 13px;border:1px solid #9bb5b1;border-radius:8px;font:inherit;margin:0 0 16px}} .teacher-agenda{{margin-top:22px}} .teacher-agenda h3{{font-family:Georgia,serif;margin:0 0 9px}} #teacher-empty{{color:#66777f}}
.teacher-link{{border:0;background:transparent;text-align:left;cursor:pointer;font:inherit}}
.copyright{{margin:36px 0 0;text-align:center;color:#66777f;font-size:12px}}
@media print{{body{{background:white}} main{{padding:0}} .card,.table-wrap{{box-shadow:none}} h2{{break-before:page}}}}
</style></head>
<body><main>
<section class="intro"><h1>Vejledningsplan</h1><p>Detaljeret tidsplan for elever og lærere baseret på den beregnede SOP-fordeling. Fordelingen er foretaget ud fra SOPtima – en algoritme udviklet af Henrik (hst@nextkhb.dk). Stadig i beta, så skriv gerne, hvis der er mangler eller fejl.</p><ul>{setting_lines}</ul></section>
<section class="card"><h2>Vejledningsplan for lærere</h2><label for="teacher-query">Søg på lærerens navn eller initialer</label><input class="teacher-search" id="teacher-query" type="search" autocomplete="off" placeholder="Fx Anne eller AB"><p id="teacher-prompt">Skriv lærerens navn eller initialer for at se den specifikke plan.</p>{gantt_html}
<section class="teacher-agenda" id="teacher-agenda" hidden><h3>Elever med vejledning hos den valgte lærer</h3><div class="table-wrap"><table><thead><tr>{teacher_headers}</tr></thead><tbody id="teacher-agenda-rows">{''.join(teacher_agenda_rows)}</tbody></table></div><p id="teacher-empty" hidden>Ingen lærer matcher søgningen.</p></section></section>
<div id="student-data" hidden aria-hidden="true">{schedule["students"].to_html(index=False, escape=True, border=0)}</div>
<footer class="copyright">{html.escape(COPYRIGHT)}</footer>
<script>
const teacherQuery=document.getElementById('teacher-query');
const ganttRows=[...document.querySelectorAll('.gantt-row')], agendaRows=[...document.querySelectorAll('#teacher-agenda-rows tr')], teacherAgenda=document.getElementById('teacher-agenda'), teacherEmpty=document.getElementById('teacher-empty');
const teacherPrompt=document.getElementById('teacher-prompt');
function normaliseTeacher(value){{return String(value||'').toLocaleLowerCase('da-DK').trim().replace(/\\s+/g,' ');}}
function filterTeachers(){{const term=normaliseTeacher(teacherQuery.value);let matches=0;ganttRows.forEach(row=>{{const visible=!term||row.dataset.teacher.includes(term);row.hidden=!visible;row.style.display=visible?'':'none';}});agendaRows.forEach(row=>{{const visible=Boolean(term)&&row.dataset.teacher.includes(term);row.hidden=!visible;row.style.display=visible?'':'none';if(visible)matches++;}});teacherAgenda.hidden=!term;teacherEmpty.hidden=matches!==0;teacherPrompt.hidden=Boolean(term);}}
teacherQuery.addEventListener('input',filterTeachers);
teacherQuery.addEventListener('change',filterTeachers);
document.querySelectorAll('.teacher-link').forEach(link=>link.addEventListener('click',()=>{{teacherQuery.value=link.dataset.teacherSelect||'';filterTeachers();teacherQuery.focus();}}));
filterTeachers();
</script></main></body></html>"""


def compact_search_key(value: Any) -> str:
    """Søgenøgle uden tegn, punktummer eller mellemrum (fx 3.q == 3q)."""
    text = normal_key(value)
    text = "".join(char for char in unicodedata.normalize("NFD", text) if unicodedata.category(char) != "Mn")
    return "".join(char for char in text if char.isalnum())


def class_search_keys(value: Any) -> list[str]:
    """Returnér klassens egen nøgle og elevvenlige aliaser (fx 3q/S 2024q)."""
    key = compact_search_key(value)
    keys = [key] if key else []
    match = re.fullmatch(r"s(20\d{2})([a-z])", key)
    if match:
        # S 2024q er 3q i skoleåret 2026; beregn klassetrinnet ud fra
        # startåret, så samme logik også virker for senere årgange.
        grade = max(1, date.today().year - int(match.group(1)) + 1)
        keys.append(f"{grade}{match.group(2)}")
    return list(dict.fromkeys(keys))


def make_student_schedule_html(schedule: dict[str, Any]) -> str:
    """Lav en selvstændig elevside, der kan søges lokalt uden login eller server."""
    columns = ["Tid", "Elev", "Klasse", "Fag 1", "Vejleder 1", "Fag 2", "Vejleder 2"]
    frame = schedule["students"].reindex(columns=columns, fill_value="").fillna("")
    headers = "".join(f"<th>{html.escape(column)}</th>" for column in columns)
    rows = []
    for row in frame.to_dict("records"):
        searchable = f"{compact_search_key(row['Elev'])} {' '.join(class_search_keys(row['Klasse']))}"
        cells = "".join(f"<td>{html.escape(str(row[column]))}</td>" for column in columns)
        rows.append(f'<tr data-search="{html.escape(searchable, quote=True)}">{cells}</tr>')
    return f"""<!doctype html>
<html lang="da">
<head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Min SOP-vejledning</title>
<style>
body{{margin:0;background:#f7fbf9;color:#172126;font:16px/1.45 Arial,sans-serif}}
main{{max-width:1080px;margin:auto;padding:48px 20px 70px}}
h1{{font:700 38px Georgia,serif;margin:0 0 8px}} p{{color:#52656a;margin:0 0 24px}}
.card{{background:#fff;border:1px solid #dce4e2;border-radius:14px;padding:20px;box-shadow:0 8px 24px #1933300b}}
label{{display:block;font-weight:700;margin-bottom:7px}} input{{box-sizing:border-box;width:100%;padding:12px 14px;border:1px solid #9bb5b1;border-radius:8px;font:inherit}}
.hint{{font-size:13px;margin:8px 0 18px}} .table-wrap{{overflow:auto}} table{{border-collapse:collapse;width:100%;min-width:780px}}
th,td{{padding:10px 11px;border-bottom:1px solid #e2e9e7;text-align:left;vertical-align:top}} th{{background:#e7f4f1;color:#075f5b;white-space:nowrap}}
#empty{{display:none;margin:18px 0 0;color:#52656a}} @media print{{main{{padding:0}}.card{{border:0;box-shadow:none}}}}
.copyright{{margin:34px 0 0;text-align:center;color:#66777f;font-size:12px}}
</style></head>
<body><main><h1>Find din SOP-vejledning</h1><p>Detaljeret tidsplan for elever baseret på den beregnede SOP-fordeling. Fordelingen er foretaget ud fra SOPtima – en algoritme udviklet af Henrik (hst@nextkhb.dk). Stadig i beta, så skriv gerne, hvis der er mangler eller fejl.</p>
<section class="card"><label for="query">Elevnavn eller klasse</label><input id="query" type="search" autocomplete="off" placeholder="Fx Ane Andersen, 3q eller 3.q">
<p class="hint">Søgningen foregår kun i denne fil.</p><div class="table-wrap"><table><thead><tr>{headers}</tr></thead><tbody id="results">{''.join(rows)}</tbody></table></div><p id="empty">Ingen vejledninger matcher søgningen.</p></section>
<footer class="copyright">{html.escape(COPYRIGHT)}</footer>
<script>
const query=document.getElementById('query'), rows=[...document.querySelectorAll('#results tr')], empty=document.getElementById('empty');
function normalise(value){{return String(value||'').toLocaleLowerCase('da-DK').normalize('NFD').replace(/[\\u0300-\\u036f]/g,'').replace(/[^a-z0-9]+/g,'');}}
function filter(){{const term=normalise(query.value);let shown=0;rows.forEach(row=>{{const match=!term||row.dataset.search.includes(term);row.hidden=!match;if(match)shown++;}});empty.style.display=shown?'none':'block';}}
query.addEventListener('input',filter); query.focus();
</script></main></body></html>"""


def replace_docx_placeholders(template: bytes, values: dict[str, str]) -> bytes:
    from docx import Document
    source = io.BytesIO(template)
    document = Document(source)
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        full = "".join(run.text for run in paragraph.runs)
        for placeholder, replacement in values.items():
            if placeholder in full:
                full = full.replace(placeholder, replacement)
        if paragraph.runs:
            paragraph.runs[0].text = full
            for run in paragraph.runs[1:]:
                run.text = ""
    result = io.BytesIO()
    document.save(result)
    return result.getvalue()


def make_docx_zip(students: list[dict[str, Any]], teachers: list[dict[str, Any]], solution: dict[str, Any]) -> bytes:
    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    files = io.BytesIO()
    with zipfile.ZipFile(files, "w", zipfile.ZIP_DEFLATED) as archive:
        for index, student in enumerate(students):
            assigned = solution["assignments"][index]
            values = {"Elevnavn": student["name"], "Klasse": student.get("className", ""), "Fag1 og niveau": student["subjects"][0], "Vejleder fag 1": teacher_label(assigned[0], teacher_map), "Fag2 og niveau": student["subjects"][1], "Vejleder fag 2": teacher_label(assigned[1], teacher_map)}
            if DEFAULT_TEMPLATE.exists():
                template = DEFAULT_TEMPLATE.read_bytes()
                content = replace_docx_placeholders(template, values)
            else:
                from docx import Document
                document = Document()
                document.add_heading(f"SOP – {student['name']}", 0)
                document.add_paragraph(f"Klasse: {student.get('className', '')}")
                document.add_paragraph(f"{student['subjects'][0]}: {teacher_label(assigned[0], teacher_map)}")
                document.add_paragraph(f"{student['subjects'][1]}: {teacher_label(assigned[1], teacher_map)}")
                output = io.BytesIO(); document.save(output); content = output.getvalue()
            identity = repair_text(student.get("id")) or f"elev-{index + 1}"
            filename = re.sub(r"[<>:\"/\\|?*]", "-", f"SOP - {student['name']} - {identity}.docx")
            filename = re.sub(r"\s+", " ", filename).strip(" .")
            if filename in archive.namelist():
                filename = re.sub(r"\.docx$", f" - {index + 1}.docx", filename, flags=re.I)
            archive.writestr(filename, content)
    return files.getvalue()


def students_to_frame(students: list[dict[str, Any]], teachers: list[dict[str, Any]] | None = None) -> pd.DataFrame:
    """Lav en redigerbar tabel; vis lærerønsker som læsbare navne, når lærerlisten findes."""
    teacher_map = {teacher["id"]: teacher for teacher in (teachers or [])}
    rows = []
    for student in students:
        subjects = list(student.get("subjects", [])) + ["", ""]
        wishes = [teacher_label(normal_key(wish), teacher_map) if normal_key(wish) in teacher_map else repair_text(wish) for wish in student.get("wishes", [])]
        if not wishes:
            wishes = [NO_WISHES_LABEL]
        wishes = wishes + [""]
        rows.append(
            {
                "Elev-ID": student.get("id", ""),
                "Elevnavn": student.get("name", ""),
                "Klasse": student.get("className", ""),
                "Fag 1": subjects[0],
                "Fag 2": subjects[1],
                "Ønskevejleder 1": wishes[0],
                "Ønskevejleder 2": wishes[1],
                "Projekttitel": student.get("projectTitle", ""),
                "Projektbeskrivelse": student.get("projectDescription", ""),
            }
        )
    return pd.DataFrame(rows, columns=["Elev-ID", "Elevnavn", "Klasse", "Fag 1", "Fag 2", "Ønskevejleder 1", "Ønskevejleder 2", "Projekttitel", "Projektbeskrivelse"])


def wish_column_config(teachers: list[dict[str, Any]]) -> dict[str, Any]:
    """Giv elevarket en tydelig dropdown med alle lærere og ingen ønsker."""
    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    options = [NO_WISHES_LABEL] + sorted((teacher_label(teacher_id, teacher_map) for teacher_id in teacher_map), key=str.casefold)
    return {
        "Ønskevejleder 1": st.column_config.SelectboxColumn("Ønskevejleder 1", options=options, required=False),
        "Ønskevejleder 2": st.column_config.SelectboxColumn("Ønskevejleder 2", options=options, required=False),
    }


def frame_to_students(frame: pd.DataFrame) -> list[dict[str, Any]]:
    if frame.empty:
        return []
    students = parse_students(frame)
    ids = [normal_key(value) for value in frame.get("Elev-ID", pd.Series(dtype=object)).tolist()]
    for index, student in enumerate(students):
        if index < len(ids) and ids[index]:
            student["id"] = ids[index]
    return students


def student_data_signature(students: list[dict[str, Any]]) -> list[tuple[Any, ...]]:
    """Sammenlign kun de felter, der kan redigeres i elevarket.

    Niveauoplysninger og forskel på store/små bogstaver i Elev-ID må ikke få
    editoren til at markere data som ændret ved hver genkørsel.
    """
    return [
        (
            normal_key(student.get("id", "")),
            repair_text(student.get("name", "")),
            repair_text(student.get("className", "")),
            tuple(repair_text(subject) for subject in list(student.get("subjects", []))[:2]),
            tuple(normal_key(wish) for wish in student.get("wishes", []) if normal_key(wish)),
            repair_text(student.get("projectTitle", "")),
            repair_text(student.get("projectDescription", "")),
        )
        for student in students
    ]


def teachers_to_frame(
    teachers: list[dict[str, Any]],
    capacities: dict[str, int],
    subject_count: int | None = None,
    hold_count: int | None = None,
) -> pd.DataFrame:
    # Tre fagkolonner er standard; brugeren kan udvide ved behov.
    existing_subject_count = max((len(teacher.get("subjects", [])) for teacher in teachers), default=0)
    subject_count = max(3, subject_count or 3, existing_subject_count)
    existing_hold_count = max((len(teacher.get("holds", [])) for teacher in teachers), default=0)
    # Op til fire hold kan vises. Eksisterende holdkolonner bevares, hvis brugeren
    # midlertidigt vælger et lavere antal i kontrollen.
    hold_count = min(4, max(0, hold_count if hold_count is not None else 4, existing_hold_count))
    columns = ["Initialer", "Lærernavn"] + [f"Fag {index}" for index in range(1, subject_count + 1)] + [f"Hold {index}" for index in range(1, hold_count + 1)] + ["Max"]
    rows = []
    for teacher in teachers:
        subjects = list(teacher.get("subjects", []))
        row = {"Initialer": teacher.get("id", ""), "Lærernavn": teacher.get("name", "")}
        row.update({f"Fag {index}": subjects[index - 1] if index <= len(subjects) else "" for index in range(1, subject_count + 1)})
        holds = list(teacher.get("holds", []))
        row.update({f"Hold {index}": holds[index - 1] if index <= len(holds) else "" for index in range(1, hold_count + 1)})
        row["Max"] = int(capacities.get(teacher.get("id", ""), 0))
        rows.append(row)
    return pd.DataFrame(rows, columns=columns)


def frame_to_teachers(frame: pd.DataFrame, technical: dict[str, str] | None = None) -> tuple[list[dict[str, Any]], dict[str, int]]:
    if frame.empty:
        return [], {}
    technical = technical or {}
    subject_columns = [column for column in frame.columns if normal_key(column).startswith("fag")]
    hold_columns = [column for column in frame.columns if normal_key(column).startswith("hold")]
    teachers_by_id: dict[str, dict[str, Any]] = {}
    capacities: dict[str, int] = {}
    for _, row in frame.iterrows():
        teacher_id = normal_key(row.get("Initialer", ""))
        if not teacher_id:
            continue
        name = value_at(row, "Lærernavn") or teacher_id
        teacher = teachers_by_id.setdefault(teacher_id, {"id": teacher_id, "name": name, "subjects": [], "holds": []})
        if teacher["name"] == teacher_id and name != teacher_id:
            teacher["name"] = name
        for column in subject_columns:
            subject = subject_from_course(row.get(column), technical)
            if subject and canonical_subject(subject) not in {canonical_subject(item) for item in teacher["subjects"]}:
                teacher["subjects"].append(subject)
        for column in hold_columns:
            hold = value_at(row, column)
            if hold and normal_key(hold) not in {normal_key(item) for item in teacher["holds"]}:
                teacher["holds"].append(hold)
        raw_max = value_at(row, "Max")
        capacities[teacher_id] = parse_capacity(raw_max)
    return sorted(teachers_by_id.values(), key=lambda item: item["id"]), capacities


def resolve_wishes(students: list[dict[str, Any]], teachers: list[dict[str, Any]]) -> None:
    """Tillad både initialer og lærernavne i ønskekolonnerne."""
    lookup = {}
    name_ids: dict[str, set[str]] = defaultdict(set)
    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    for teacher in teachers:
        lookup[normal_key(teacher["id"])] = teacher["id"]
        if normal_key(teacher.get("name", "")):
            name_ids[normal_key(teacher.get("name", ""))].add(teacher["id"])
        # Elevarket viser etiketten "Navn (initialer)"; den skal også kunne læses tilbage.
        lookup[normal_key(teacher_label(teacher["id"], teacher_map))] = teacher["id"]
    for name, ids in name_ids.items():
        if len(ids) == 1:
            lookup[name] = next(iter(ids))
    no_wishes_key = normal_key(NO_WISHES_LABEL)
    for student in students:
        resolved = []
        for wish in student.get("wishes", []):
            key = normal_key(wish)
            if not key or key == no_wishes_key:
                continue
            resolved.append(lookup.get(key, key))
        student["wishes"] = resolved


def input_validation(students: list[dict[str, Any]], teachers: list[dict[str, Any]]) -> dict[str, list[dict[str, str]]]:
    teacher_ids = {teacher["id"] for teacher in teachers}
    candidates = build_candidates(teachers)
    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    unknown_wishes = []
    ambiguous_wishes = []
    missing_subjects = []
    incompatible_wishes = []
    partial_wishes = []
    teacher_names: dict[str, list[str]] = defaultdict(list)
    for teacher in teachers:
        name_key = normal_key(teacher.get("name", ""))
        if name_key and teacher["id"] not in teacher_names[name_key]:
            teacher_names[name_key].append(teacher["id"])
    ambiguous_names = {name: ids for name, ids in teacher_names.items() if len(ids) > 1}
    for student in students:
        for wish in student.get("wishes", []):
            if wish not in teacher_ids:
                if normal_key(wish) in ambiguous_names:
                    ambiguous_wishes.append({
                        "Elev": student["name"],
                        "Ønskevejleder": wish,
                        "Mulige initialer": " · ".join(ambiguous_names[normal_key(wish)]),
                    })
                else:
                    unknown_wishes.append({"Elev": student["name"], "Ønskevejleder": wish})
        for subject in student.get("subjects", []):
            if subject and not candidates.get(canonical_subject(subject), []):
                missing_subjects.append({"Elev": student["name"], "Fag": subject})
        for wish in student.get("wishes", []):
            if wish in teacher_ids:
                covered = [subject for subject in student.get("subjects", []) if wish in candidates.get(canonical_subject(subject), [])]
                uncovered = [subject for subject in student.get("subjects", []) if subject not in covered]
                if not covered:
                    incompatible_wishes.append({"Elev": student["name"], "Ønskevejleder": teacher_label(wish, teacher_map), "Elevens fag": " · ".join(student.get("subjects", [])), "Status": "Underviser ikke i elevens fag"})
                elif uncovered:
                    partial_wishes.append({"Elev": student["name"], "Ønskevejleder": teacher_label(wish, teacher_map), "Fag uden match": " · ".join(uncovered), "Status": "Underviser kun i ét af elevens fag"})
    def unique_rows(rows: list[dict[str, str]]) -> list[dict[str, str]]:
        seen = set()
        result = []
        for row in rows:
            marker = tuple(sorted(row.items()))
            if marker not in seen:
                seen.add(marker)
                result.append(row)
        return result

    # Saml flere ukendte ønsker fra samme elev i én række, så eleven ikke
    # optræder dobbelt i advarslen.
    unknown_by_student: dict[str, list[str]] = defaultdict(list)
    for row in unknown_wishes:
        if row["Ønskevejleder"] not in unknown_by_student[row["Elev"]]:
            unknown_by_student[row["Elev"]].append(row["Ønskevejleder"])
    unknown_wishes = [
        {"Elev": student_name, "Ønskevejleder": " · ".join(wishes)}
        for student_name, wishes in unknown_by_student.items()
    ]

    return {
        "unknown_wishes": unknown_wishes,
        "ambiguous_wishes": unique_rows(ambiguous_wishes),
        "missing_subjects": unique_rows(missing_subjects),
        "incompatible_wishes": unique_rows(incompatible_wishes),
        "partial_wishes": unique_rows(partial_wishes),
    }


def incompatible_wish_suggestions(students: list[dict[str, Any]], teachers: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """Lav redigerbare forslag til ønsker, der ikke matcher nogen af elevens fag."""
    teacher_ids = {teacher["id"] for teacher in teachers}
    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    candidates = build_candidates(teachers)
    rows = []
    for student_index, student in enumerate(students):
        for wish_index, wish in enumerate(student.get("wishes", [])):
            if wish in teacher_ids and any(wish in candidates.get(canonical_subject(subject), []) for subject in student.get("subjects", [])):
                continue
            possible_ids = []
            for subject in student.get("subjects", []):
                for teacher_id in candidates.get(canonical_subject(subject), []):
                    if teacher_id not in possible_ids:
                        possible_ids.append(teacher_id)
            rows.append({
                "_student_index": student_index,
                "_wish_index": wish_index,
                "_possible_ids": possible_ids,
                "Elev": student["name"],
                "Nuværende ønske": teacher_label(wish, teacher_map),
                "Elevens fag": " · ".join(student.get("subjects", [])),
                "Mulige lærere": " · ".join(teacher_label(teacher_id, teacher_map) for teacher_id in possible_ids) or "Ingen",
                "Vælg ny lærer": "",
            })
    return rows


def subject_coverage(students: list[dict[str, Any]], teachers: list[dict[str, Any]]) -> pd.DataFrame:
    """Viser tydeligt hvilke elevfag der har mindst én mulig lærer."""
    demand = Counter(canonical_subject(subject) for student in students for subject in student.get("subjects", []) if subject)
    candidates = build_candidates(teachers)
    teacher_map = {teacher["id"]: teacher for teacher in teachers}
    labels = {teacher_id: teacher_label(teacher_id, teacher_map) for teacher_id in teacher_map}
    subjects = {}
    for student in students:
        for subject in student.get("subjects", []):
            if subject:
                subjects.setdefault(canonical_subject(subject), subject)
    rows = []
    for key, subject in sorted(subjects.items(), key=lambda item: item[1].casefold()):
        ids = candidates.get(key, [])
        rows.append({
            "Fag": subject,
            "Antal elevpladser": demand.get(key, 0),
            "Mulige lærere": " · ".join(labels.get(teacher_id, teacher_id) for teacher_id in ids) or "Ingen",
            "Status": "Dækket" if ids else "MANGLER LÆRER",
        })
    return pd.DataFrame(rows, columns=["Fag", "Antal elevpladser", "Mulige lærere", "Status"])


def data_readiness(students: list[dict[str, Any]], teachers: list[dict[str, Any]]) -> list[str]:
    """Returnér fejl, som skal være rettet, før en fordeling må beregnes."""
    errors = []
    if not students:
        errors.append("Der er ikke indlæst nogen elever.")
    if not teachers:
        errors.append("Der er ikke indlæst nogen lærere.")
    for index, student in enumerate(students, 1):
        if not repair_text(student.get("name")):
            errors.append(f"Elev række {index} mangler elevnavn.")
        subjects = list(student.get("subjects", []))
        if len(subjects) < 2 or any(not repair_text(subject) for subject in subjects[:2]):
            errors.append(f"Elev {student.get('name') or index} skal have både Fag 1 og Fag 2.")
    student_ids = [normal_key(student.get("id", "")) for student in students if normal_key(student.get("id", ""))]
    duplicate_student_ids = sorted(item for item, count in Counter(student_ids).items() if count > 1)
    if duplicate_student_ids:
        errors.append("Elev-ID skal være entydige. Dubletter: " + ", ".join(duplicate_student_ids[:8]) + ".")
    for index, teacher in enumerate(teachers, 1):
        if not normal_key(teacher.get("id")):
            errors.append(f"Lærer række {index} mangler initialer.")
        if not any(repair_text(subject) for subject in teacher.get("subjects", [])):
            errors.append(f"Lærer {teacher.get('name') or index} mangler mindst ét fag.")

    validation = input_validation(students, teachers)
    if validation["unknown_wishes"]:
        errors.append(f"{len(validation['unknown_wishes'])} ønskevejleder(e) findes ikke i lærerlisten.")
    if validation["ambiguous_wishes"]:
        errors.append(f"{len(validation['ambiguous_wishes'])} ønskevejleder(e) er tvetydige og skal vælges med initialer.")
    if validation["missing_subjects"]:
        errors.append(f"{len(validation['missing_subjects'])} elevfag har ingen lærer.")
    if validation["incompatible_wishes"]:
        errors.append(f"{len(validation['incompatible_wishes'])} ønsker peger på lærere uden elevens fag.")
    return errors


def assignment_assessment(
    students: list[dict[str, Any]],
    teachers: list[dict[str, Any]],
    capacities: dict[str, int],
    assignments: list[list[str | None]],
    K: int,
    double_limit: int,
    use_global_k: bool,
    allow_over_capacity: bool,
    lock_teacher_max: bool,
) -> dict[str, Any]:
    """Kontrollér en automatisk eller manuel fordeling mod alle hårde regler."""
    teacher_ids = {teacher["id"] for teacher in teachers}
    candidates = build_candidates(teachers)
    loads = {teacher_id: 0 for teacher_id in teacher_ids}
    double_loads = {teacher_id: 0 for teacher_id in teacher_ids}
    errors: list[str] = []
    warnings: list[str] = []
    incomplete = 0

    if len(assignments) != len(students):
        errors.append("Antallet af tildelinger svarer ikke til antallet af elever.")
    for index, student in enumerate(students):
        assigned = list(assignments[index]) if index < len(assignments) else [None, None]
        assigned = (assigned + [None, None])[:2]
        if any(not teacher_id for teacher_id in assigned):
            incomplete += 1
        for slot, teacher_id in enumerate(assigned):
            if not teacher_id:
                continue
            teacher_id = normal_key(teacher_id)
            subject = student.get("subjects", ["", ""])[slot]
            if teacher_id not in teacher_ids:
                errors.append(f"{student['name']}: læreren {teacher_id} findes ikke i lærerlisten.")
            elif teacher_id not in candidates.get(canonical_subject(subject), []):
                errors.append(f"{student['name']}: {teacher_id} underviser ikke i {subject}.")
        for teacher_id in set(normal_key(item) for item in assigned if item and normal_key(item) in loads):
            loads[teacher_id] += 1
        if assigned[0] and normal_key(assigned[0]) == normal_key(assigned[1]):
            teacher_id = normal_key(assigned[0])
            if teacher_id in double_loads:
                double_loads[teacher_id] += 1

    permit_over = use_global_k and allow_over_capacity and not lock_teacher_max
    for teacher_id, load in sorted(loads.items()):
        individual = int(capacities.get(teacher_id, 0))
        normal_limit = min(K, individual) if use_global_k else individual
        hard_limit = K if permit_over else normal_limit
        if load > hard_limit:
            errors.append(f"{teacher_id} har {load} elever, men den hårde grænse er {hard_limit}.")
        elif load > normal_limit:
            warnings.append(f"{teacher_id} har {load} elever og ligger {load - normal_limit} over sit individuelle max.")
        if double_loads[teacher_id] > double_limit:
            errors.append(
                f"{teacher_id} har {double_loads[teacher_id]} dobbeltvejledninger, men I-grænsen er {double_limit}."
            )
    if incomplete:
        warnings.append(f"{incomplete} elev(er) mangler mindst én vejleder.")
    return {
        "errors": errors,
        "warnings": warnings,
        "loads": loads,
        "double_loads": double_loads,
        "incomplete": incomplete,
    }


def stable_signature(payload: Any) -> str:
    encoded = json.dumps(payload, sort_keys=True, ensure_ascii=False, default=str, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def distribution_signature(
    students: list[dict[str, Any]],
    teachers: list[dict[str, Any]],
    capacities: dict[str, int],
    K: int,
    double_limit: int,
    use_global_k: bool,
    allow_over_capacity: bool,
    lock_teacher_max: bool,
    prioritize_pairs: bool,
    prioritize_classes: bool,
    attempts: int,
) -> str:
    return stable_signature({
        "students": students,
        "teachers": teachers,
        "capacities": {key: int(value) for key, value in capacities.items()},
        "rules": [K, double_limit, use_global_k, allow_over_capacity, lock_teacher_max, prioritize_pairs, prioritize_classes, attempts],
    })


def schedule_signature(
    students: list[dict[str, Any]],
    teachers: list[dict[str, Any]],
    solution: dict[str, Any],
    settings: list[Any],
) -> str:
    return stable_signature({
        "students": students,
        "teachers": teachers,
        "assignments": solution.get("assignments", []),
        "settings": settings,
    })


def invalidate_derived_state(reason: str = "", *, require_approval: bool = False) -> None:
    """Fjern resultater, der ikke længere svarer til de aktuelle data eller regler."""
    had_result = st.session_state.get("solution") is not None or st.session_state.get("v2_schedule") is not None
    st.session_state["solution"] = None
    st.session_state["v2_schedule"] = None
    st.session_state["v2_schedule_signature"] = None
    st.session_state["v2_docx_zip"] = None
    if require_approval:
        st.session_state["input_loaded"] = False
    if reason and had_result:
        st.session_state["v2_stale_notice"] = reason


def approve_data_and_navigate(next_step: str | None = None) -> None:
    st.session_state["input_loaded"] = True
    invalidate_derived_state()
    st.session_state["v2_stale_notice"] = ""
    if next_step:
        st.session_state["v2_active_step"] = next_step


def navigate_to_step(step: str) -> None:
    st.session_state["v2_active_step"] = step


def make_input_export(students: list[dict[str, Any]], teachers: list[dict[str, Any]], capacities: dict[str, int]) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        safe_excel_frame(students_to_frame(students, teachers)).to_excel(writer, index=False, sheet_name="Elevdata")
        safe_excel_frame(teachers_to_frame(teachers, capacities)).to_excel(writer, index=False, sheet_name="Lærerdata")
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
    return output.getvalue()


def make_teacher_export(
    teachers: list[dict[str, Any]],
    capacities: dict[str, int],
    subject_count: int | None = None,
    hold_count: int | None = None,
) -> bytes:
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        safe_excel_frame(teachers_to_frame(teachers, capacities, subject_count, hold_count)).to_excel(writer, index=False, sheet_name="Lærerdata")
        sheet = writer.book["Lærerdata"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            width = min(55, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    return output.getvalue()


def try_excel_export(factory: Any) -> bytes | None:
    try:
        return factory()
    except ImportError:
        return None


def load_default_state(load_students: bool = True, load_teachers: bool = True) -> None:
    students, teachers, capacities = embedded_data()
    st.session_state.setdefault("students", students if load_students else [])
    st.session_state.setdefault("teachers", teachers if load_teachers else [])
    st.session_state.setdefault("capacities", capacities if load_teachers else {})
    st.session_state.setdefault("solution", None)


def activate_demo_data() -> None:
    """Indlæs et komplet, valideret demosæt i den aktuelle app-session."""
    demo_students = parse_students(read_path_table(TEST_STUDENTS_FILE))
    demo_teachers, demo_capacities = parse_teachers(read_path_table(TEST_TEACHERS_FILE))
    resolve_wishes(demo_students, demo_teachers)
    readiness_errors = data_readiness(demo_students, demo_teachers)
    st.session_state["students"] = demo_students
    st.session_state["teachers"] = demo_teachers
    st.session_state["capacities"] = {
        teacher["id"]: int(demo_capacities.get(teacher["id"], 0))
        for teacher in demo_teachers
    }
    st.session_state["student_revision"] = st.session_state.get("student_revision", 0) + 1
    st.session_state["teacher_revision"] = st.session_state.get("teacher_revision", 0) + 1
    st.session_state["capacity_revision"] = st.session_state.get("capacity_revision", 0) + 1
    st.session_state["input_loaded"] = not readiness_errors
    st.session_state["solution"] = None
    st.session_state["v2_schedule"] = None
    st.session_state["v2_schedule_signature"] = None
    st.session_state["v2_docx_zip"] = None
    st.session_state["v2_stale_notice"] = ""
    st.session_state["v2_demo_mode"] = True


def legacy_main() -> None:
    st.set_page_config(page_title="SOPtima · demo med fiktive data", page_icon="🎓", layout="wide", initial_sidebar_state="expanded")
    st.markdown("""
    <style>
      :root { --ink:#172126; --muted:#66777f; --brand:#075f5b; --soft:#e7f4f1; }
      .stApp { background: radial-gradient(circle at 8% 0%, rgba(15,129,122,.12), transparent 31rem), linear-gradient(180deg,#f7fbf9 0,#fffdf8 24rem); }
      h1,h2,h3 { font-family: Georgia, 'Times New Roman', serif; letter-spacing:-.025em; }
      .hero-kicker { color:var(--brand); font-size:.76rem; font-weight:800; letter-spacing:.14em; text-transform:uppercase; margin-bottom:.35rem; }
      .info-card { border:1px solid #dce4e2; border-radius:14px; padding:1rem 1.15rem; background:rgba(255,255,255,.78); color:#52625f; }
      .metric-card { border:1px solid #dce4e2; border-radius:14px; padding:1rem 1.1rem; background:rgba(255,255,255,.78); min-height:92px; }
      .metric-value { display:block; font:700 2rem/1 Georgia,serif; color:#172126; }
      .metric-label { color:#66777f; font-size:.72rem; font-weight:800; letter-spacing:.07em; text-transform:uppercase; }
      div[data-testid='stDataFrame'] { border:1px solid #dce4e2; border-radius:12px; overflow:hidden; }
      div[data-testid='stDataEditor'] { border:1px solid #dce4e2; border-radius:12px; overflow:hidden; }
    </style>
    """, unsafe_allow_html=True)
    load_default_state()
    st.session_state.setdefault("capacity_revision", 0)
    students = st.session_state["students"]
    teachers = st.session_state["teachers"]
    capacities = st.session_state["capacities"]
    teacher_map = {teacher["id"]: teacher for teacher in teachers}

    st.markdown('<div class="hero-kicker">SOPtima · planlægningsværktøj</div>', unsafe_allow_html=True)
    st.caption(f"{len(students)} elever · {len(teachers)} lærere · data lokalt på denne computer")
    st.markdown('<div class="info-card"><strong>Sådan fungerer fordelingen:</strong> Først reserveres gyldige elevønsker til elever med få muligheder. Derefter fordeles de mest begrænsede fag først under hensyn til K, individuelle max-tal og grænsen for dobbeltvejledning I. Der beregnes flere variationer, og den bedste løsning vises.</div>', unsafe_allow_html=True)

    st.markdown("### 1 · Indlæs data")
    with st.sidebar:
        st.header("Proces")
        st.markdown("**1 - Data**\n\nIndlæs elever og lærere")
        st.markdown("**2 - Regler**\n\nVælg K, I og prioriteringer")
        st.markdown("**3 - Lærermax**\n\nKontrollér kapaciteter")
        st.markdown("**4 - Resultat**\n\nGennemgå og eksportér")
        st.divider()
        st.caption("Alt behandles lokalt på denne computer.")

    with st.expander("1 - Indlæs egne data", expanded=False):
        st.write("Elevarket skal indeholde Navn, Klasse, Fag 1, Fag 2 og gerne to kolonner med vejlederønsker.")
        student_upload = st.file_uploader("Elevdata (.xlsx, .xlsm eller .csv)", type=["xlsx", "xlsm", "csv"], key="students_upload")
        teacher_upload = st.file_uploader("Lærerdata (.xlsx, .xlsm eller .csv) – valgfri", type=["xlsx", "xlsm", "csv"], key="teachers_upload")
        if st.button("Indlæs og klargør data", type="primary"):
            try:
                if student_upload is not None:
                    st.session_state["students"] = parse_students(read_uploaded_table(student_upload))
                if teacher_upload is not None:
                    parsed_teachers, parsed_caps = parse_teacher_upload(teacher_upload)
                    if parsed_teachers:
                        st.session_state["teachers"] = parsed_teachers
                        st.session_state["capacities"] = {teacher["id"]: parsed_caps.get(teacher["id"], 0) for teacher in parsed_teachers}
                        st.session_state["capacity_revision"] += 1
                st.session_state["solution"] = None
                st.success(f"Data klargjort: {len(st.session_state['students'])} elever og {len(st.session_state['teachers'])} lærere.")
                st.rerun()
            except Exception as error:
                st.error(str(error))

    st.markdown("### 2 · Vælg fordelingsregler")
    with st.expander("Regler for K, I og prioriteringer", expanded=True):
        rule_cols = st.columns(2)
        with rule_cols[0]:
            use_global = st.checkbox("Brug global K-grænse", value=True, key="main_use_global")
            K = st.slider("Maksimalt antal elever pr. lærer (K)", 1, 50, 18, key="main_k")
            double_limit = st.slider("Maksimale dobbeltvejledninger pr. lærer (I)", 0, 50, 0, key="main_double_limit")
            allow_over = st.checkbox("Tillad overskridelse af max (op til K)", value=True, disabled=not use_global, key="main_allow_over")
        with rule_cols[1]:
            lock_max = st.checkbox("Lås lærernes max-tal fast", value=False, key="main_lock_max")
            prioritize_pairs = st.checkbox("Prioritér samme vejlederpar", value=True, key="main_prioritize_pairs")
            prioritize_classes = st.checkbox("Saml elever fra samme klasse/hold", value=True, key="main_prioritize_classes")
            attempts = st.slider("Algoritmedybde for fordeling", 20, 180, 120, 10, key="main_attempts", help="Algoritmen forsøger at placere eleverne hos lærere, der dækker fagene, samtidig med at kapacitet og elevønsker respekteres. Tallet angiver, hvor mange forslag der afprøves; højere værdi kan give et bedre resultat, men tager længere tid.")
        st.caption("K er den globale grænse. I bestemmer, hvor mange elever der må få samme lærer til begge fag.")

    st.markdown("### 3 · Kontrollér lærermax og beregn")
    st.caption("Redigér kun kolonnen Max. Brug søgning for at finde en bestemt lærer eller et fag.")
    capacity_query = st.text_input("Søg i lærere", placeholder="Navn, initialer eller fag", key="main_capacity_query").casefold()
    sorted_teachers = sorted(teachers, key=lambda item: teacher_label(item["id"], teacher_map))
    capacity_rows = [{"Lærer": teacher_label(teacher["id"], teacher_map), "Initialer": teacher["id"], "Fag": " · ".join(teacher["subjects"]), "Max": int(capacities.get(teacher["id"], 0))} for teacher in sorted_teachers]
    if capacity_query:
        capacity_rows = [row for row in capacity_rows if capacity_query in " ".join(str(value) for value in row.values()).casefold()]
    edited_capacities = st.data_editor(
        pd.DataFrame(capacity_rows),
        width="stretch",
        hide_index=True,
        height=min(430, 38 + max(1, len(capacity_rows)) * 35),
        disabled=["Lærer", "Initialer", "Fag"],
        column_config={"Max": st.column_config.NumberColumn("Max", min_value=0, max_value=50, step=1, required=True)},
        key=f"main_capacity_editor_{st.session_state['capacity_revision']}",
    )
    for _, row in edited_capacities.iterrows():
        teacher_id = normal_key(row["Initialer"])
        if teacher_id in capacities:
            capacities[teacher_id] = max(0, min(50, int(row["Max"] or 0)))

    cap_action_cols = st.columns([1, 1, 1, 3])
    with cap_action_cols[0]:
        if st.button("Nulstil max"):
            _, _, default_caps = embedded_data()
            st.session_state["capacities"] = {teacher["id"]: default_caps.get(teacher["id"], 0) for teacher in teachers}
            st.session_state["capacity_revision"] += 1
            st.rerun()
    with cap_action_cols[1]:
        if st.button("Foreslå kapaciteter"):
            st.info("Forslag beregnes ved at optimere med ensartet loft. Tryk derefter på Beregn fordeling for at kontrollere resultatet.")
            progress_bar = st.progress(0.0, text="Beregner kapacitetsforslag: 0 %")
            suggestion = optimize(
                students, teachers, {teacher["id"]: 50 for teacher in teachers}, K, double_limit, use_global, allow_over,
                lock_max, prioritize_pairs, prioritize_classes, max(30, attempts // 2),
                progress_callback=lambda value: update_algorithm_progress(progress_bar, "Beregner kapacitetsforslag", value),
            )
            update_algorithm_progress(progress_bar, "Kapacitetsforslag færdigt", 1.0)
            st.session_state["capacities"] = dict(suggestion["loads"])
            st.session_state["capacity_revision"] += 1
            st.rerun()
    with cap_action_cols[2]:
        if st.button("Beregn fordeling", type="primary"):
            progress_bar = st.progress(0.0, text="Beregner fordeling: 0 %")
            with st.spinner("Beregner flere mulige fordelinger …"):
                st.session_state["solution"] = optimize(
                    students, teachers, capacities, K, double_limit, use_global, allow_over, lock_max,
                    prioritize_pairs, prioritize_classes, attempts,
                    progress_callback=lambda value: update_algorithm_progress(progress_bar, "Beregner fordeling", value),
                )
            update_algorithm_progress(progress_bar, "Fordeling færdig", 1.0)
            st.rerun()

    solution = st.session_state.get("solution")
    if solution is None:
        st.info("Vælg kapaciteter og tryk **Beregn fordeling** for at se resultatet.")
        return
    st.markdown("### 4 - Resultat og eksport")
    stats = solution["stats"]
    if stats["unassigned"] or stats["over_teachers"] or stats["capacity_blocked_slots"]:
        messages = []
        if stats["unassigned"]:
            messages.append(f"{stats['unassigned']} elever mangler mindst én vejleder.")
        if stats["capacity_blocked_slots"]:
            messages.append(f"{stats['capacity_blocked_slots']} fagpladser blev blokeret af lærermax.")
        if stats["over_teachers"]:
            messages.append(f"{stats['over_teachers']} lærere ligger over deres individuelle max.")
        st.warning(" ".join(messages))

    metric_data = [(len(students) - stats["unassigned"], "Elever fordelt"), (stats["both"], "Begge ønsker"), (stats["one"] + stats["both"], "Mindst ét ønske"), (stats["none"], "Ingen ønsker"), (stats["capacity_blocked_students"], "Kapacitetsblokerede")]
    cols = st.columns(5)
    for column, (value, label) in zip(cols, metric_data):
        column.markdown(f'<div class="metric-card"><span class="metric-value">{value}</span><span class="metric-label">{label}</span></div>', unsafe_allow_html=True)
    st.caption(f"Resultat ved {'K=' + str(solution['K']) if solution['use_global_k'] else 'lærernes egne max-tal'} og I={solution['double_limit']}. {stats['both']} elever får begge ønsker opfyldt.")

    tab_students, tab_teachers, tab_subjects, tab_unassigned, tab_analysis = st.tabs(["Elever", "Lærere", "Fagstatistik", "Ikke tildelte", "Analyse"])
    with tab_students:
        query = st.text_input("Søg i elever", key="student_query").casefold()
        student_rows = []
        for index, student in enumerate(students):
            assigned = solution["assignments"][index]
            search = " ".join([student["name"], student.get("className", ""), *student["subjects"], *(assigned_item or "" for assigned_item in assigned)]).casefold()
            if query and query not in search:
                continue
            student_rows.append({"Elev": student["name"], "Klasse": student.get("className", ""), "Fag 1": student["subjects"][0], "Vejleder 1": teacher_label(assigned[0], teacher_map), "Fag 2": student["subjects"][1], "Vejleder 2": teacher_label(assigned[1], teacher_map), "Ønsker": f"{wished_count(student, assigned)}/2"})
        st.dataframe(pd.DataFrame(student_rows), width="stretch", hide_index=True, height=560)
        st.caption("Manuel justering foretages via knapperne nedenfor. Værdierne skal være lærerinitialer.")
        with st.expander("Manuel justering af vejledere"):
            edit_rows = []
            for index, student in enumerate(students):
                assigned = solution["assignments"][index]
                edit_rows.append({"Elev": student["name"], "Fag 1": student["subjects"][0], "Vejleder 1": assigned[0] or "", "Fag 2": student["subjects"][1], "Vejleder 2": assigned[1] or ""})
            edited = st.data_editor(pd.DataFrame(edit_rows), width="stretch", hide_index=True, disabled=["Elev", "Fag 1", "Fag 2"], key="manual_editor")
            if st.button("Gem manuelle ændringer"):
                valid_ids = set(teacher_map)
                for index, row in edited.iterrows():
                    for slot in (1, 2):
                        value = normal_key(row[f"Vejleder {slot}"])
                        if value and value not in valid_ids:
                            st.error(f"Ukendt lærerinitial på række {index + 1}: {value}")
                            break
                        solution["assignments"][index][slot - 1] = value or None
                solution["loads"] = {teacher_id: 0 for teacher_id in teacher_map}
                solution["double_loads"] = {teacher_id: 0 for teacher_id in teacher_map}
                for assigned in solution["assignments"]:
                    for teacher_id in set(item for item in assigned if item):
                        solution["loads"][teacher_id] += 1
                    if assigned[0] and assigned[0] == assigned[1]:
                        solution["double_loads"][assigned[0]] += 1
                solution["stats"] = score_solution(students, teachers, capacities, solution["assignments"], solution["loads"], solution["K"], solution["use_global_k"], prioritize_pairs, prioritize_classes)
                st.session_state["solution"] = solution
                st.success("Fordelingen er opdateret.")
                st.rerun()
    with tab_teachers:
        teacher_rows = []
        for teacher in sorted(teachers, key=lambda item: (-solution["loads"].get(item["id"], 0), teacher_label(item["id"], teacher_map))):
            teacher_id = teacher["id"]
            limit = min(solution["K"], capacities.get(teacher_id, 0)) if solution["use_global_k"] else capacities.get(teacher_id, 0)
            pupils = [students[index]["name"] for index, assigned in enumerate(solution["assignments"]) if teacher_id in assigned]
            teacher_rows.append({"Lærer": teacher_label(teacher_id, teacher_map), "Fag": " · ".join(teacher["subjects"]), "Elever": len(pupils), "Max": limit, "Status": "Over max" if len(pupils) > limit else "OK", "Dobbelt": solution["double_loads"].get(teacher_id, 0)})
        st.dataframe(pd.DataFrame(teacher_rows), width="stretch", hide_index=True, height=600)
    with tab_subjects:
        subject_frame = subject_stats(students, teachers, capacities, solution["K"], solution["use_global_k"])
        st.dataframe(subject_frame, width="stretch", hide_index=True, height=560)
    with tab_unassigned:
        missing = []
        for index, student in enumerate(students):
            for slot, subject in enumerate(student["subjects"]):
                if not solution["assignments"][index][slot]:
                    missing.append({"Elev": student["name"], "Klasse": student.get("className", ""), "Fag": subject, "Status": "Mangler vejleder"})
        st.dataframe(pd.DataFrame(missing), width="stretch", hide_index=True, height=560)
        if not missing:
            st.success("Alle fag er tildelt.")
    with tab_analysis:
        st.write("K/I-afvejningen viser, hvordan flere kapacitetsvariationer påvirker antallet af opfyldte ønsker.")
        if st.button("Beregn K-kurve", key="calculate_curve"):
            curve = []
            progress = st.progress(0)
            for index, curve_k in enumerate(range(1, 51)):
                curve_solution = optimize(students, teachers, capacities, curve_k, double_limit, True, allow_over, lock_max, prioritize_pairs, prioritize_classes, max(3, attempts // 30))
                curve.append({"K": curve_k, "Begge ønsker": curve_solution["stats"]["both"], "Mindst ét ønske": curve_solution["stats"]["one"] + curve_solution["stats"]["both"], "Ingen ønsker": curve_solution["stats"]["none"]})
                progress.progress((index + 1) / 50)
            st.session_state["curve"] = pd.DataFrame(curve).set_index("K")
        if st.session_state.get("curve") is not None:
            st.line_chart(st.session_state["curve"], color=["#075f5b", "#e69f35", "#b44438"])
        else:
            st.info("Tryk på Beregn K-kurve for at lave analysen.")

    st.divider()
    st.subheader("Eksport")
    export_col, docx_col = st.columns(2)
    with export_col:
        st.download_button("Download fordeling som Excel", data=make_export(students, teachers, solution, capacities, solution["use_global_k"]), file_name="vejlederfordeling.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary")
    with docx_col:
        if st.button("Forbered Word-filer"):
            with st.spinner("Genererer Word-filer …"):
                try:
                    st.session_state["docx_zip"] = make_docx_zip(students, teachers, solution)
                except ImportError:
                    st.error("Word-eksport kræver python-docx. Kør start_app.cmd igen, så installeres pakken.")
        if st.session_state.get("docx_zip"):
            st.download_button("Download SOP-filer (.zip)", data=st.session_state["docx_zip"], file_name="SOP-filer.zip", mime="application/zip")


if False and __name__ == "__main__":
    legacy_main()


def main() -> None:
    st.set_page_config(
        page_title="SOPtima · intelligent vejlederfordeling",
        page_icon="🎓",
        layout="wide",
        initial_sidebar_state="expanded",
    )
    st.markdown("""
    <style>
      :root {
        --ink:#152326; --muted:#637477; --brand:#08756f; --brand-dark:#055852;
        --brand-soft:#e8f5f2; --line:#dbe6e3; --paper:#ffffff; --warm:#fffaf1;
      }
      .stApp {
        background:
          radial-gradient(circle at 8% -4%, rgba(17,145,134,.14), transparent 28rem),
          radial-gradient(circle at 94% 8%, rgba(233,174,73,.10), transparent 24rem),
          linear-gradient(180deg,#f7fbfa 0,#fffdf8 30rem,#f9fbfa 100%);
        color:var(--ink);
      }
      [data-testid='stHeader'] { background:rgba(247,251,250,.72); backdrop-filter:blur(12px); }
      [data-testid='stAppViewBlockContainer'] { max-width:1480px; padding-top:1.5rem; padding-bottom:4rem; }
      section[data-testid='stSidebar'] { background:linear-gradient(180deg,#f0f8f6 0,#fbfcfa 58%,#fffaf1 100%); border-right:1px solid var(--line); }
      section[data-testid='stSidebar'] [data-testid='stSidebarContent'] { padding-top:1.1rem; }
      h1,h2,h3 { font-family:Georgia,'Times New Roman',serif; letter-spacing:-.025em; color:var(--ink); }
      h2 { margin-top:1.35rem; }
      .soptima-brand { display:flex; align-items:center; gap:.7rem; margin:.2rem 0 1rem; }
      .soptima-brand-icon { display:grid; place-items:center; width:2.25rem; height:2.25rem; border-radius:.75rem; background:linear-gradient(135deg,var(--brand),#12a195); color:white; box-shadow:0 8px 20px rgba(8,117,111,.22); font-size:1.05rem; }
      .soptima-brand-name { font:700 1.34rem/1 Georgia,serif; color:var(--ink); }
      .soptima-brand-sub { color:var(--muted); font-size:.72rem; margin-top:.12rem; }
      .soptima-hero { position:relative; overflow:hidden; display:flex; gap:1rem; align-items:flex-start; padding:1.35rem 1.45rem; margin:0 0 1.35rem; border:1px solid rgba(8,117,111,.16); border-radius:20px; background:linear-gradient(120deg,rgba(255,255,255,.96),rgba(232,245,242,.86)); box-shadow:0 16px 38px rgba(22,54,51,.07); }
      .soptima-hero::after { content:''; position:absolute; width:11rem; height:11rem; right:-4.5rem; top:-5.5rem; border-radius:50%; background:rgba(18,161,149,.08); }
      .soptima-hero-icon { flex:0 0 auto; display:grid; place-items:center; width:2.8rem; height:2.8rem; border-radius:.9rem; background:var(--brand); color:white; font-size:1.25rem; box-shadow:0 8px 22px rgba(8,117,111,.22); }
      .soptima-kicker { color:var(--brand-dark); font-size:.7rem; font-weight:800; letter-spacing:.14em; text-transform:uppercase; margin:.05rem 0 .3rem; }
      .soptima-hero h1 { font-size:clamp(1.65rem,3vw,2.45rem); line-height:1.06; margin:0 0 .45rem; padding:0; }
      .soptima-hero p { max-width:52rem; color:var(--muted); font-size:.98rem; line-height:1.55; margin:0; }
      .source-pill { display:inline-flex; align-items:center; gap:.35rem; margin-top:.65rem; padding:.28rem .58rem; border-radius:999px; background:rgba(8,117,111,.09); color:var(--brand-dark); font-size:.72rem; font-weight:700; }
      .sidebar-card { margin:.55rem 0 .8rem; padding:.78rem .85rem; border:1px solid var(--line); border-radius:13px; background:rgba(255,255,255,.72); }
      .sidebar-card-title { color:var(--ink); font-weight:750; font-size:.82rem; margin-bottom:.45rem; }
      .sidebar-row { display:flex; justify-content:space-between; gap:.7rem; color:var(--muted); font-size:.76rem; padding:.18rem 0; }
      .sidebar-row strong { color:var(--ink); font-weight:700; text-align:right; }
      .info-card,.metric-card { border:1px solid var(--line); border-radius:15px; padding:1rem 1.15rem; background:rgba(255,255,255,.86); box-shadow:0 8px 24px rgba(22,54,51,.04); }
      .metric-card { min-height:96px; transition:transform .16s ease,box-shadow .16s ease; }
      .metric-card:hover { transform:translateY(-2px); box-shadow:0 12px 28px rgba(22,54,51,.08); }
      .metric-value { display:block; font:700 2rem/1 Georgia,serif; color:var(--ink); }
      .metric-label { display:block; margin-top:.45rem; color:var(--muted); font-size:.68rem; font-weight:800; letter-spacing:.075em; text-transform:uppercase; }
      div[data-testid='stDataFrame'],div[data-testid='stDataEditor'] { border:1px solid var(--line); border-radius:14px; overflow:hidden; box-shadow:0 7px 22px rgba(22,54,51,.035); }
      [data-testid='stFileUploaderDropzone'] { border:1px dashed #9bbdb8; border-radius:14px; background:rgba(232,245,242,.48); }
      div.stButton > button,div.stDownloadButton > button { border-radius:11px; min-height:2.55rem; font-weight:700; transition:transform .15s ease,box-shadow .15s ease,border-color .15s ease; }
      div.stButton > button:hover,div.stDownloadButton > button:hover { transform:translateY(-1px); border-color:var(--brand); box-shadow:0 7px 18px rgba(8,117,111,.12); }
      div.stButton > button[kind='primary'] { border:0; background:linear-gradient(135deg,var(--brand-dark),#0a8d84); box-shadow:0 8px 20px rgba(8,117,111,.19); }
      .stTabs [data-baseweb='tab-list'] { gap:.35rem; border-bottom:1px solid var(--line); }
      .stTabs [data-baseweb='tab'] { border-radius:9px 9px 0 0; padding:.45rem .78rem; }
      [data-testid='stAlert'] { border-radius:13px; }
      #MainMenu,footer { visibility:hidden; }
      @media (max-width:700px) { .soptima-hero { padding:1rem; border-radius:16px; } .soptima-hero-icon { width:2.4rem; height:2.4rem; } }
    </style>
    """, unsafe_allow_html=True)
    load_default_state()
    for key, default in (("capacity_revision", 0), ("student_revision", 0), ("teacher_revision", 0), ("input_loaded", False)):
        st.session_state.setdefault(key, default)

    students = st.session_state["students"]
    teachers = st.session_state["teachers"]
    capacities = st.session_state["capacities"]
    resolve_wishes(students, teachers)
    teacher_map = {teacher["id"]: teacher for teacher in teachers}

    st.markdown('<div class="hero-kicker">SOPtima · planlægningsværktøj</div>', unsafe_allow_html=True)
    st.caption(f"{len(students)} elever · {len(teachers)} lærere · data behandles lokalt")
    st.markdown('<div class="info-card"><strong>Arbejdsgang:</strong> Indlæs og ret elev- og lærerdata i første fane. Vælg derefter regler og max-tal. Til sidst beregnes fordelingen og vises i en overskuelig tabel, som også kan downloades som Excel.</div>', unsafe_allow_html=True)

    with st.sidebar:
        st.header("Proces")
        st.markdown("**1 · Data**\n\nIndlæs og ret elever/lærere")
        st.markdown("**2 · Regler**\n\nVælg K, I og max-tal")
        st.markdown("**3 · Fordeling**\n\nBeregn og gennemgå resultatet")
        st.divider()
        st.caption("Alt behandles lokalt på denne computer.")

    data_tab, rules_tab, result_tab = st.tabs(["1 · Elev- og lærerdata", "2 · Fordelingsregler og max", "3 · Beregn fordeling"])

    with data_tab:
        st.subheader("Indlæs data")
        st.write("Indlæs elevdata og lærerdata hver for sig. Du kan derefter rette alle celler direkte i tabellerne nedenfor.")
        upload_cols = st.columns(2)
        with upload_cols[0]:
            student_upload = st.file_uploader("Elevdata (.xlsx, .xlsm eller .csv)", type=["xlsx", "xlsm", "csv"], key="students_upload")
        with upload_cols[1]:
            teacher_upload = st.file_uploader("Lærerdata (.xlsx, .xlsm eller .csv)", type=["xlsx", "xlsm", "csv"], key="teachers_upload")
        if st.button("Indlæs valgte filer", type="primary", key="load_input_files"):
            try:
                if student_upload is not None:
                    st.session_state["students"] = parse_students(read_uploaded_table(student_upload))
                    st.session_state["student_revision"] += 1
                if teacher_upload is not None:
                    parsed_teachers, parsed_caps = parse_teacher_upload(teacher_upload)
                    st.session_state["teachers"] = parsed_teachers
                    st.session_state["capacities"] = {teacher["id"]: parsed_caps.get(teacher["id"], 0) for teacher in parsed_teachers}
                    st.session_state["teacher_revision"] += 1
                    st.session_state["capacity_revision"] += 1
                resolve_wishes(st.session_state["students"], st.session_state["teachers"])
                readiness_errors = data_readiness(st.session_state["students"], st.session_state["teachers"])
                st.session_state["input_loaded"] = not readiness_errors
                st.session_state["solution"] = None
                st.success(f"Data indlæst: {len(st.session_state['students'])} elever og {len(st.session_state['teachers'])} lærere.")
                st.rerun()
            except Exception as error:
                st.error(str(error))

        input_export = try_excel_export(lambda: make_input_export(students, teachers, capacities))
        if input_export is not None:
            st.download_button("Download aktuelle inputark som Excel", data=input_export, file_name="vejlederfordeling_input.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="download_input")
        else:
            st.info("Excel-download kræver, at projektets afhængigheder er installeret (openpyxl mangler i dette miljø).")
        st.divider()
        st.subheader("Elevdata – redigerbar")
        st.caption("Påkrævede felter er Elevnavn, Klasse, Fag 1 og Fag 2. Ønskevejlederne skal være initialer eller lærernavne.")
        edited_students = st.data_editor(
            students_to_frame(students, teachers), width="stretch", hide_index=True, num_rows="dynamic",
            height=min(600, 80 + max(1, len(students)) * 38), disabled=["Elev-ID"],
            column_config={**wish_column_config(teachers), "Projektbeskrivelse": st.column_config.TextColumn("Projektbeskrivelse", width="large")},
            key=f"students_editor_{st.session_state['student_revision']}",
        )
        try:
            new_students = frame_to_students(edited_students)
            resolve_wishes(new_students, teachers)
            if student_data_signature(new_students) != student_data_signature(students):
                st.session_state["students"] = new_students
                st.session_state["input_loaded"] = False
                st.session_state["solution"] = None
                students = new_students
        except ValueError as error:
            st.error(str(error))

        st.subheader("Lærerdata – redigerbar")
        st.caption("Én række pr. lærer. Ret initialer, navn, fag og Max direkte. Brug pilen nederst i tabellen for at tilføje lærere.")
        edited_teachers = st.data_editor(
            teachers_to_frame(teachers, capacities), width="stretch", hide_index=True, num_rows="dynamic",
            height=min(500, 80 + max(1, len(teachers)) * 38),
            column_config={"Max": st.column_config.NumberColumn("Max", min_value=0, max_value=50, step=1, required=True)},
            key=f"teachers_editor_{st.session_state['teacher_revision']}",
        )
        new_teachers, new_capacities = frame_to_teachers(edited_teachers)
        if json.dumps(new_teachers, sort_keys=True, ensure_ascii=False) != json.dumps(teachers, sort_keys=True, ensure_ascii=False) or new_capacities != capacities:
            st.session_state["teachers"] = new_teachers
            st.session_state["capacities"] = new_capacities
            st.session_state["input_loaded"] = False
            st.session_state["solution"] = None
            teachers, capacities = new_teachers, new_capacities
        resolve_wishes(students, teachers)
        teacher_map = {teacher["id"]: teacher for teacher in teachers}

        validation = input_validation(students, teachers)
        if validation["unknown_wishes"]:
            st.warning(f"{len(validation['unknown_wishes'])} ønske(r) peger på en lærer, der ikke findes i lærerarket. Ret initialerne/navnet i elevtabellen.")
            st.dataframe(pd.DataFrame(validation["unknown_wishes"]), width="stretch", hide_index=True)
        if validation["missing_subjects"]:
            st.warning(f"{len(validation['missing_subjects'])} elevfag har ingen lærer i lærerarket. Ret fagene i elevtabellen eller tilføj faget på en lærer.")
            st.dataframe(pd.DataFrame(validation["missing_subjects"]).drop_duplicates(), width="stretch", hide_index=True)
        if validation["incompatible_wishes"]:
            st.info("Nogle ønsker findes, men den ønskede lærer har ikke et af elevens fag. Kontrollér fag og ønsker i tabellerne.")
            st.dataframe(pd.DataFrame(validation["incompatible_wishes"]).drop_duplicates(), width="stretch", hide_index=True)
        if not any(validation.values()):
            st.success("Elevønsker, elevfag og lærernes fag ser konsistente ud.")

        readiness_errors = data_readiness(students, teachers)
        if st.session_state["input_loaded"] and not readiness_errors:
            st.success("Data er godkendt og klar til beregning.")
        else:
            if readiness_errors:
                st.error("Fordeling er låst, indtil data er rettet: " + " ".join(readiness_errors[:8]))
            if st.button("Godkend data til beregning", type="primary", key="approve_input_data"):
                if readiness_errors:
                    st.error("Data kan ikke godkendes endnu. Ret fejlene ovenfor først.")
                else:
                    st.session_state["input_loaded"] = True
                    st.session_state["solution"] = None
                    st.success("Data er godkendt. Gå til fanen Beregn fordeling.")
                    st.rerun()

    with rules_tab:
        st.subheader("Fordelingsregler")
        rule_cols = st.columns(2)
        with rule_cols[0]:
            use_global = st.checkbox("Brug global K-grænse", value=True, key="main_use_global")
            K = st.slider("Maksimalt antal elever pr. lærer (K)", 1, 50, 18, key="main_k")
            double_limit = st.slider("Maksimale dobbeltvejledninger pr. lærer (I)", 0, 50, 0, key="main_double_limit")
            allow_over = st.checkbox("Tillad overskridelse af max (op til K)", value=True, disabled=not use_global, key="main_allow_over")
        with rule_cols[1]:
            lock_max = st.checkbox("Lås lærernes max-tal fast", value=False, key="main_lock_max")
            prioritize_pairs = st.checkbox("Prioritér samme vejlederpar", value=True, key="main_prioritize_pairs")
            prioritize_classes = st.checkbox("Saml elever fra samme klasse/hold", value=True, key="main_prioritize_classes")
            attempts = st.slider("Algoritmedybde for fordeling", 20, 180, 120, 10, key="main_attempts", help="Algoritmen forsøger at placere eleverne hos lærere, der dækker fagene, samtidig med at kapacitet og elevønsker respekteres. Tallet angiver, hvor mange forslag der afprøves; højere værdi kan give et bedre resultat, men tager længere tid.")
        st.caption("K er den globale grænse. I bestemmer, hvor mange elever der må få samme lærer til begge fag.")

        st.subheader("Lærernes max-tal")
        st.caption("Redigér Max direkte. Søgning filtrerer kun visningen.")
        capacity_query = st.text_input("Søg i lærere", placeholder="Navn, initialer eller fag", key="main_capacity_query").casefold()
        sorted_teachers = sorted(teachers, key=lambda item: teacher_label(item["id"], teacher_map))
        capacity_rows = [{"Lærer": teacher_label(teacher["id"], teacher_map), "Initialer": teacher["id"], "Fag": " · ".join(teacher["subjects"]), "Max": int(capacities.get(teacher["id"], 0))} for teacher in sorted_teachers]
        if capacity_query:
            capacity_rows = [row for row in capacity_rows if capacity_query in " ".join(str(value) for value in row.values()).casefold()]
        edited_capacities = st.data_editor(
            pd.DataFrame(capacity_rows), width="stretch", hide_index=True,
            height=min(430, 80 + max(1, len(capacity_rows)) * 35), disabled=["Lærer", "Initialer", "Fag"],
            column_config={"Max": st.column_config.NumberColumn("Max", min_value=0, max_value=50, step=1, required=True)},
            key=f"main_capacity_editor_{st.session_state['capacity_revision']}",
        )
        for _, row in edited_capacities.iterrows():
            teacher_id = normal_key(row["Initialer"])
            if teacher_id in capacities:
                capacities[teacher_id] = max(0, min(50, int(row["Max"] or 0)))
        action_cols = st.columns([1, 1, 2])
        with action_cols[0]:
            if st.button("Nulstil max", key="reset_max"):
                _, _, default_caps = embedded_data()
                st.session_state["capacities"] = {teacher["id"]: default_caps.get(teacher["id"], 0) for teacher in teachers}
                st.session_state["capacity_revision"] += 1
                st.rerun()
        with action_cols[1]:
            if st.button("Foreslå kapaciteter", key="suggest_max") and students and teachers:
                progress_bar = st.progress(0.0, text="Beregner kapacitetsforslag: 0 %")
                with st.spinner("Beregner forslag til kapaciteter …"):
                    suggestion = optimize(
                        students, teachers, {teacher["id"]: 50 for teacher in teachers}, K, double_limit, use_global, allow_over,
                        lock_max, prioritize_pairs, prioritize_classes, max(30, attempts // 2),
                        progress_callback=lambda value: update_algorithm_progress(progress_bar, "Beregner kapacitetsforslag", value),
                    )
                update_algorithm_progress(progress_bar, "Kapacitetsforslag færdigt", 1.0)
                st.session_state["capacities"] = dict(suggestion["loads"])
                st.session_state["capacity_revision"] += 1
                st.rerun()
        with action_cols[2]:
            st.caption("Ændringer i data eller regler påvirker resultatet, når du beregner igen.")

    with result_tab:
        st.subheader("Beregn fordeling")
        st.write("Gennemgå eventuelle advarsler i datafanen, og beregn derefter fordelingen. Resultatet kan downloades som et Excel-ark med flere faner.")
        validation = input_validation(students, teachers)
        readiness_errors = data_readiness(students, teachers)
        can_calculate = st.session_state["input_loaded"] and not readiness_errors
        if not st.session_state["input_loaded"]:
            st.warning("Data er ikke godkendt endnu. Gå til datafanen, indlæs eller ret data, og tryk **Godkend data til beregning**.")
        elif readiness_errors:
            st.error("Fordeling er låst, fordi data ikke er på korrekt form: " + " ".join(readiness_errors[:8]))
        if st.button("Beregn fordeling", type="primary", key="calculate_distribution", disabled=not can_calculate):
            if not students:
                st.error("Der er ingen elever at fordele.")
            elif not teachers:
                st.error("Der er ingen lærere at fordele på.")
            else:
                progress_bar = st.progress(0.0, text="Beregner fordeling: 0 %")
                with st.spinner("Beregner flere mulige fordelinger …"):
                    st.session_state["solution"] = optimize(
                        students, teachers, capacities, K, double_limit, use_global, allow_over, lock_max,
                        prioritize_pairs, prioritize_classes, attempts,
                        progress_callback=lambda value: update_algorithm_progress(progress_bar, "Beregner fordeling", value),
                    )
                update_algorithm_progress(progress_bar, "Fordeling færdig", 1.0)
                st.rerun()

        solution = st.session_state.get("solution")
        if solution is None:
            st.info("Tryk på **Beregn fordeling** for at se resultatet.")
            return

        stats = solution["stats"]
        if stats["unassigned"] or stats["over_teachers"] or stats["capacity_blocked_slots"]:
            messages = []
            if stats["unassigned"]:
                messages.append(f"{stats['unassigned']} elever mangler mindst én vejleder.")
            if stats["capacity_blocked_slots"]:
                messages.append(f"{stats['capacity_blocked_slots']} fagpladser blev blokeret af lærermax.")
            if stats["over_teachers"]:
                messages.append(f"{stats['over_teachers']} lærere ligger over deres individuelle max.")
            st.warning(" ".join(messages))

        metric_data = [(len(students) - stats["unassigned"], "Elever fordelt"), (stats["both"], "Begge ønsker"), (stats["one"] + stats["both"], "Mindst ét ønske"), (stats["none"], "Ingen ønsker"), (stats["capacity_blocked_students"], "Kapacitetsblokerede")]
        cols = st.columns(5)
        for column, (value, label) in zip(cols, metric_data):
            column.markdown(f'<div class="metric-card"><span class="metric-value">{value}</span><span class="metric-label">{label}</span></div>', unsafe_allow_html=True)
        st.caption(f"Resultat ved {'K=' + str(solution['K']) if solution['use_global_k'] else 'lærernes egne max-tal'} og I={solution['double_limit']}. {stats['both']} elever får begge ønsker opfyldt.")
        export_sheet_options = ["Fordeling", "Lærerbelastning", "Fagstatistik", "Ikke tildelte", "Elevdata", "Lærerdata"]
        selected_export_sheets = st.multiselect(
            "Vælg ark til Excel-filen",
            options=export_sheet_options,
            default=export_sheet_options[:4],
            help="Vælg ét eller flere ark. Rækkefølgen i Excel følger listen ovenfor.",
            key="selected_export_sheets",
        )
        if not selected_export_sheets:
            st.warning("Vælg mindst ét ark, før Excel-filen kan downloades.")
            result_export = None
        else:
            result_export = try_excel_export(lambda: make_export(students, teachers, solution, capacities, solution["use_global_k"], selected_export_sheets))
        if result_export is not None:
            st.download_button("Download fordeling som Excel", data=result_export, file_name="vejlederfordeling.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary", key="download_result")
        elif selected_export_sheets:
            st.error("Excel-download kræver openpyxl. Installer projektets requirements.txt for at aktivere eksporten.")

        result_students, result_teachers, result_subjects, result_missing = st.tabs(["Fordeling", "Lærere", "Fagstatistik", "Ikke tildelte"])
        with result_students:
            query = st.text_input("Søg i elever", key="student_query").casefold()
            rows = []
            for index, student in enumerate(students):
                assigned = solution["assignments"][index]
                search = " ".join([student["name"], student.get("className", ""), *student["subjects"], *(assigned_item or "" for assigned_item in assigned)]).casefold()
                if query and query not in search:
                    continue
                rows.append({"Elev": student["name"], "Klasse": student.get("className", ""), "Fag 1": student["subjects"][0], "Vejleder 1": teacher_label(assigned[0], teacher_map), "Fag 2": student["subjects"][1], "Vejleder 2": teacher_label(assigned[1], teacher_map), "Ønsker": f"{wished_count(student, assigned)}/2", "Projekttitel": student.get("projectTitle", "")})
            st.dataframe(pd.DataFrame(rows), width="stretch", hide_index=True, height=560)
            st.caption("Fordelingsarket i Excel indeholder også projektbeskrivelse og vejlederinitialer.")
            with st.expander("Manuel justering af vejledere"):
                edit_rows = []
                for index, student in enumerate(students):
                    assigned = solution["assignments"][index]
                    edit_rows.append({"Elev": student["name"], "Fag 1": student["subjects"][0], "Vejleder 1": assigned[0] or "", "Fag 2": student["subjects"][1], "Vejleder 2": assigned[1] or ""})
                edited = st.data_editor(pd.DataFrame(edit_rows), width="stretch", hide_index=True, disabled=["Elev", "Fag 1", "Fag 2"], key="manual_editor")
                if st.button("Gem manuelle ændringer", key="save_manual"):
                    errors = []
                    valid_ids = set(teacher_map)
                    for index, row in edited.iterrows():
                        values = []
                        for slot in (1, 2):
                            value = normal_key(row[f"Vejleder {slot}"])
                            if value and value not in valid_ids:
                                errors.append(f"Række {index + 1}: ukendt lærer {value}")
                            values.append(value or None)
                        if not errors:
                            solution["assignments"][index] = values
                    if errors:
                        st.error("; ".join(errors[:5]))
                    else:
                        solution["loads"] = {teacher_id: 0 for teacher_id in teacher_map}
                        solution["double_loads"] = {teacher_id: 0 for teacher_id in teacher_map}
                        for assigned in solution["assignments"]:
                            for teacher_id in set(item for item in assigned if item):
                                solution["loads"][teacher_id] += 1
                            if assigned[0] and assigned[0] == assigned[1]:
                                solution["double_loads"][assigned[0]] += 1
                        solution["stats"] = score_solution(students, teachers, capacities, solution["assignments"], solution["loads"], solution["K"], solution["use_global_k"], prioritize_pairs, prioritize_classes)
                        st.session_state["solution"] = solution
                        st.success("Fordelingen er opdateret.")
                        st.rerun()
        with result_teachers:
            teacher_rows = []
            for teacher in sorted(teachers, key=lambda item: (-solution["loads"].get(item["id"], 0), teacher_label(item["id"], teacher_map))):
                teacher_id = teacher["id"]
                limit = min(solution["K"], capacities.get(teacher_id, 0)) if solution["use_global_k"] else capacities.get(teacher_id, 0)
                pupils = [students[index]["name"] for index, assigned in enumerate(solution["assignments"]) if teacher_id in assigned]
                teacher_rows.append({"Lærer": teacher_label(teacher_id, teacher_map), "Initialer": teacher_id, "Fag": " · ".join(teacher["subjects"]), "Elever": len(pupils), "Max": limit, "Status": "Over max" if len(pupils) > limit else "OK", "Dobbelt": solution["double_loads"].get(teacher_id, 0), "Elevliste": ", ".join(pupils)})
            st.dataframe(pd.DataFrame(teacher_rows), width="stretch", hide_index=True, height=600)
        with result_subjects:
            st.dataframe(subject_stats(students, teachers, capacities, solution["K"], solution["use_global_k"]), width="stretch", hide_index=True, height=560)
        with result_missing:
            missing = []
            overview = {}
            candidates = build_candidates(teachers)
            for index, student in enumerate(students):
                for slot, subject in enumerate(student["subjects"]):
                    if not solution["assignments"][index][slot]:
                        missing.append({
                            "Nr.": len(missing) + 1,
                            "student_index": index,
                            "slot": slot,
                            "Elev": student["name"],
                            "Klasse": student.get("className", ""),
                            "Fag": subject,
                            "Ønskevejledere": " · ".join(student.get("wishes", [])) or "Ingen registrerede ønsker",
                            "Mulige lærere": " · ".join(candidates.get(canonical_subject(subject), [])) or "Ingen",
                            "Tildel vejleder": "",
                        })
                        item = overview.setdefault(index, {"Elev": student["name"], "Klasse": student.get("className", ""), "Fag der mangler": [], "Antal": 0})
                        item["Fag der mangler"].append(subject)
                        item["Antal"] += 1

            if not missing:
                st.success("Alle elever og fag er tildelt.")
            else:
                st.warning(f"{len(overview)} elever mangler vejleder til {len(missing)} fagplads(er). Vælg en lærer direkte nedenfor og tryk Gem tildelinger.")
                st.markdown("**Overblik pr. elev**")
                overview_rows = list(overview.values())
                for item in overview_rows:
                    item["Fag der mangler"] = " · ".join(item["Fag der mangler"])
                st.dataframe(pd.DataFrame(overview_rows), width="stretch", hide_index=True, height=min(280, 70 + len(overview_rows) * 38))

                st.markdown("**Tildel manglende vejledere direkte**")
                st.caption("Vælg lærerinitialer i kolonnen Tildel vejleder. Kun lærere, der har det valgte fag, kan gemmes.")
                missing_frame = pd.DataFrame(missing)
                edited_missing = st.data_editor(
                    missing_frame[["Nr.", "Elev", "Klasse", "Fag", "Ønskevejledere", "Mulige lærere", "Tildel vejleder"]],
                    width="stretch", hide_index=True, height=min(560, 90 + len(missing) * 38),
                    disabled=["Nr.", "Elev", "Klasse", "Fag", "Ønskevejledere", "Mulige lærere"],
                    column_config={
                        "Tildel vejleder": st.column_config.SelectboxColumn(
                            "Tildel vejleder", options=[""] + sorted(teacher_map), required=False,
                            help="Vælg en lærer, der har dette fag.",
                        )
                    },
                    key="missing_assignments_editor",
                )
                if st.button("Gem tildelinger", type="primary", key="save_missing_assignments"):
                    errors = []
                    changes = []
                    for _, row in edited_missing.iterrows():
                        row_number = int(row["Nr."]) - 1
                        if row_number < 0 or row_number >= len(missing):
                            continue
                        item = missing[row_number]
                        teacher_id = normal_key(row["Tildel vejleder"])
                        if not teacher_id:
                            continue
                        if teacher_id not in teacher_map:
                            errors.append(f"{item['Elev']}: ukendt lærer {teacher_id}.")
                        elif teacher_id not in candidates.get(canonical_subject(item["Fag"]), []):
                            errors.append(f"{item['Elev']}: {teacher_id} har ikke faget {item['Fag']}.")
                        else:
                            changes.append((item["student_index"], item["slot"], teacher_id))
                    if errors:
                        st.error(" ".join(errors[:8]))
                    else:
                        for student_index, slot, teacher_id in changes:
                            solution["assignments"][student_index][slot] = teacher_id
                        solution["loads"] = {teacher_id: 0 for teacher_id in teacher_map}
                        solution["double_loads"] = {teacher_id: 0 for teacher_id in teacher_map}
                        for assigned in solution["assignments"]:
                            for teacher_id in set(item for item in assigned if item):
                                solution["loads"][teacher_id] += 1
                            if assigned[0] and assigned[0] == assigned[1]:
                                solution["double_loads"][assigned[0]] += 1
                        solution["stats"] = score_solution(students, teachers, capacities, solution["assignments"], solution["loads"], solution["K"], solution["use_global_k"], prioritize_pairs, prioritize_classes)
                        st.session_state["solution"] = solution
                        st.success(f"{len(changes)} tildeling(er) er gemt i den aktuelle fordeling.")
                        st.rerun()

        st.divider()
        st.subheader("Eksport")
        if st.button("Forbered Word-filer", key="prepare_docx"):
            with st.spinner("Genererer Word-filer …"):
                try:
                    st.session_state["docx_zip"] = make_docx_zip(students, teachers, solution)
                except ImportError:
                    st.error("Word-eksport kræver python-docx.")
        if st.session_state.get("docx_zip"):
            st.download_button("Download SOP-filer (.zip)", data=st.session_state["docx_zip"], file_name="SOP-filer.zip", mime="application/zip", key="download_docx")


def main_v2() -> None:
    """Trinopdelt arbejdsgang med navigation i venstremenuen."""
    st.set_page_config(page_title="SOPtima · vejlederfordeling og tidsplan", page_icon="🎓", layout="wide", initial_sidebar_state="expanded")
    st.markdown("""
    <style>
      :root {
        --ink:#152326; --muted:#5c6f72; --brand:#08756f; --brand-dark:#055852;
        --brand-soft:#e8f5f2; --line:#d8e5e2; --paper:#ffffff; --warm:#fffaf1;
      }
      .stApp {
        background:
          radial-gradient(circle at 8% -4%,rgba(17,145,134,.14),transparent 28rem),
          radial-gradient(circle at 94% 8%,rgba(233,174,73,.10),transparent 24rem),
          linear-gradient(180deg,#f7fbfa 0,#fffdf8 30rem,#f9fbfa 100%);
        color:var(--ink);
      }
      [data-testid='stHeader'] { background:rgba(247,251,250,.82); backdrop-filter:blur(12px); }
      [data-testid='stAppViewBlockContainer'] { max-width:1480px; padding-top:1.25rem; padding-bottom:4rem; }
      section[data-testid='stSidebar'] { background:linear-gradient(180deg,#eff8f6 0,#fbfcfa 58%,#fffaf1 100%); border-right:1px solid var(--line); }
      section[data-testid='stSidebar'] [data-testid='stSidebarContent'] { padding-top:1rem; }
      h1,h2,h3 { font-family:Georgia,'Times New Roman',serif; letter-spacing:-.025em; color:var(--ink); }
      h2 { margin-top:1.35rem; }
      p,li,label { line-height:1.52; }
      .soptima-brand { display:flex; align-items:center; gap:.7rem; margin:.2rem 0 1rem; }
      .soptima-brand-icon { display:grid; place-items:center; width:2.35rem; height:2.35rem; border-radius:.78rem; background:linear-gradient(135deg,var(--brand-dark),#12a195); color:white; box-shadow:0 8px 20px rgba(8,117,111,.22); font:bold 1.08rem Georgia,serif; }
      .soptima-brand-name { font:700 1.34rem/1 Georgia,serif; color:var(--ink); }
      .soptima-brand-sub { color:var(--muted); font-size:.72rem; margin-top:.12rem; }
      .soptima-hero { position:relative; overflow:hidden; display:flex; gap:1rem; align-items:flex-start; padding:1.35rem 1.45rem; margin:0 0 1.35rem; border:1px solid rgba(8,117,111,.16); border-radius:20px; background:linear-gradient(120deg,rgba(255,255,255,.97),rgba(232,245,242,.9)); box-shadow:0 16px 38px rgba(22,54,51,.07); }
      .soptima-hero::after { content:''; position:absolute; width:11rem; height:11rem; right:-4.5rem; top:-5.5rem; border-radius:50%; background:rgba(18,161,149,.08); pointer-events:none; }
      .soptima-hero-icon { flex:0 0 auto; display:grid; place-items:center; width:2.8rem; height:2.8rem; border-radius:.9rem; background:var(--brand); color:white; font-size:1.25rem; box-shadow:0 8px 22px rgba(8,117,111,.22); }
      .soptima-kicker { color:var(--brand-dark); font-size:.7rem; font-weight:800; letter-spacing:.14em; text-transform:uppercase; margin:.05rem 0 .3rem; }
      .soptima-hero h1 { font-size:clamp(1.65rem,3vw,2.45rem); line-height:1.06; margin:0 0 .45rem; padding:0; }
      .soptima-hero p { max-width:54rem; color:var(--muted); font-size:.98rem; line-height:1.55; margin:0; }
      .source-pill { display:inline-flex; align-items:center; gap:.35rem; margin-top:.65rem; padding:.3rem .62rem; border-radius:999px; background:rgba(8,117,111,.09); color:var(--brand-dark); font-size:.72rem; font-weight:750; }
      .sidebar-card { margin:.6rem 0 .8rem; padding:.8rem .86rem; border:1px solid var(--line); border-radius:13px; background:rgba(255,255,255,.76); box-shadow:0 7px 20px rgba(22,54,51,.035); }
      .sidebar-card-title { color:var(--ink); font-weight:750; font-size:.82rem; margin-bottom:.45rem; }
      .sidebar-row { display:flex; justify-content:space-between; gap:.7rem; color:var(--muted); font-size:.76rem; padding:.2rem 0; }
      .sidebar-row strong { color:var(--ink); font-weight:700; text-align:right; }
      .metric-card { border:1px solid var(--line); border-radius:15px; padding:1rem 1.12rem; background:rgba(255,255,255,.9); box-shadow:0 8px 24px rgba(22,54,51,.04); min-height:96px; transition:transform .16s ease,box-shadow .16s ease; }
      .metric-card:hover { transform:translateY(-2px); box-shadow:0 12px 28px rgba(22,54,51,.08); }
      .metric-value { display:block; font:700 2rem/1 Georgia,serif; color:var(--ink); }
      .metric-label { display:block; margin-top:.45rem; color:var(--muted); font-size:.68rem; font-weight:800; letter-spacing:.07em; text-transform:uppercase; }
      .section-note { padding:.82rem 1rem; margin:.4rem 0 1rem; border-left:4px solid var(--brand); border-radius:0 10px 10px 0; background:rgba(232,245,242,.72); color:#344f51; }
      div[data-testid='stDataFrame'],div[data-testid='stDataEditor'] { border:1px solid var(--line); border-radius:14px; overflow:hidden; box-shadow:0 7px 22px rgba(22,54,51,.035); }
      [data-testid='stFileUploaderDropzone'] { border:1px dashed #8db6b0; border-radius:14px; background:rgba(232,245,242,.52); }
      div.stButton > button,div.stDownloadButton > button { border-radius:11px; min-height:2.55rem; font-weight:700; transition:transform .15s ease,box-shadow .15s ease,border-color .15s ease; }
      div.stButton > button:hover,div.stDownloadButton > button:hover { transform:translateY(-1px); border-color:var(--brand); box-shadow:0 7px 18px rgba(8,117,111,.12); }
      div.stButton > button[kind='primary'] { border:0; background:linear-gradient(135deg,var(--brand-dark),#0a8d84); box-shadow:0 8px 20px rgba(8,117,111,.19); }
      div.stButton > button:focus-visible,div.stDownloadButton > button:focus-visible,[role='radio']:focus-visible { outline:3px solid #e09a2f!important; outline-offset:3px; }
      .stTabs [data-baseweb='tab-list'] { gap:.35rem; border-bottom:1px solid var(--line); }
      .stTabs [data-baseweb='tab'] { border-radius:9px 9px 0 0; padding:.48rem .8rem; }
      [data-testid='stAlert'] { border-radius:13px; }
      [data-testid='stSidebar'] [role='radiogroup'] label { border-radius:10px; padding:.3rem .45rem; margin:.08rem 0; }
      [data-testid='stSidebar'] [role='radiogroup'] label:hover { background:rgba(8,117,111,.07); }
      footer { visibility:hidden; }
      @media (max-width:700px) {
        [data-testid='stAppViewBlockContainer'] { padding-left:1rem; padding-right:1rem; }
        .soptima-hero { padding:1rem; border-radius:16px; }
        .soptima-hero-icon { width:2.4rem; height:2.4rem; }
        .metric-card { min-height:auto; }
      }
      @media (prefers-reduced-motion:reduce) { * { scroll-behavior:auto!important; transition:none!important; animation:none!important; } }
    </style>
    """, unsafe_allow_html=True)
    for key, default in (
        ("students", []),
        ("teachers", []),
        ("capacities", {}),
        ("capacity_revision", 0),
        ("student_revision", 0),
        ("teacher_revision", 0),
        ("input_loaded", False),
        ("solution", None),
        ("v2_schedule", None),
        ("v2_schedule_signature", None),
        ("v2_docx_zip", None),
        ("v2_stale_notice", ""),
        ("v2_confirm_demo_restore", False),
        ("v2_demo_mode", False),
        ("v2_app_initialized", False),
        ("v2_startup_error", ""),
    ):
        st.session_state.setdefault(key, default)

    if not st.session_state["v2_app_initialized"]:
        if not st.session_state["students"] and not st.session_state["teachers"]:
            try:
                activate_demo_data()
            except Exception as error:
                st.session_state["v2_startup_error"] = f"Demodata kunne ikke indlæses: {error}"
        st.session_state["v2_app_initialized"] = True

    if st.session_state.get("v2_flash_message"):
        st.toast(st.session_state.pop("v2_flash_message"), icon="✅")
    if st.session_state.get("v2_flash_warning"):
        st.toast(st.session_state.pop("v2_flash_warning"), icon="⚠️")

    students = st.session_state["students"]
    teachers = st.session_state["teachers"]
    capacities = st.session_state["capacities"]
    resolve_wishes(students, teachers)
    teacher_map = {teacher["id"]: teacher for teacher in teachers}

    K = int(st.session_state.get("v2_k", 18))
    double_limit = int(st.session_state.get("v2_double_limit", 0))
    use_global = bool(st.session_state.get("v2_use_global", True))
    allow_over = bool(st.session_state.get("v2_allow_over", True))
    lock_max = bool(st.session_state.get("v2_lock_max", False))
    prioritize_pairs = bool(st.session_state.get("v2_prioritize_pairs", True))
    prioritize_classes = bool(st.session_state.get("v2_prioritize_classes", True))
    attempts = int(st.session_state.get("v2_attempts", 120))
    current_distribution_signature = distribution_signature(
        students, teachers, capacities, K, double_limit, use_global, allow_over,
        lock_max, prioritize_pairs, prioritize_classes, attempts,
    )
    existing_solution = st.session_state.get("solution")
    if existing_solution is not None and existing_solution.get("configuration_signature") != current_distribution_signature:
        invalidate_derived_state("Indstillingerne eller datagrundlaget er ændret. Beregn en ny fordeling, før resultatet bruges.")

    readiness_errors = data_readiness(students, teachers)
    if not students:
        current_step = 1
        next_step = "Start med at indlæse elevdata."
    elif not teachers:
        current_step = 2
        next_step = "Næste skridt er at indlæse lærerdata."
    elif not st.session_state["input_loaded"] or readiness_errors:
        current_step = 2
        next_step = "Godkend data øverst, og gå derefter til trin 4 for at beregne."
    elif st.session_state.get("solution") is None:
        current_step = 4
        next_step = "Data er godkendt — gå videre til Beregn."
    else:
        current_step = 5
        next_step = "Fordelingen er beregnet — gennemgå resultatet og eksportér."

    process_steps = [
        "1 · Elevdata",
        "2 · Lærerdata",
        "3 · Regler og max",
        "4 · Beregn",
        "5 · Resultat og eksport",
        "6 · Tidsplan",
    ]
    step_labels = {
        "1 · Elevdata": "👥  Elevdata",
        "2 · Lærerdata": "🧑‍🏫  Lærerdata",
        "3 · Regler og max": "⚙️  Regler og max",
        "4 · Beregn": "✨  Beregn",
        "5 · Resultat og eksport": "📊  Resultat og eksport",
        "6 · Tidsplan": "🗓️  Tidsplan",
    }

    with st.sidebar:
        st.markdown(
            '<div class="soptima-brand"><div class="soptima-brand-icon" aria-hidden="true">S</div>'
            '<div><div class="soptima-brand-name">SOPtima</div>'
            '<div class="soptima-brand-sub">Vejlederfordeling og tidsplan</div></div></div>',
            unsafe_allow_html=True,
        )
        st.progress(current_step / 6, text=f"Arbejdsstatus · trin {current_step} af 6")
        st.caption(next_step)
        active_step = st.radio(
            "Gå til trin",
            process_steps,
            index=current_step - 1,
            key="v2_active_step",
            label_visibility="collapsed",
            format_func=lambda step: step_labels[step],
        )
        if students:
            student_detail = f"{len(students)} elever"
            if readiness_errors:
                student_detail += f" · {len(readiness_errors)} fejl"
            elif st.session_state["input_loaded"]:
                student_detail += " · godkendt"
            else:
                student_detail += " · ikke godkendt"
        else:
            student_detail = "Ikke indlæst"
        teacher_detail = f"{len(teachers)} lærere" if teachers else "Ikke indlæst"
        solution_detail = "Beregnet" if st.session_state.get("solution") is not None else "Ikke beregnet"
        source_detail = "Fiktive demodata" if st.session_state.get("v2_demo_mode") else "Egne eller redigerede data"
        st.markdown(
            '<div class="sidebar-card"><div class="sidebar-card-title">Status</div>'
            f'<div class="sidebar-row"><span>👥 Elever</span><strong>{html.escape(student_detail)}</strong></div>'
            f'<div class="sidebar-row"><span>🧑‍🏫 Lærere</span><strong>{html.escape(teacher_detail)}</strong></div>'
            f'<div class="sidebar-row"><span>✨ Fordeling</span><strong>{html.escape(solution_detail)}</strong></div>'
            f'<div class="sidebar-row"><span>◉ Datakilde</span><strong>{html.escape(source_detail)}</strong></div></div>',
            unsafe_allow_html=True,
        )
        if students and teachers and not readiness_errors and not st.session_state["input_loaded"]:
            if st.button("✓ Godkend data", type="primary", key="v2_sidebar_approve_data", use_container_width=True):
                st.session_state["input_loaded"] = True
                invalidate_derived_state()
                st.session_state["v2_stale_notice"] = ""
                st.rerun()
        st.divider()
        st.markdown("#### Demo")
        st.caption("350 fiktive elever og 85 fiktive lærere indlæses automatisk første gang.")
        if not st.session_state["v2_confirm_demo_restore"]:
            if st.button("↻ Gendan demodata", key="v2_restore_demo", use_container_width=True):
                st.session_state["v2_confirm_demo_restore"] = True
                st.rerun()
        else:
            st.warning("Dette erstatter aktuelle elev- og lærerdata i sessionen.")
            confirm_columns = st.columns(2)
            with confirm_columns[0]:
                if st.button("Gendan", type="primary", key="v2_restore_demo_confirm", use_container_width=True):
                    try:
                        activate_demo_data()
                        st.session_state["v2_confirm_demo_restore"] = False
                        st.session_state["v2_flash_message"] = "Demodata er gendannet og godkendt."
                        st.rerun()
                    except Exception as error:
                        st.error(f"Demodata kunne ikke gendannes: {error}")
            with confirm_columns[1]:
                if st.button("Annuller", key="v2_restore_demo_cancel", use_container_width=True):
                    st.session_state["v2_confirm_demo_restore"] = False
                    st.rerun()
        st.caption("Filer behandles kun i den aktuelle app-session. Download dine resultater, før sessionen lukkes.")

    if st.session_state.get("v2_startup_error"):
        st.error(st.session_state["v2_startup_error"])
    if st.session_state.get("v2_stale_notice"):
        st.warning(st.session_state["v2_stale_notice"], icon="⚠️")

    page_meta = {
        "1 · Elevdata": ("👥", "Datagrundlag", "Elevdata", "Indlæs, kontrollér og tilpas elevernes fag, ønsker og projektoplysninger."),
        "2 · Lærerdata": ("🧑‍🏫", "Datagrundlag", "Lærerdata", "Indlæs lærernes fag, hold og maksimumstal, og kontrollér at alle elevfag er dækket."),
        "3 · Regler og max": ("⚙️", "Indstillinger", "Regler og kapacitet", "Fastlæg de rammer, som den endelige vejlederfordeling skal overholde."),
        "4 · Beregn": ("✨", "Optimering", "Beregn fordeling", "Afprøv flere gyldige fordelinger og vælg automatisk den løsning, der samlet scorer bedst."),
        "5 · Resultat og eksport": ("📊", "Resultat", "Fordeling og eksport", "Gennemgå kvalitet, belastning og mangler, foretag justeringer og hent resultatet."),
        "6 · Tidsplan": ("🗓️", "Planlægning", "Tidsplan for vejledning", "Placér vejledninger parallelt uden lærerkonflikter, og tilføj pauser og frokost."),
    }
    page_icon, page_kicker, page_title, page_description = page_meta[active_step]
    source_badge = "● Demo klar" if st.session_state.get("v2_demo_mode") else "● Egne data"
    st.markdown(
        f'<section class="soptima-hero"><div class="soptima-hero-icon" aria-hidden="true">{page_icon}</div>'
        f'<div><div class="soptima-kicker">{html.escape(page_kicker)}</div>'
        f'<h1>{html.escape(page_title)}</h1><p>{html.escape(page_description)}</p>'
        f'<span class="source-pill">{html.escape(source_badge)}</span></div></section>',
        unsafe_allow_html=True,
    )

    if active_step == "2 · Lærerdata":
        st.subheader("Indlæs lærerdata")
        if students and teachers and not readiness_errors:
            if st.session_state["input_loaded"]:
                st.success("Data er godkendt og klar til beregning. Gå til trin 4 for at beregne fordelingen.")
            else:
                st.info("Data ser korrekt ud. Godkend data med knappen øverst i sidepanelet, og gå derefter til trin 4.")
        st.caption("Brug eventuelt elevdataens holdnavne som reference i lærerarket.")
        with st.expander("Se eksempel på lærerarket (kun illustration)", expanded=False):
            st.dataframe(pd.DataFrame([
                {"Initialer": "AB", "Lærernavn": "Anna Borg", "Fag 1": "Matematik", "Fag 2": "Fysik", "Fag 3": "Informatik", "Hold 1": "3a", "Hold 2": "3b", "Hold 3": "", "Hold 4": "", "Max": 18},
                {"Initialer": "CD", "Lærernavn": "Carsten Dahl", "Fag 1": "Dansk", "Fag 2": "Historie", "Fag 3": "Samfundsfag", "Hold 1": "2x", "Hold 2": "", "Hold 3": "", "Hold 4": "", "Max": 16},
            ]), width="stretch", hide_index=True)
            st.caption("Format: Initialer + mindst ét fag. Fag 1–3 vises som standard; op til 20 fag og 4 hold kan tilføjes.")
        st.caption("Næste skridt: vælg en lærerfil.")
        teacher_upload = st.file_uploader("Upload lærerdata (.xlsx, .xlsm eller .csv)", type=["xlsx", "xlsm", "csv"], key="v2_teachers_upload")
        if st.button("Indlæs lærerdata", type="primary" if teacher_upload is not None else "secondary", disabled=teacher_upload is None, key="v2_load_teachers"):
            try:
                replacing_demo = bool(st.session_state.get("v2_demo_mode"))
                parsed_teachers, parsed_caps = parse_teacher_upload(teacher_upload)
                if replacing_demo:
                    st.session_state["students"] = []
                    st.session_state["student_revision"] += 1
                st.session_state["teachers"] = parsed_teachers
                st.session_state["capacities"] = {teacher["id"]: parsed_caps.get(teacher["id"], 0) for teacher in parsed_teachers}
                st.session_state["teacher_revision"] += 1
                st.session_state["capacity_revision"] += 1
                invalidate_derived_state(require_approval=True)
                st.session_state["v2_demo_mode"] = False
                st.session_state["v2_stale_notice"] = ""
                st.session_state["v2_flash_message"] = (
                    f"Lærerdata indlæst: {len(parsed_teachers)} lærere. Demoeleverne er fjernet; indlæs nu elevdata."
                    if replacing_demo else f"Lærerdata indlæst: {len(parsed_teachers)} lærere."
                )
                st.rerun()
            except Exception as error:
                st.error(str(error))

        st.subheader("Lærerdata – redigerbart Excel-ark")
        st.caption("Redigér direkte i arket. Ét fag pr. kolonne; brug pilen nederst for flere lærere.")
        current_subject_count = max(3, max((len(teacher.get("subjects", [])) for teacher in teachers), default=0))
        current_hold_count = max((len(teacher.get("holds", [])) for teacher in teachers), default=0)
        setting_columns = st.columns(2)
        with setting_columns[0]:
            subject_column_count = st.number_input(
                "Antal fagkolonner i lærerarket",
                min_value=3,
                max_value=20,
                value=min(20, current_subject_count),
                step=1,
                help="Start med 3. Vælg et højere antal, hvis en lærer skal have flere fag.",
                key=f"v2_subject_column_count_{st.session_state['teacher_revision']}",
            )
        with setting_columns[1]:
            hold_column_count = st.number_input(
                "Antal holdkolonner i lærerarket",
                min_value=0,
                max_value=4,
                value=min(4, current_hold_count),
                step=1,
                help="Vælg 0–4 holdkolonner. Eksisterende hold bevares, hvis du vælger færre kolonner.",
                key=f"v2_hold_column_count_{st.session_state['teacher_revision']}",
            )
        st.caption("Fag: mindst 3 kolonner og op til 20. Hold: valgfrit, fra 0 til 4 kolonner. Ryd en holdcelle i arket, hvis et eksisterende hold skal fjernes.")
        edited_teachers = st.data_editor(
            teachers_to_frame(teachers, capacities, int(subject_column_count), int(hold_column_count)), width="stretch", hide_index=True, num_rows="dynamic",
            height=min(500, 80 + max(1, len(teachers)) * 38),
            column_config={"Max": st.column_config.NumberColumn("Max", min_value=0, max_value=50, step=1, required=True)},
            key=f"v2_teachers_editor_{st.session_state['teacher_revision']}",
        )
        new_teachers, new_capacities = frame_to_teachers(edited_teachers)
        if json.dumps(new_teachers, sort_keys=True, ensure_ascii=False) != json.dumps(teachers, sort_keys=True, ensure_ascii=False) or new_capacities != capacities:
            st.session_state["teachers"] = new_teachers
            st.session_state["capacities"] = new_capacities
            invalidate_derived_state(require_approval=True)
            st.session_state["v2_demo_mode"] = False
            teachers, capacities = new_teachers, new_capacities
            teacher_map = {teacher["id"]: teacher for teacher in teachers}
        no_subject_teachers = [teacher["id"] for teacher in teachers if not teacher.get("subjects")]
        if no_subject_teachers:
            st.warning(f"{len(no_subject_teachers)} lærer(e) mangler fag og kan ikke få elever: {', '.join(no_subject_teachers)}")
        zero_capacity_teachers = [teacher_label(teacher["id"], teacher_map) for teacher in teachers if int(capacities.get(teacher["id"], 0)) == 0]
        if zero_capacity_teachers:
            preview = ", ".join(zero_capacity_teachers[:8])
            suffix = " …" if len(zero_capacity_teachers) > 8 else ""
            st.warning(f"{len(zero_capacity_teachers)} lærer(e) har max 0 og kan ikke få elever: {preview}{suffix}. Ret Max eller brug max-forslaget i trin 3.")
        teacher_data_ready = bool(teachers) and not no_subject_teachers and all(normal_key(teacher.get("id")) for teacher in teachers)
        if teacher_data_ready:
            teacher_export = try_excel_export(
                lambda: make_teacher_export(teachers, capacities, int(subject_column_count), int(hold_column_count))
            )
            if teacher_export is not None:
                st.download_button(
                    "Download lærerdata som Excel",
                    data=teacher_export,
                    file_name="lærerdata.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="v2_download_teachers",
                )
        if students and teachers:
            st.subheader("Fagdækning blandt elevernes fag")
            st.dataframe(subject_coverage(students, teachers), width="stretch", hide_index=True)

        if students:
            readiness_errors = data_readiness(students, teachers)
            if st.session_state["input_loaded"] and not readiness_errors:
                st.success("Data er godkendt og klar til fordeling.")
            else:
                if readiness_errors:
                    st.error(f"Data kan ikke godkendes endnu. Der er {len(readiness_errors)} problem(er), som skal rettes.")
                    with st.expander("Vis alle problemer", expanded=True):
                        for error in readiness_errors:
                            st.markdown(f"- {error}")
                if st.button("Godkend data til fordeling", type="primary" if not readiness_errors else "secondary", disabled=bool(readiness_errors), key="v2_approve_data"):
                    st.session_state["input_loaded"] = True
                    invalidate_derived_state()
                    st.session_state["v2_stale_notice"] = ""
                    st.success("Data er godkendt. Gå videre til trin 3 · Regler og max.")
                    st.rerun()

    if active_step == "1 · Elevdata":
        st.subheader("Indlæs elevdata først")
        with st.expander("Se eksempel på elevarket (kun illustration)", expanded=False):
            st.dataframe(pd.DataFrame([
                {"Elevnavn": "Emma Jensen", "Klasse": "3a", "Fag 1": "Matematik", "Fag 2": "Fysik", "Ønskevejleder 1": "AB", "Ønskevejleder 2": "CD"},
                {"Elevnavn": "Noah Hansen", "Klasse": "3a", "Fag 1": "Dansk", "Fag 2": "Historie", "Ønskevejleder 1": "CD", "Ønskevejleder 2": ""},
            ]), width="stretch", hide_index=True)
            st.caption("Format: Elevnavn og Fag 1–2 er påkrævet. Klasse/Hold og ønskevejledere er valgfri; en manglende klasse vises som ukendt.")
        st.caption("Første skridt: vælg en elevfil.")
        student_upload = st.file_uploader("Upload elevdata (.xlsx, .xlsm eller .csv)", type=["xlsx", "xlsm", "csv"], key="v2_students_upload")
        if st.button("Indlæs elevdata", type="primary" if student_upload is not None else "secondary", disabled=student_upload is None, key="v2_load_students"):
            try:
                replacing_demo = bool(st.session_state.get("v2_demo_mode"))
                st.session_state["students"] = parse_students(read_uploaded_table(student_upload))
                if replacing_demo:
                    st.session_state["teachers"] = []
                    st.session_state["capacities"] = {}
                    st.session_state["teacher_revision"] += 1
                    st.session_state["capacity_revision"] += 1
                st.session_state["student_revision"] += 1
                invalidate_derived_state(require_approval=True)
                st.session_state["v2_demo_mode"] = False
                st.session_state["v2_stale_notice"] = ""
                st.session_state["v2_flash_message"] = (
                    f"Elevdata indlæst: {len(st.session_state['students'])} elever. Demolærerne er fjernet; indlæs nu lærerdata."
                    if replacing_demo else f"Elevdata indlæst: {len(st.session_state['students'])} elever."
                )
                st.rerun()
            except Exception as error:
                st.error(str(error))

        st.subheader("Elevdata – redigerbart Excel-ark")
        st.caption("Redigér direkte i arket. Vælg en lærer eller **Ingen ønsker**.")
        edited_students = st.data_editor(
            students_to_frame(students, teachers), width="stretch", hide_index=True, num_rows="dynamic",
            height=min(600, 80 + max(1, len(students)) * 38), disabled=["Elev-ID"],
            column_config={**wish_column_config(teachers), "Projektbeskrivelse": st.column_config.TextColumn("Projektbeskrivelse", width="large")},
            key=f"v2_students_editor_{st.session_state['student_revision']}",
        )
        try:
            new_students = frame_to_students(edited_students)
            resolve_wishes(new_students, teachers)
            if student_data_signature(new_students) != student_data_signature(students):
                st.session_state["students"] = new_students
                invalidate_derived_state(require_approval=True)
                st.session_state["v2_demo_mode"] = False
                students = new_students
        except ValueError as error:
            st.error(str(error))

        resolve_wishes(students, teachers)
        teacher_map = {teacher["id"]: teacher for teacher in teachers}
        validation = input_validation(students, teachers)
        if validation["unknown_wishes"]:
            unknown_frame = pd.DataFrame(validation["unknown_wishes"]).drop_duplicates()
            st.warning(f"{len(unknown_frame)} ønsker kan ikke genkendes. Vælg en lærer eller **Ingen ønsker** i elevarket:")
            st.dataframe(unknown_frame, width="stretch", hide_index=True)
        if validation["ambiguous_wishes"]:
            st.warning("Disse ønsker matcher flere lærere med samme navn. Vælg læreren med initialer:")
            st.dataframe(pd.DataFrame(validation["ambiguous_wishes"]), width="stretch", hide_index=True)
        if validation["missing_subjects"]:
            st.warning("Disse elevfag har ingen lærer endnu:")
            st.dataframe(pd.DataFrame(validation["missing_subjects"]).drop_duplicates(), width="stretch", hide_index=True)
        if validation["incompatible_wishes"] or validation["unknown_wishes"] or validation["ambiguous_wishes"]:
            if validation["incompatible_wishes"]:
                st.warning("Disse ønskede lærere underviser ikke i elevens fag:")
                st.dataframe(pd.DataFrame(validation["incompatible_wishes"]).drop_duplicates(), width="stretch", hide_index=True)
            suggestion_rows = incompatible_wish_suggestions(students, teachers)
            if suggestion_rows:
                st.markdown("**Ret ønskerne med forslag fra lærerlisten**")
                st.caption("Vælg en foreslået lærer, eller vælg **Ingen ønsker** for at fjerne ønsket. Forslagene dækker mindst ét af elevens fag.")
                teacher_labels = {teacher["id"]: teacher_label(teacher["id"], teacher_map) for teacher in teachers}
                selected_suggestions = {}
                suggestion_table = pd.DataFrame(suggestion_rows)[["Elev", "Nuværende ønske", "Elevens fag", "Mulige lærere"]]
                st.dataframe(suggestion_table, width="stretch", hide_index=True)
                st.markdown("**Vælg ny lærer**")
                for row_index, item in enumerate(suggestion_rows):
                    possible_options = [NO_WISHES_LABEL] + [teacher_labels[teacher_id] for teacher_id in item["_possible_ids"]]
                    default_index = 1 if len(possible_options) > 1 else 0
                    selected_suggestions[row_index] = st.selectbox(
                        f"{item['Elev']} · {item['Nuværende ønske']}", possible_options, index=default_index,
                        key=f"v2_wish_suggestion_{st.session_state['teacher_revision']}_{st.session_state['student_revision']}_{row_index}",
                    )
                if st.button("Gem foreslåede lærere", type="primary", key="v2_save_wish_suggestions"):
                    label_to_id = {normal_key(label): teacher_id for teacher_id, label in teacher_labels.items()}
                    candidate_map = build_candidates(teachers)
                    changed_wishes = {index: list(student.get("wishes", [])) for index, student in enumerate(students)}
                    errors = []
                    changes = 0
                    for row_index, selected in selected_suggestions.items():
                        selected = repair_text(selected)
                        if not selected:
                            continue
                        item = suggestion_rows[row_index]
                        student_index = item["_student_index"]
                        wish_index = item["_wish_index"]
                        if normal_key(selected) == normal_key(NO_WISHES_LABEL):
                            changed_wishes[student_index][wish_index] = None
                            changes += 1
                            continue
                        teacher_id = label_to_id.get(normal_key(selected))
                        if not teacher_id:
                            errors.append(f"{item['Elev']}: kunne ikke genkende {selected}.")
                            continue
                        valid_for_subject = any(teacher_id in candidate_map.get(canonical_subject(subject), []) for subject in students[student_index].get("subjects", []))
                        if not valid_for_subject:
                            errors.append(f"{item['Elev']}: {selected} underviser ikke i elevens fag.")
                            continue
                        changed_wishes[student_index][wish_index] = teacher_id
                        changes += 1
                    if errors:
                        st.error(" ".join(errors[:8]))
                    elif changes:
                        for student_index, wishes in changed_wishes.items():
                            students[student_index]["wishes"] = [wish for wish in wishes if wish]
                        invalidate_derived_state(require_approval=True)
                        st.session_state["v2_demo_mode"] = False
                        st.success(f"{changes} ønske(r) er opdateret med de valgte lærere.")
                        st.rerun()
        if validation["partial_wishes"]:
            st.caption(f"{len(validation['partial_wishes'])} ønsker matcher kun ét af elevens fag.")
            with st.expander("Se ønsker med delvist fagmatch"):
                st.dataframe(pd.DataFrame(validation["partial_wishes"]).drop_duplicates(), width="stretch", hide_index=True)
        if students and not any(validation[key] for key in ("unknown_wishes", "ambiguous_wishes", "missing_subjects", "incompatible_wishes")):
            st.success("Elevønsker, elevfag og lærernes fag ser konsistente ud.")
        if students and teachers:
            st.subheader("Kontrol: har hvert fag en mulig lærer?")
            st.dataframe(subject_coverage(students, teachers), width="stretch", hide_index=True)

        input_export = try_excel_export(lambda: make_input_export(students, teachers, capacities))
        if input_export is not None:
            st.download_button("Download aktuelle inputark som Excel", data=input_export, file_name="vejlederfordeling_input.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", key="v2_download_input")
        readiness_errors = data_readiness(students, teachers)
        if students and teachers and not readiness_errors and not st.session_state["input_loaded"]:
            st.divider()
            st.markdown("#### Klar til næste trin")
            st.caption("Data er konsistente. Godkend dem for at låse datagrundlaget til den næste beregning.")
            st.button(
                "Godkend data og fortsæt",
                type="primary",
                key="v2_approve_from_students",
                on_click=approve_data_and_navigate,
                args=("3 · Regler og max",),
            )

    if active_step == "3 · Regler og max":
        st.subheader("Fordelingsregler")
        st.markdown('<div class="section-note"><strong>Sådan hænger reglerne sammen:</strong> K er den absolutte globale grænse. Lærerens eget max er den normale grænse. En overskridelse af lærerens max kan kun ske op til K, når den er tilladt og max ikke er låst. I begrænser dobbeltvejledninger.</div>', unsafe_allow_html=True)
        rule_cols = st.columns(2)
        with rule_cols[0]:
            use_global = st.checkbox("Brug global K-grænse", value=True, key="v2_use_global")
            K = st.slider("Maksimalt antal elever pr. lærer (K)", 1, 50, 18, key="v2_k")
            double_limit = st.slider("Maksimale dobbeltvejledninger pr. lærer (I)", 0, 50, 0, key="v2_double_limit")
            allow_over = st.checkbox("Tillad overskridelse af max (op til K)", value=True, disabled=not use_global, key="v2_allow_over")
        with rule_cols[1]:
            lock_max = st.checkbox("Lås lærernes max-tal fast", value=False, key="v2_lock_max")
            prioritize_pairs = st.checkbox("Prioritér samme vejlederpar", value=True, key="v2_prioritize_pairs")
            prioritize_classes = st.checkbox("Saml elever fra samme klasse/hold", value=True, key="v2_prioritize_classes")
            attempts = st.slider("Algoritmedybde for fordeling", 20, 180, 120, 10, key="v2_attempts", help="Algoritmen forsøger at placere eleverne hos lærere, der dækker fagene, samtidig med at kapacitet og elevønsker respekteres. Tallet angiver, hvor mange forslag der afprøves; højere værdi kan give et bedre resultat, men tager længere tid.")
        st.caption("K er den globale grænse. I bestemmer, hvor mange elever der må få samme lærer til begge fag.")
        st.subheader("Lærernes max-tal")
        max_action_cols = st.columns([1, 3])
        with max_action_cols[0]:
            if st.button(
                "Foreslå lærernes max-tal",
                type="primary",
                key="v2_suggest_max_top",
                disabled=not bool(students and teachers),
                help="Beregner et forslag til hvert maksimum ud fra elevfordelingen. Du kan efterfølgende redigere tallene manuelt.",
            ):
                progress_bar = st.progress(0.0, text="Beregner max-forslag: 0 %")
                with st.spinner("Beregner forslag til lærernes max-tal …"):
                    suggestion = optimize(
                        students, teachers, {teacher["id"]: 50 for teacher in teachers}, K, double_limit, use_global, allow_over,
                        lock_max, prioritize_pairs, prioritize_classes, max(30, attempts // 2),
                        progress_callback=lambda value: update_algorithm_progress(progress_bar, "Beregner max-forslag", value),
                    )
                update_algorithm_progress(progress_bar, "Max-forslag færdigt", 1.0)
                st.session_state["capacities"] = dict(suggestion["loads"])
                st.session_state["capacity_revision"] += 1
                invalidate_derived_state("Max-tallene er opdateret. Beregn fordelingen igen efter din kontrol.")
                st.session_state["v2_demo_mode"] = False
                st.session_state["v2_flash_message"] = "Forslag til lærernes max-tal er indsat. Kontrollér og tilpas dem efter behov."
                st.rerun()
        with max_action_cols[1]:
            st.caption("Forslaget fordeler belastningen så jævnt som muligt. Kontrollér altid tallene, før du beregner den endelige fordeling.")
        capacity_query = st.text_input("Søg i lærere", placeholder="Navn, initialer eller fag", key="v2_capacity_query").casefold()
        capacity_rows = [{"Lærer": teacher_label(teacher["id"], teacher_map), "Initialer": teacher["id"], "Fag": " · ".join(teacher["subjects"]), "Max": int(capacities.get(teacher["id"], 0))} for teacher in sorted(teachers, key=lambda item: teacher_label(item["id"], teacher_map))]
        if capacity_query:
            capacity_rows = [row for row in capacity_rows if capacity_query in " ".join(str(value) for value in row.values()).casefold()]
        edited_capacities = st.data_editor(
            pd.DataFrame(capacity_rows), width="stretch", hide_index=True, disabled=["Lærer", "Initialer", "Fag"],
            column_config={"Max": st.column_config.NumberColumn("Max", min_value=0, max_value=50, step=1, required=True)}, key="v2_capacity_editor",
        )
        capacities_before_edit = dict(capacities)
        for _, row in edited_capacities.iterrows():
            teacher_id = normal_key(row["Initialer"])
            if teacher_id in capacities:
                capacities[teacher_id] = max(0, min(50, int(row["Max"] or 0)))
        if capacities != capacities_before_edit:
            st.session_state["capacities"] = capacities
            invalidate_derived_state("Lærernes max-tal er ændret. Beregn en ny fordeling for at få et aktuelt resultat.")
            st.session_state["v2_demo_mode"] = False

    can_calculate = st.session_state["input_loaded"] and not data_readiness(students, teachers)
    if active_step == "4 · Beregn":
        st.subheader("Beregn fordeling")
        summary_columns = st.columns(4)
        summary_columns[0].metric("Elever", len(students))
        summary_columns[1].metric("Lærere", len(teachers))
        summary_columns[2].metric("Global K", K if use_global else "Fra")
        summary_columns[3].metric("Dobbeltvejledning I", double_limit)
        with st.expander("Kontrollér indstillinger før beregning", expanded=False):
            st.write(
                f"**Kapacitetsregel:** {'Global K = ' + str(K) if use_global else 'Lærernes egne max-tal'} · "
                f"**Overskridelse op til K:** {'Ja' if use_global and allow_over and not lock_max else 'Nej'} · "
                f"**Låste max-tal:** {'Ja' if lock_max else 'Nej'} · **Algoritmedybde:** {attempts}"
            )
            st.write(f"**Prioritér lærerpar:** {'Ja' if prioritize_pairs else 'Nej'} · **Saml klasser/hold:** {'Ja' if prioritize_classes else 'Nej'}")
        if not st.session_state["input_loaded"]:
            st.warning("Godkend først data i trin 1 eller 2.")
        elif not can_calculate:
            calculation_errors = data_readiness(students, teachers)
            st.error("Fordeling er låst, fordi datagrundlaget indeholder fejl.")
            with st.expander("Vis fejl i datagrundlaget", expanded=True):
                for error in calculation_errors:
                    st.markdown(f"- {error}")
        else:
            st.success("Data er godkendt. Du kan beregne fordelingen.")
        if st.button("Beregn fordeling", type="primary" if can_calculate else "secondary", key="v2_calculate", disabled=not can_calculate):
            progress_bar = st.progress(0.0, text="Beregner fordeling: 0 %")
            with st.spinner("Beregner flere mulige fordelinger …"):
                calculated_solution = optimize(
                    students, teachers, capacities, K, double_limit, use_global, allow_over, lock_max,
                    prioritize_pairs, prioritize_classes, attempts,
                    progress_callback=lambda value: update_algorithm_progress(progress_bar, "Beregner fordeling", value),
                )
                calculated_solution["configuration_signature"] = distribution_signature(
                    students, teachers, capacities, K, double_limit, use_global, allow_over,
                    lock_max, prioritize_pairs, prioritize_classes, attempts,
                )
                assessment = assignment_assessment(
                    students, teachers, capacities, calculated_solution["assignments"], K, double_limit,
                    use_global, allow_over, lock_max,
                )
                if assessment["errors"]:
                    st.session_state["solution"] = None
                    st.error("Fordelingen bestod ikke den afsluttende regelkontrol: " + " ".join(assessment["errors"][:8]))
                else:
                    st.session_state["solution"] = calculated_solution
                    st.session_state["v2_stale_notice"] = ""
                    st.session_state["v2_schedule"] = None
                    st.session_state["v2_schedule_signature"] = None
            update_algorithm_progress(progress_bar, "Fordeling færdig", 1.0)
            if st.session_state.get("solution") is not None:
                st.success("Fordelingen er beregnet og regelkontrolleret. Gå til trin 5 for at gennemgå resultatet.")

    if active_step == "5 · Resultat og eksport":
        solution = st.session_state.get("solution")
        st.subheader("Resultat og eksport")
        if solution is None:
            if st.session_state.get("v2_stale_notice"):
                st.info("Resultatet er fjernet, fordi grundlaget er ændret. Gå til trin 4 og beregn igen.")
            else:
                st.info("Der er endnu ikke et resultat. Gå til trin 4 · Beregn.")
        else:
            stats = solution["stats"]
            result_assessment = assignment_assessment(
                students, teachers, capacities, solution["assignments"], solution["K"], solution["double_limit"],
                solution["use_global_k"], solution.get("allow_over_capacity", False), solution.get("lock_teacher_max", False),
            )
            if result_assessment["errors"]:
                st.error("Fordelingen har kritiske regelbrud: " + " ".join(result_assessment["errors"][:8]))
            elif stats["unassigned"] == 0 and not result_assessment["warnings"]:
                st.success("Fordelingen er komplet og overholder alle valgte hårde grænser.", icon="✅")
            if stats["unassigned"] or stats["over_teachers"] or stats["capacity_blocked_slots"]:
                messages = []
                if stats["unassigned"]: messages.append(f"{stats['unassigned']} elever mangler mindst én vejleder.")
                if stats["capacity_blocked_slots"]: messages.append(f"{stats['capacity_blocked_slots']} fagpladser blev blokeret af lærermax.")
                if stats["over_teachers"]: messages.append(f"{stats['over_teachers']} lærere ligger over deres normale individuelle max; kontrollér at dette er tilsigtet.")
                st.warning(" ".join(messages))
            metric_data = [(len(students) - stats["unassigned"], "Elever komplet fordelt"), (stats.get("all_wishes", stats["both"]), "Alle ønsker opfyldt"), (stats["one"] + stats["both"], "Mindst ét ønske opfyldt"), (stats["none"], "Uden opfyldt ønske"), (stats["capacity_blocked_students"], "Kapacitetsblokerede")]
            cols = st.columns(5)
            for column, (value, label) in zip(cols, metric_data):
                column.markdown(f'<div class="metric-card"><span class="metric-value">{value}</span><span class="metric-label">{label}</span></div>', unsafe_allow_html=True)
            st.caption(f"Ønskestatistikken omfatter {stats.get('wish_students', len(students))} elever med mindst ét angivet lærerønske. {stats.get('no_wishes', 0)} elever har ikke angivet ønsker.")
            export_sheet_options = ["Fordeling", "Lærerbelastning", "Fagstatistik", "Ikke tildelte", "Elevdata", "Lærerdata"]
            selected_export_sheets = st.multiselect("Vælg ark til Excel-filen", export_sheet_options, default=export_sheet_options[:4], key="v2_export_sheets")
            result_export = try_excel_export(lambda: make_export(students, teachers, solution, capacities, solution["use_global_k"], selected_export_sheets)) if selected_export_sheets else None
            if result_export is not None:
                st.download_button("Download fordeling som Excel", data=result_export, file_name="vejlederfordeling.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", type="primary", key="v2_download_result")
            elif not selected_export_sheets:
                st.info("Vælg mindst ét ark for at aktivere Excel-download.")

            result_students, result_teachers, result_subjects, result_missing = st.tabs(["Fordeling", "Lærere", "Fagstatistik", "Ikke tildelte"])
            with result_students:
                rows = []
                for index, student in enumerate(students):
                    assigned = solution["assignments"][index]
                    rows.append({"Elev": student["name"], "Klasse": student.get("className", ""), "Fag 1": student["subjects"][0], "Vejleder 1": teacher_label(assigned[0], teacher_map), "Fag 2": student["subjects"][1], "Vejleder 2": teacher_label(assigned[1], teacher_map), "Ønsker opfyldt": wish_status_label(student, assigned)})
                st.dataframe(pd.DataFrame(rows), width="stretch", hide_index=True, height=560)
            with result_teachers:
                rows = []
                for teacher in sorted(teachers, key=lambda item: (-solution["loads"].get(item["id"], 0), item["id"])):
                    teacher_id = teacher["id"]
                    individual_limit = int(capacities.get(teacher_id, 0))
                    normal_limit = min(solution["K"], individual_limit) if solution["use_global_k"] else individual_limit
                    permit_over = solution["use_global_k"] and solution.get("allow_over_capacity", False) and not solution.get("lock_teacher_max", False)
                    hard_limit = solution["K"] if permit_over else normal_limit
                    load = solution["loads"].get(teacher_id, 0)
                    status = "BRUD" if load > hard_limit else "Tilladt over eget max" if load > normal_limit else "OK"
                    rows.append({"Lærer": teacher_label(teacher_id, teacher_map), "Initialer": teacher_id, "Fag": " · ".join(teacher["subjects"]), "Elever": load, "Eget max": individual_limit, "Hård grænse": hard_limit, "Status": status})
                st.dataframe(pd.DataFrame(rows), width="stretch", hide_index=True, height=560)
            with result_subjects:
                st.dataframe(subject_stats(students, teachers, capacities, solution["K"], solution["use_global_k"], solution.get("allow_over_capacity", False), solution.get("lock_teacher_max", False)), width="stretch", hide_index=True)
            with result_missing:
                missing = [{"Elev": student["name"], "Klasse": student.get("className", ""), "Fag": student["subjects"][slot], "Status": "Mangler vejleder"} for index, student in enumerate(students) for slot in (0, 1) if not solution["assignments"][index][slot]]
                if missing:
                    st.dataframe(pd.DataFrame(missing), width="stretch", hide_index=True)
                else:
                    st.success("Alle elever har to vejledere.")
            with st.expander("Manuel justering af vejledere"):
                st.caption("Vælg en anden lærer i dropdownen. Ved gemning kontrolleres fagmatch, K, lærernes max og I-grænsen. Vælg **Ingen vejleder** for at fjerne en tildeling.")
                assignment_options = [UNASSIGNED_LABEL] + sorted((teacher_label(teacher["id"], teacher_map) for teacher in teachers), key=str.casefold)
                edit_rows = [{"Elev": student["name"], "Fag 1": student["subjects"][0], "Vejleder 1": teacher_label(solution["assignments"][index][0], teacher_map) if solution["assignments"][index][0] else UNASSIGNED_LABEL, "Fag 2": student["subjects"][1], "Vejleder 2": teacher_label(solution["assignments"][index][1], teacher_map) if solution["assignments"][index][1] else UNASSIGNED_LABEL} for index, student in enumerate(students)]
                edited = st.data_editor(
                    pd.DataFrame(edit_rows), width="stretch", hide_index=True, disabled=["Elev", "Fag 1", "Fag 2"],
                    column_config={
                        "Vejleder 1": st.column_config.SelectboxColumn("Vejleder 1", options=assignment_options),
                        "Vejleder 2": st.column_config.SelectboxColumn("Vejleder 2", options=assignment_options),
                    }, key="v2_manual_editor",
                )
                if st.button("Gem manuelle ændringer", key="v2_save_manual"):
                    candidates = build_candidates(teachers)
                    label_to_id = {normal_key(teacher_label(teacher["id"], teacher_map)): teacher["id"] for teacher in teachers}
                    errors = []
                    assignments = []
                    for index, row in edited.iterrows():
                        values = []
                        for slot in (1, 2):
                            selected = repair_text(row[f"Vejleder {slot}"])
                            teacher_id = None if not selected or normal_key(selected) == normal_key(UNASSIGNED_LABEL) else label_to_id.get(normal_key(selected), normal_key(selected))
                            subject = students[index]["subjects"][slot - 1]
                            if teacher_id and teacher_id not in candidates.get(canonical_subject(subject), []):
                                errors.append(f"{students[index]['name']}: {selected} underviser ikke i {subject} eller findes ikke i lærerlisten.")
                                values.append(None)
                            else:
                                values.append(teacher_id)
                        assignments.append(values)
                    manual_assessment = assignment_assessment(
                        students, teachers, capacities, assignments, solution["K"], solution["double_limit"],
                        solution["use_global_k"], solution.get("allow_over_capacity", False), solution.get("lock_teacher_max", False),
                    )
                    errors.extend(manual_assessment["errors"])
                    if errors:
                        st.error(" ".join(errors[:8]))
                    else:
                        solution["assignments"] = assignments
                        solution["loads"] = manual_assessment["loads"]
                        solution["double_loads"] = manual_assessment["double_loads"]
                        solution["stats"] = score_solution(students, teachers, capacities, assignments, solution["loads"], solution["K"], solution["use_global_k"], prioritize_pairs, prioritize_classes)
                        st.session_state["solution"] = solution
                        st.session_state["v2_schedule"] = None
                        st.session_state["v2_schedule_signature"] = None
                        st.session_state["v2_docx_zip"] = None
                        if manual_assessment["warnings"]:
                            st.session_state["v2_flash_warning"] = "Fordelingen er gemt med advarsel: " + " ".join(manual_assessment["warnings"][:5])
                        else:
                            st.session_state["v2_flash_message"] = "Fordelingen er opdateret og regelkontrolleret."
                        st.rerun()
            st.divider()
            if st.button("Forbered Word-filer", key="v2_prepare_docx"):
                with st.spinner("Genererer Word-filer …"):
                    try:
                        st.session_state["v2_docx_zip"] = make_docx_zip(students, teachers, solution)
                    except ImportError:
                        st.error("Word-eksport kræver python-docx.")
            if st.session_state.get("v2_docx_zip"):
                st.download_button("Download SOP-filer (.zip)", data=st.session_state["v2_docx_zip"], file_name="SOP-filer.zip", mime="application/zip", key="v2_download_docx")

    if active_step == "6 · Tidsplan":
        st.subheader("1 · Vælg datagrundlag")
        st.caption("Brug den aktuelle beregnede fordeling, eller upload en færdig, godkendt liste. En uploadet liste har forrang og markeres tydeligt som aktiv kilde.")
        schedule_upload = st.file_uploader(
            "Godkendt liste (.xlsx eller .xlsm)",
            type=["xlsx", "xlsm"],
            key="v2_schedule_upload",
            help="Påkrævede kolonner: Elev, Klasse, Lærer 1, Fag 1, Lærer 2 og Fag 2.",
        )
        template_columns = st.columns([1, 2])
        with template_columns[0]:
            schedule_template = try_excel_export(make_schedule_input_template)
            if schedule_template is not None:
                st.download_button(
                    "Download Excel-skabelon",
                    data=schedule_template,
                    file_name="tidsplan_input.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="v2_download_schedule_template",
                )
        with template_columns[1]:
            st.caption("Den sidste kolonne kan bruges til bemærkninger eller ændringer fra lærerne. Den tomme kolonne i eksemplet er valgfri.")

        uploaded_schedule = None
        if schedule_upload is not None:
            try:
                uploaded_schedule = parse_schedule_upload(schedule_upload)
            except (ValueError, ImportError) as error:
                st.error(str(error))

        solution = st.session_state.get("solution")
        schedule_source_ready = False
        if uploaded_schedule is not None:
            schedule_students_source = uploaded_schedule["students"]
            schedule_teachers_source = uploaded_schedule["teachers"]
            schedule_solution_source = uploaded_schedule["solution"]
            schedule_source_ready = True
            st.success(f"Aktiv kilde: **{uploaded_schedule['source_name']}** · {len(uploaded_schedule['students'])} elever · {len(uploaded_schedule['teachers'])} lærere.")
        elif solution is None:
            st.info("Upload en godkendt liste her, eller beregn først en fordeling i trin 4.")
        else:
            schedule_students_source = students
            schedule_teachers_source = teachers
            schedule_solution_source = solution
            missing_schedule_assignments = sum(any(not teacher_id for teacher_id in assigned) for assigned in solution.get("assignments", []))
            schedule_source_ready = missing_schedule_assignments == 0
            st.success(f"Aktiv kilde: **Aktuel beregnet fordeling** · {len(students)} elever · {len(teachers)} lærere.")
            st.caption("Hver elev får én samlet vejledningstid med sine tildelte vejledere. Brug Elevplanen som elevens opslag og Lærerplanen som lærerens dagsorden.")
            if missing_schedule_assignments:
                st.error(f"{missing_schedule_assignments} elev(er) mangler en komplet vejlederfordeling. Ret dem i trin 5, før tidsplanen genereres.")

        if uploaded_schedule is not None or solution is not None:
            # Opdatér eksisterende sessioner én gang til de aftalte
            # standardindstillinger; derefter er alle felter fortsat frie at
            # justere som normalt.
            if st.session_state.get("v2_schedule_defaults_version") != "2026-09-15":
                st.session_state.update({
                    "v2_schedule_start": dt_time(8, 15),
                    "v2_schedule_end": dt_time(16, 15),
                    "v2_schedule_transition_minutes": 0,
                    "v2_schedule_pause_count": 2,
                    "v2_schedule_pause_minutes": 10,
                    "v2_schedule_floating_pauses": True,
                    "v2_schedule_group_pairs": True,
                    "v2_schedule_avoid_teacher_gaps": True,
                    "v2_schedule_defaults_version": "2026-09-15",
                })
            st.subheader("2 · Indstil dagen")
            time_columns = st.columns(4)
            with time_columns[0]:
                schedule_start = st.time_input("Starttidspunkt", value=dt_time(8, 15), key="v2_schedule_start")
            with time_columns[1]:
                schedule_end = st.time_input("Sluttidspunkt", value=dt_time(16, 15), key="v2_schedule_end")
            with time_columns[2]:
                schedule_student_minutes = st.number_input("Minutter pr. elev", min_value=1, max_value=180, value=20, step=5, key="v2_schedule_student_minutes")
            with time_columns[3]:
                schedule_transition_minutes = st.number_input("Minutter mellem elever", min_value=0, max_value=60, value=0, step=1, key="v2_schedule_transition_minutes")
            pause_columns = st.columns(3)
            with pause_columns[0]:
                schedule_pause_count = st.number_input("Antal pauser", min_value=0, max_value=20, value=2, step=1, key="v2_schedule_pause_count")
            with pause_columns[1]:
                schedule_pause_minutes = st.number_input("Minutter pr. pause", min_value=0, max_value=120, value=10, step=5, key="v2_schedule_pause_minutes")
            with pause_columns[2]:
                schedule_group_pairs = st.checkbox("Saml samme lærerpar mest muligt", value=True, key="v2_schedule_group_pairs", help="Elever med samme lærerpar lægges i sammenhængende blokke, så lærerne skifter færre gange.")
            schedule_floating_pauses = st.checkbox(
                "Fordel pauser før og efter frokost",
                value=True,
                key="v2_schedule_floating_pauses",
                help="Med to pauser placeres én efter en vejledning før frokost og én efter frokost, men før dagens sidste vejledning.",
            )
            schedule_avoid_teacher_gaps = st.checkbox(
                "Forsøg at undgå huller i lærernes vejledningstider",
                value=True,
                key="v2_schedule_avoid_teacher_gaps",
                help="Bevarer det lavest mulige antal vejledningsrunder, men prøver at samle hver lærers runder, så der er færre tomme mellemrum.",
            )
            st.markdown("#### Lærernes spærringer")
            st.caption("Vælg de lærere, der har en spærring, og vælg derefter fra- og til-klokkeslet. Tiderne kan kun vælges fra listen og ligger inden for den valgte dag.")
            teacher_block_seed = stable_signature([teacher["id"] for teacher in schedule_teachers_source])
            if st.session_state.get("v2_schedule_teacher_block_seed") != teacher_block_seed:
                st.session_state["v2_schedule_teacher_block_seed"] = teacher_block_seed
                st.session_state.pop("v2_schedule_blocked_teachers", None)
            teacher_map_for_blocks = {teacher["id"]: teacher for teacher in schedule_teachers_source}
            teacher_ids_for_blocks = sorted(teacher_map_for_blocks, key=lambda teacher_id: teacher_label(teacher_id, teacher_map_for_blocks).casefold())
            blocked_teacher_ids = st.multiselect(
                "Lærere med spærring",
                teacher_ids_for_blocks,
                format_func=lambda teacher_id: teacher_label(teacher_id, teacher_map_for_blocks),
                placeholder="Vælg eventuelt en eller flere lærere",
                key="v2_schedule_blocked_teachers",
            )
            teacher_blocks: dict[str, list[tuple[dt_time, dt_time]]] = {}
            teacher_blocks_error = ""
            time_choices = schedule_time_choices(schedule_start, schedule_end)
            if not time_choices:
                teacher_blocks_error = "Sluttidspunktet skal ligge efter starttidspunktet."
                st.error(teacher_blocks_error)
            elif blocked_teacher_ids:
                for teacher_id in blocked_teacher_ids:
                    label = teacher_label(teacher_id, teacher_map_for_blocks)
                    with st.container(border=True):
                        st.markdown(f"**{label}**")
                        block_columns = st.columns(2)
                        with block_columns[0]:
                            blocked_start = st.selectbox(
                                "Spærret fra",
                                time_choices[:-1],
                                format_func=lambda value: value.strftime("%H:%M"),
                                key=f"v2_schedule_block_start_{teacher_id}",
                            )
                        valid_end_choices = [value for value in time_choices if value > blocked_start]
                        end_key = f"v2_schedule_block_end_{teacher_id}"
                        if end_key in st.session_state and st.session_state[end_key] not in valid_end_choices:
                            st.session_state[end_key] = valid_end_choices[0]
                        with block_columns[1]:
                            blocked_end = st.selectbox(
                                "Spærret til",
                                valid_end_choices,
                                format_func=lambda value: value.strftime("%H:%M"),
                                key=end_key,
                            )
                    teacher_blocks[teacher_id] = [(blocked_start, blocked_end)]
                st.info(f"{len(teacher_blocks)} lærerspærring(er) er aktiv(e) og vil påvirke tidsplanen.")
            lunch_columns = st.columns(3)
            with lunch_columns[0]:
                schedule_lunch_minutes = st.number_input(
                    "Frokostpause (minutter)", min_value=1, max_value=120, value=30, step=5,
                    key="v2_schedule_lunch_minutes",
                )
            with lunch_columns[1]:
                schedule_lunch_mode = st.selectbox(
                    "Frokostpause for lærerne",
                    ["Fast tidspunkt for alle lærere", "Flydende for alle lærere"],
                    key="v2_schedule_lunch_mode",
                    help="Fast betyder samme klokkeslæt for alle. Flydende placerer pausen midt i den samlede vejledningsplan.",
                )
            with lunch_columns[2]:
                if schedule_lunch_mode == "Fast tidspunkt for alle lærere":
                    schedule_lunch_start = st.time_input("Fast frokosttid", value=dt_time(12, 0), key="v2_schedule_lunch_start")
                else:
                    schedule_lunch_start = dt_time(12, 0)
                    st.caption("Flydende: placeres automatisk midt i planen.")
            schedule_depth = st.slider(
                "Algoritmedybde for tidsplan",
                10,
                300,
                80,
                10,
                key="v2_schedule_depth",
                help="Algoritmen afprøver forskellige måder at placere uafhængige lærerpar parallelt. En højere værdi kan give færre vejledningsrunder, men tager længere tid.",
            )
            current_schedule_signature = schedule_signature(
                schedule_students_source,
                schedule_teachers_source,
                schedule_solution_source,
                [schedule_start, schedule_end, int(schedule_student_minutes), int(schedule_pause_count),
                 int(schedule_pause_minutes), int(schedule_transition_minutes), bool(schedule_group_pairs),
                 bool(schedule_floating_pauses), bool(schedule_avoid_teacher_gaps), schedule_lunch_mode, schedule_lunch_start,
                 int(schedule_lunch_minutes), int(schedule_depth), teacher_blocks],
            )
            if st.session_state.get("v2_schedule") is not None and st.session_state.get("v2_schedule_signature") != current_schedule_signature:
                st.session_state["v2_schedule"] = None
                st.session_state["v2_schedule_signature"] = None
                st.info("Tidsindstillingerne eller datakilden er ændret. Generér planen igen for at se et aktuelt resultat.")
            st.subheader("3 · Generér og kontrollér")
            if st.button("Generér tidsplan" if uploaded_schedule is not None else "Lav eller opdater tidsplan", type="primary", key="v2_make_schedule", disabled=not schedule_source_ready or bool(teacher_blocks_error)):
                try:
                    progress_bar = st.progress(0.0, text="Beregner tidsplan: 0 %")
                    st.session_state["v2_schedule"] = make_schedule(
                        schedule_students_source,
                        schedule_teachers_source,
                        schedule_solution_source,
                        schedule_start,
                        schedule_end,
                        int(schedule_student_minutes),
                        int(schedule_pause_count),
                        int(schedule_pause_minutes),
                        int(schedule_transition_minutes),
                        schedule_group_pairs,
                        lunch_mode=schedule_lunch_mode,
                        lunch_start_time=schedule_lunch_start,
                        lunch_minutes=int(schedule_lunch_minutes),
                        search_attempts=int(schedule_depth),
                        avoid_teacher_gaps=schedule_avoid_teacher_gaps,
                        teacher_blocks=teacher_blocks,
                        floating_pauses=schedule_floating_pauses,
                        progress_callback=lambda value: update_algorithm_progress(progress_bar, "Beregner tidsplan", value),
                    )
                    st.session_state["v2_schedule_signature"] = current_schedule_signature
                    update_algorithm_progress(progress_bar, "Tidsplan færdig", 1.0)
                    st.success("Tidsplanen er genereret fra Excel-filen." if uploaded_schedule is not None else "Tidsplanen er opdateret.")
                except ValueError as error:
                    st.session_state["v2_schedule"] = None
                    st.session_state["v2_schedule_signature"] = None
                    st.error(str(error))
            schedule = st.session_state.get("v2_schedule")
            if schedule is not None:
                settings = schedule["settings"]
                metric_columns = st.columns(4)
                metric_columns[0].metric("Elever", len(schedule["students"]))
                metric_columns[1].metric("Lærerpar", len(schedule["pairs"]))
                metric_columns[2].metric("Planlagt tidsforbrug", f"{settings['Planlagt tidsforbrug (min.)']} min.")
                metric_columns[3].metric("Sidste sluttid", settings["Slut"])
                st.info("Elever finder deres tidspunkt i Elevplanen. Lærere kan vælge deres navn i Lærerplanen og se elev, fag og eventuel medvejleder.")
                schedule_students, schedule_teachers, schedule_timeline, schedule_pairs = st.tabs(["Elevplan", "Lærerplan", "Samlet tidslinje", "Lærerpar"])
                with schedule_students:
                    st.dataframe(schedule["students"], width="stretch", hide_index=True, height=560)
                with schedule_teachers:
                    teacher_query = st.text_input("Søg på lærerens navn", placeholder="Skriv navn eller initialer", key="v2_schedule_teacher_search")
                    all_teacher_names = sorted(schedule["teachers"]["Lærer"].dropna().unique().tolist(), key=str.casefold)
                    query_key = normal_key(teacher_query)
                    matching_teacher_names = [
                        name for name in all_teacher_names
                        if not query_key or query_key in normal_key(name)
                    ]
                    teacher_frame = schedule["teachers"][schedule["teachers"]["Lærer"].isin(matching_teacher_names)]
                    st.dataframe(teacher_frame, width="stretch", hide_index=True, height=560)
                    if settings.get("Forsøg at undgå lærerhuller") == "Ja":
                        st.caption(f"Huloptimering: {settings.get('Lærerhuller (runder)', 0)} tomme runder mellem lærernes første og sidste vejledning.")
                    st.divider()
                    visual_schedule = {**schedule, "teachers": teacher_frame}
                    st.markdown(make_teacher_gantt_html(visual_schedule, None), unsafe_allow_html=True)
                with schedule_timeline:
                    st.dataframe(schedule["timeline"], width="stretch", hide_index=True, height=560)
                with schedule_pairs:
                    st.dataframe(schedule["pairs"], width="stretch", hide_index=True)
                download_columns = st.columns(3)
                with download_columns[0]:
                    st.download_button(
                        "Download tidsplan som HTML",
                        data=make_schedule_html(schedule).encode("utf-8"),
                        file_name="Vejledningsplan (SOPtima beta af Henrik Sterner).html",
                        mime="text/html",
                        type="primary",
                        key="v2_download_schedule_html",
                    )
                with download_columns[1]:
                    st.download_button(
                        "Download elevopslag som HTML",
                        data=make_student_schedule_html(schedule).encode("utf-8"),
                        file_name="elevopslag-vejledning.html",
                        mime="text/html",
                        type="primary",
                        key="v2_download_student_schedule_html",
                    )
                with download_columns[2]:
                    schedule_export = try_excel_export(lambda: make_schedule_excel(schedule))
                    if schedule_export is not None:
                        st.download_button(
                            "Download tidsplan som Excel",
                            data=schedule_export,
                            file_name="vejledningsplan.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            type="primary",
                            key="v2_download_schedule_excel",
                        )
                    else:
                        st.error("Excel-download kræver openpyxl. Installer projektets requirements.txt for at aktivere eksporten.")

    st.divider()
    active_index = process_steps.index(active_step)
    navigation_columns = st.columns([1, 2, 1])
    with navigation_columns[0]:
        if active_index > 0:
            st.button(
                f"← {step_labels[process_steps[active_index - 1]].strip()}",
                key="v2_previous_step",
                use_container_width=True,
                on_click=navigate_to_step,
                args=(process_steps[active_index - 1],),
            )
    with navigation_columns[1]:
        st.caption(f"Du er i trin {active_index + 1} af 6. Du kan altid gå direkte til et trin i venstremenuen.")
    with navigation_columns[2]:
        if active_index < len(process_steps) - 1:
            st.button(
                f"{step_labels[process_steps[active_index + 1]].strip()} →",
                key="v2_next_step",
                use_container_width=True,
                on_click=navigate_to_step,
                args=(process_steps[active_index + 1],),
            )
    st.caption(COPYRIGHT)


if __name__ == "__main__":
    main_v2()

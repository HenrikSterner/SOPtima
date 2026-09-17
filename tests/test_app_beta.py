from __future__ import annotations

import io
import base64
import html
import re
import shutil
import subprocess
import zipfile
from datetime import date, datetime, time as clock

import pandas as pd
import pytest
from docx import Document
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

import app
import app_alpha as alpha
import app_beta as beta


def student(index: int, *, name: str | None = None, email: str = "") -> dict:
    return {
        "id": f"E{index:03d}",
        "name": name or f"Elev {index}",
        "className": "S 2024q",
        "subjects": ["Dansk A", "Historie B"],
        "subjectsWithLevel": ["Dansk A", "Historie B"],
        "wishes": [],
        "projectTitle": f"Projekt {index}",
        "projectDescription": "",
        "scheduleNotes": "",
        "teams_email": email,
    }


def teacher(teacher_id: str, name: str, email: str = "") -> dict:
    return {
        "id": teacher_id,
        "name": name,
        "subjects": ["Dansk", "Historie"],
        "holds": [],
        "teams_email": email,
    }


def make_beta_schedule(
    pairs: list[tuple[str, str]],
    *,
    start: clock = clock(8, 0),
    end: clock = clock(12, 0),
    minutes: int = 20,
    transition: int = 0,
    pause_count: int = 0,
    pause_minutes: int = 0,
    teacher_blocks=None,
    locks=None,
    students_override=None,
    teachers_override=None,
):
    ids = sorted({teacher_id for pair in pairs for teacher_id in pair})
    teachers = teachers_override or [teacher(teacher_id, teacher_id.upper(), f"{teacher_id}@nextkbh.dk") for teacher_id in ids]
    students = students_override or [student(index + 1, email=f"next{27000 + index}@edu.nextkbh.dk") for index in range(len(pairs))]
    solution = {"assignments": [list(pair) for pair in pairs]}
    schedule = beta.make_schedule_beta(
        students,
        teachers,
        solution,
        start,
        end,
        minutes,
        pause_count,
        pause_minutes,
        transition,
        True,
        lunch_mode="Fast tidspunkt for alle lærere",
        lunch_start_time=clock(10, 0) if end > clock(10, 5) else clock(8, 40),
        lunch_minutes=5,
        search_attempts=5,
        teacher_blocks=teacher_blocks,
        floating_pauses=True,
        locked_activities=locks or {},
    )
    context = {
        "students": students,
        "teachers": teachers,
        "solution": solution,
        "start_time": start,
        "end_time": end,
        "student_minutes": minutes,
        "pause_count": pause_count,
        "pause_minutes": pause_minutes,
        "transition_minutes": transition,
        "group_pairs": True,
        "lunch_mode": "Fast tidspunkt for alle lærere",
        "lunch_start_time": clock(10, 0) if end > clock(10, 5) else clock(8, 40),
        "lunch_minutes": 5,
        "search_attempts": 5,
        "avoid_teacher_gaps": True,
        "teacher_blocks": teacher_blocks,
        "floating_pauses": True,
    }
    return schedule, context


def minutes(value: str) -> int:
    parsed = datetime.strptime(value, "%H:%M")
    return parsed.hour * 60 + parsed.minute


def test_alpha_preserves_core_pair_blocks_when_rebuilding_schedule():
    pairs = [("a", "b"), ("c", "d"), ("a", "b"), ("a", "e"), ("c", "d"), ("f", "g"), ("f", "g")]
    ids = sorted({teacher_id for pair in pairs for teacher_id in pair})
    teachers = [teacher(teacher_id, teacher_id.upper()) for teacher_id in ids]
    students = [student(index + 1) for index in range(len(pairs))]
    schedule = alpha.make_schedule_beta(
        students, teachers, {"assignments": [list(pair) for pair in pairs]},
        clock(8), clock(12), 20, 0, 0, 0, True,
        lunch_mode="Fast tidspunkt for alle lærere", lunch_start_time=clock(10),
        lunch_minutes=5, search_attempts=40, avoid_teacher_gaps=True,
    )

    frame = schedule["students"]
    round_by_start = {start: index for index, start in enumerate(sorted(frame["Start"].unique()))}
    for _, rows in frame.groupby("Lærerpar"):
        positions = sorted(round_by_start[start] for start in rows["Start"])
        assert not positions or positions == list(range(positions[0], positions[-1] + 1))


def test_alpha_teams_link_uses_student_email_column_and_teacher_initials():
    pupil = student(1, email="")
    pupil["Email"] = "NEXT27888@EDU.NEXTKBH.DK"
    pupil["teams_email"] = "ikke-en-email"
    teachers = [teacher("hst", "Henrik", ""), teacher("matr", "Mathias", "")]
    teachers[0]["teams_email"] = "ugyldig"
    schedule = alpha.make_schedule_beta(
        [pupil], teachers, {"assignments": [["hst", "matr"]]},
        clock(8), clock(12), 20, 0, 0, 0, True,
        lunch_mode="Fast tidspunkt for alle lærere", lunch_start_time=clock(10),
        lunch_minutes=5, search_attempts=5,
    )

    activity_id = schedule["activities"].iloc[0]["activity_id"]
    link, error = alpha.teams_chat_link(schedule, activity_id)
    assert not error
    assert "hst%40nextkbh.dk" in link
    assert "matr%40nextkbh.dk" in link
    assert "next27888%40edu.nextkbh.dk" in link
    assert "users=hst%40nextkbh.dk,matr%40nextkbh.dk,next27888%40edu.nextkbh.dk" in link
    teacher_html = alpha.make_schedule_html_beta(schedule)
    assert f'href="{link.replace("&", "&amp;")}"' in teacher_html
    assert teacher_html.count('class="action action-teams"') == 2


def dynamic_word_bytes(teacher_html: str) -> bytes:
    """Kør den indlejrede browser-generator i Node.js til eksporttest."""
    node_script = r"""
const input=await new Promise(resolve=>{let value='';process.stdin.setEncoding('utf8');process.stdin.on('data',part=>value+=part);process.stdin.on('end',()=>resolve(value));});
const scripts=[...input.matchAll(/<script(?:\s[^>]*)?>([\s\S]*?)<\/script>/g)].map(match=>match[1]);
const values=input.match(/<script id="dynamic-word-values" type="application\/json">([\s\S]*?)<\/script>/)?.[1];
const template=input.match(/<script id="dynamic-word-template" type="text\/plain">([\s\S]*?)<\/script>/)?.[1];
const generator=scripts.find(script=>script.includes('function wordTemplateFiles'));
if(!values||!template||!generator)throw new Error('Mangler indlejret Word-generator');
global.document={getElementById:id=>({textContent:id==='dynamic-word-values'?values:template}),querySelectorAll:()=>[]};
const item=Object.values(JSON.parse(values))[0];
const bytes=await eval(`(async()=>{${generator};return await createWordDocument(${JSON.stringify(item)});})()`);
process.stdout.write(Buffer.from(bytes).toString('base64'));
"""
    result = subprocess.run(
        ["node", "--input-type=module", "-e", node_script], input=teacher_html,
        text=True, capture_output=True, check=False,
    )
    assert result.returncode == 0, result.stderr
    return base64.b64decode(result.stdout)


def assert_no_teacher_conflicts(schedule: dict) -> None:
    frame = schedule["teachers"]
    frame = frame[frame["Type"] == "Vejledning"]
    for _, group in frame.groupby("Initialer"):
        rows = group.sort_values("Start").to_dict("records")
        for left, right in zip(rows, rows[1:]):
            assert minutes(left["Slut"]) <= minutes(right["Start"])


def test_importing_beta_does_not_patch_production_module():
    assert app.make_schedule is beta._CORE_MAKE_SCHEDULE
    assert app.make_schedule_html is beta._CORE_MAKE_SCHEDULE_HTML


def test_partial_schedule_returns_planned_and_unplaced_instead_of_error():
    schedule, _ = make_beta_schedule([("a", "b")] * 4, end=clock(9, 0))
    assert 0 < len(schedule["students"]) < 4
    assert len(schedule["students"]) + len(schedule["unplaced"]) == 4
    assert set(schedule["unplaced"]["Status"]) == {"Uplaceret"}
    assert schedule["settings"]["Slut"] == "09:00"
    assert_no_teacher_conflicts(schedule)


def test_beta_deduplicates_columns_from_a_previous_streamlit_rerun(monkeypatch):
    students = [student(1)]
    teachers = [teacher("a", "A"), teacher("b", "B")]
    solution = {"assignments": [["a", "b"]]}
    original_make_schedule = beta._CORE_MAKE_SCHEDULE
    def compatible_base(*args, **kwargs):
        kwargs["lunch_mode"] = "Fast tidspunkt for alle l" + chr(230) + "rere"
        return original_make_schedule(*args, **kwargs)
    monkeypatch.setattr(beta, "_CORE_MAKE_SCHEDULE", compatible_base)
    original_beta_schedule = beta.make_schedule_beta
    def compatible_beta(*args, **kwargs):
        kwargs["lunch_mode"] = "Fast tidspunkt for alle l" + chr(230) + "rere"
        return original_beta_schedule(*args, **kwargs)
    monkeypatch.setattr(beta, "make_schedule_beta", compatible_beta)
    base = beta._CORE_MAKE_SCHEDULE(
        students, teachers, solution, clock(8), clock(12), 20, 0, 0, 0, True,
        lunch_mode="Fast tidspunkt for alle lÃ¦rere", lunch_start_time=clock(10),
        lunch_minutes=5, search_attempts=2,
    )
    # Efter en script-genkørsel kunne en gammel beta-plan tidligere blive brugt
    # som grundplan og få de samme kolonner tilføjet igen.
    base["students"] = pd.concat([base["students"], base["students"]], axis=1)
    base["teachers"] = pd.concat([base["teachers"], base["teachers"]], axis=1)
    base["timeline"] = pd.concat([base["timeline"], base["timeline"]], axis=1)
    monkeypatch.setattr(beta, "_CORE_MAKE_SCHEDULE", lambda *args, **kwargs: base)
    schedule = beta.make_schedule_beta(
        students, teachers, solution, clock(8), clock(12), 20, 0, 0, 0, True,
        lunch_mode="Fast tidspunkt for alle lÃ¦rere", lunch_start_time=clock(10),
        lunch_minutes=5, search_attempts=2,
    )
    assert schedule["students"].columns.is_unique
    assert schedule["teachers"].columns.is_unique
    assert schedule["timeline"].columns.is_unique


def test_activity_schema_and_ids_are_stable_when_input_order_changes():
    students = [student(1), student(2), student(3)]
    teachers = [teacher("a", "A"), teacher("b", "B")]
    first, _ = make_beta_schedule([("a", "b")] * 3, students_override=students, teachers_override=teachers)
    second_students = [students[2], students[0], students[1]]
    second, _ = make_beta_schedule([("a", "b")] * 3, students_override=second_students, teachers_override=teachers)
    required = {"activity_id", "elev_id", "lærer_ids", "varighed", "start", "slut", "status", "låst", "årsag"}
    assert required <= set(first["activities"].columns)
    mapping_first = dict(zip(first["activities"]["elev_id"], first["activities"]["activity_id"]))
    mapping_second = dict(zip(second["activities"]["elev_id"], second["activities"]["activity_id"]))
    assert mapping_first == mapping_second


def test_parallel_rounds_transition_and_teacher_blocks_are_hard_constraints():
    schedule, _ = make_beta_schedule(
        [("a", "b"), ("c", "d"), ("a", "b")],
        transition=5,
        teacher_blocks={"a": [(clock(8, 0), clock(8, 30))]},
    )
    a_rows = schedule["teachers"]
    a_rows = a_rows[(a_rows["Type"] == "Vejledning") & (a_rows["Initialer"] == "a")].sort_values("Start")
    assert all(minutes(value) >= 8 * 60 + 30 for value in a_rows["Start"])
    rows = a_rows.to_dict("records")
    for left, right in zip(rows, rows[1:]):
        assert minutes(right["Start"]) - minutes(left["Slut"]) >= 5
    assert_no_teacher_conflicts(schedule)


def test_independent_teacher_pairs_can_run_in_parallel():
    schedule, _ = make_beta_schedule([("a", "b"), ("c", "d")])
    assert schedule["students"]["Start"].nunique() == 1


def test_morning_teacher_block_keeps_unaffected_pairs_as_early_as_possible():
    schedule, _ = make_beta_schedule(
        [("a", "b")] * 4 + [("c", "d")] * 4,
        end=clock(12),
        teacher_blocks={"a": [(clock(8), clock(10))]},
    )
    rows = schedule["students"]
    unaffected = rows[rows["Vejleder 1"].str.contains("C", case=False, regex=False)]
    affected = rows[rows["Vejleder 1"].str.contains("A", case=False, regex=False)]
    assert unaffected["Start"].map(minutes).max() < 10 * 60
    assert affected["Start"].map(minutes).min() >= 10 * 60


def test_move_assessment_and_dynamic_repair():
    schedule, context = make_beta_schedule([("a", "b")] * 3)
    ids = schedule["activities"]["activity_id"].tolist()
    first_start = schedule["activities"].set_index("activity_id").loc[ids[0], "start"]
    assessment = beta.assess_move(schedule, ids[1], first_start)
    assert assessment["status"] == "yellow"
    with pytest.raises(ValueError, match="Streng tilstand"):
        beta.apply_manual_move(schedule, context, {}, ids[1], first_start, "strict")
    updated, locks, result = beta.apply_manual_move(schedule, context, {}, ids[1], first_start, "dynamic")
    assert result["status"] == "yellow"
    assert locks[ids[1]] == first_start
    moved = updated["activities"].set_index("activity_id")
    assert moved.loc[ids[1], "start"] == first_start
    assert moved.loc[ids[1], "status"] == "låst"
    assert_no_teacher_conflicts(updated)


def test_red_move_is_rejected_for_break_and_five_minute_grid():
    schedule, context = make_beta_schedule([("a", "b")] * 3)
    activity_id = schedule["activities"].iloc[0]["activity_id"]
    lunch_start = schedule["beta"]["breaks"][0][0]
    assert beta.assess_move(schedule, activity_id, lunch_start)["status"] == "red"
    assert beta.assess_move(schedule, activity_id, "08:03")["status"] == "red"
    with pytest.raises(ValueError):
        beta.apply_manual_move(schedule, context, {}, activity_id, lunch_start, "dynamic")


def test_move_validity_map_contains_all_colours_needed_by_component():
    schedule, _ = make_beta_schedule([("a", "b")] * 3)
    activity_id = schedule["activities"].iloc[1]["activity_id"]
    validity = beta.move_validity_map(schedule)[activity_id]
    assert "yellow" in set(validity.values())
    assert "red" in set(validity.values())
    assert "green" in set(validity.values())


def test_pauses_are_only_exported_between_guidance_sessions():
    schedule, _ = make_beta_schedule(
        [("a", "b")] * 6,
        end=clock(12, 30),
        pause_count=2,
        pause_minutes=10,
    )
    guidance = schedule["students"]
    pause_rows = schedule["timeline"][schedule["timeline"]["Type"] == "Pause"]
    for _, pause in pause_rows.iterrows():
        assert any(guidance["Slut"].map(minutes) <= minutes(pause["Start"]))
        assert any(guidance["Start"].map(minutes) >= minutes(pause["Slut"]))


def test_extreme_overflow_uses_fallback_and_never_discards_all_results():
    pairs = [("a", "b")] * 55
    schedule, _ = make_beta_schedule(pairs, end=clock(9, 0))
    assert len(schedule["students"]) + len(schedule["unplaced"]) == 55
    assert len(schedule["students"]) > 0
    assert len(schedule["unplaced"]) > 0


def test_teams_link_contains_student_and_both_teacher_addresses():
    students = [student(1, email="next27888@edu.nextkbh.dk")]
    teachers = [
        teacher("hst", "Henrik Sterner", "hst@nextkbh.dk"),
        teacher("matr", "Mathias Ramberg", "matr@nextkbh.dk"),
    ]
    schedule, _ = make_beta_schedule(
        [("hst", "matr")], students_override=students, teachers_override=teachers
    )
    activity_id = schedule["activities"].iloc[0]["activity_id"]
    link, error = beta.teams_chat_link(schedule, activity_id)
    assert not error
    assert "hst%40nextkbh.dk" in link
    assert "matr%40nextkbh.dk" in link
    assert "next27888%40edu.nextkbh.dk" in link
    assert "topicName=SOP%20-%20Elev%201%20-%20S%202024q" in link


def test_missing_teams_email_disables_action_with_explanation():
    students = [student(1, email="")]
    teachers = [teacher("a1", "A", "a1@nextkbh.dk"), teacher("longteacher", "B", "")]
    schedule, _ = make_beta_schedule([("a1", "longteacher")], students_override=students, teachers_override=teachers)
    activity_id = schedule["activities"].iloc[0]["activity_id"]
    link, error = beta.teams_chat_link(schedule, activity_id)
    assert link == ""
    assert "mangler Teams-adresse" in error
    teacher_html = beta.make_schedule_html_beta(schedule)
    assert "action-disabled" in teacher_html
    assert "mangler Teams-adresse" in teacher_html


def test_standard_teacher_initials_get_a_teams_address_and_render_a_real_link():
    students = [student(1, email="next27888@edu.nextkbh.dk")]
    teachers = [teacher("bok", "Bo", ""), teacher("polk", "Poul", "")]
    schedule, _ = make_beta_schedule([("bok", "polk")], students_override=students, teachers_override=teachers)
    activity_id = schedule["activities"].iloc[0]["activity_id"]
    link, error = beta.teams_chat_link(schedule, activity_id)
    assert not error
    assert "bok%40nextkbh.dk" in link and "polk%40nextkbh.dk" in link
    teacher_html = beta.make_schedule_html_beta(schedule)
    assert f'href="{link.replace("&", "&amp;")}"' in teacher_html


def test_teacher_html_has_word_and_teams_actions_but_unplaced_has_none():
    schedule, _ = make_beta_schedule([("a", "b")] * 4, end=clock(9, 0))
    teacher_html = beta.make_schedule_html_beta(schedule)
    assert "Download Word" in teacher_html
    assert "Åbn Teams-chat" in teacher_html
    assert 'id="dynamic-word-values"' in teacher_html
    assert "data-dynamic-word-id=" in teacher_html
    assert "createWordDocument" in teacher_html
    assert 'id="dynamic-word-template"' in teacher_html
    assert "word/document.xml" in teacher_html
    assert 'id="word-documents"' not in teacher_html
    assert "Kan ikke placeres inden for det valgte tidsrum" in teacher_html
    assert teacher_html.count("Download Word") == len(schedule["teachers"][schedule["teachers"]["Type"] == "Vejledning"])


def test_teacher_export_package_contains_only_html_with_dynamic_unique_word_downloads():
    students = [
        student(1, name="Samme/Navn", email="one@edu.nextkbh.dk"),
        student(2, name="Samme/Navn", email="two@edu.nextkbh.dk"),
    ]
    teachers = [teacher("a", "Anna", "a@nextkbh.dk"), teacher("b", "Bo", "b@nextkbh.dk")]
    schedule, _ = make_beta_schedule(
        [("a", "b"), ("a", "b")], students_override=students, teachers_override=teachers
    )
    package = beta.make_teacher_export_package(schedule)
    with zipfile.ZipFile(io.BytesIO(package)) as archive:
        names = archive.namelist()
        assert names == ["Vejledningsplan.html"]
        package_html = archive.read("Vejledningsplan.html").decode("utf-8")
        assert "createWordDocument" in package_html
        assert "data-dynamic-word-id=" in package_html
        assert "SOP-Samme-Navn-S 2024q-A-E001.docx" in package_html
        assert "SOP-Samme-Navn-S 2024q-A-E002.docx" in package_html
        assert "Word/" not in package_html
        embedded = re.search(r'<script id="dynamic-word-template" type="text/plain">([^<]+)</script>', package_html)
        assert embedded
        assert base64.b64decode(embedded.group(1)) == app.DEFAULT_TEMPLATE.read_bytes()


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js kræves for at afprøve den indlejrede Word-generator")
def test_dynamic_word_download_uses_template_and_replaces_all_schedule_fields():
    schedule, _ = make_beta_schedule([("a", "b")])
    teacher_html = beta.make_schedule_html_beta(schedule)
    word_bytes = dynamic_word_bytes(teacher_html)
    with zipfile.ZipFile(io.BytesIO(word_bytes)) as document:
        assert "word/media/image1.png" in document.namelist()
        xml = document.read("word/document.xml").decode("utf-8")
    visible_text = html.unescape(re.sub(r"<[^>]+>", "", xml))
    Document(io.BytesIO(word_bytes))
    assert "Elev: Elev 1" in visible_text
    assert "Klasse: 3.q" in visible_text
    assert "<Elev 1>" not in visible_text and "<3.q>" not in visible_text
    assert "Dansk A" in xml and "Historie B" in xml
    first = schedule["students"].iloc[0]
    assert str(first["Vejleder 1"]) in xml
    assert str(first["Vejleder 2"]) in xml
    for placeholder in ("Elevnavn", "Fag1 og niveau", "Vejleder fag 1", "Fag2 og niveau", "Vejleder fag 2"):
        assert placeholder not in xml


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js kræves for at afprøve den indlejrede Word-generator")
def test_sop_template_fields_follow_the_documented_replacement_rules(tmp_path, monkeypatch):
    fields = [
        "Elevnavn", "Klasse", "Fag1 og niveau", "Vejleder fag 1", "Fag2 og niveau", "Vejleder fag 2",
    ]
    template = Document()
    template.add_paragraph("Elev: <Elevnavn> | Klasse: <Klasse> | " + " | ".join(f"<{field}>" for field in fields[2:]))
    template_path = tmp_path / "SOP-SKABELON.docx"
    template.save(template_path)
    values = {
        "Elevnavn": "Ane Andersen", "Klasse": "3.q", "Fag1 og niveau": "Dansk A",
        "Vejleder fag 1": "Anna (ann)", "Fag2 og niveau": "Historie B", "Vejleder fag 2": "Bo (bo)",
    }
    completed = Document(io.BytesIO(app.replace_docx_placeholders(template_path.read_bytes(), values)))
    completed_text = "\n".join(paragraph.text for paragraph in completed.paragraphs)
    for field, value in values.items():
        assert value in completed_text
        assert f"<{field}>" not in completed_text
        assert f"<{value}>" not in completed_text
    assert "Klasse: 3.q" in completed_text

    monkeypatch.setattr(app, "DEFAULT_TEMPLATE", template_path)
    schedule, _ = make_beta_schedule([("a", "b")])
    word_bytes = dynamic_word_bytes(beta.make_schedule_html_beta(schedule))
    with zipfile.ZipFile(io.BytesIO(word_bytes)) as document:
        xml = document.read("word/document.xml").decode("utf-8")
    visible_text = html.unescape(re.sub(r"<[^>]+>", "", xml))
    for field in fields:
        assert f"<{field}>" not in visible_text
    assert "Elev: Elev 1" in visible_text and "Klasse: 3.q" in visible_text
    assert "<Elev 1>" not in visible_text and "<3.q>" not in visible_text
    assert "Dansk A" in xml and "Historie B" in xml


def test_sop_class_label_uses_the_pupil_facing_class_format():
    assert app.sop_class_label("S 2024y", date(2026, 9, 17)) == "3.y"
    assert app.sop_class_label("3.q", date(2026, 9, 17)) == "3.q"


def test_batch_word_utility_reports_progress_for_each_planned_student():
    schedule, _ = make_beta_schedule([("a", "b"), ("a", "b")])
    progress = []
    documents = beta._word_documents(schedule, progress_callback=lambda completed, total: progress.append((completed, total)))
    assert len(documents) == 2
    assert progress == [(1, 2), (2, 2)]


def test_teacher_package_never_prebuilds_word_documents(monkeypatch):
    schedule, _ = make_beta_schedule([("a", "b")])

    def fail_if_called(*_args, **_kwargs):
        raise AssertionError("Word-dokumenter må ikke bygges under HTML-eksporten")

    monkeypatch.setattr(beta, "_word_documents", fail_if_called)
    package = beta.make_teacher_export_package(schedule)
    assert package


def test_excel_and_both_html_exports_include_unplaced_block():
    schedule, _ = make_beta_schedule([("a", "b")] * 4, end=clock(9, 0))
    workbook = load_workbook(io.BytesIO(beta.make_schedule_excel_beta(schedule)), read_only=True)
    assert "Kræver anden dag" in workbook.sheetnames
    assert "Kan ikke placeres inden for det valgte tidsrum" in beta.make_schedule_html_beta(schedule)
    assert "Kan ikke placeres inden for det valgte tidsrum" in beta.make_student_schedule_html_beta(schedule)


def test_email_columns_are_parsed_and_known_teacher_addresses_are_filled():
    student_frame = pd.DataFrame({
        "Elevnavn": ["Ane"], "Klasse": ["3q"], "Fag 1": ["Dansk A"],
        "Fag 2": ["Historie B"], "Elev Teams-email": ["NEXT1@EDU.NEXTKBH.DK"],
    })
    parsed_students = beta.parse_students_beta(student_frame)
    assert parsed_students[0]["teams_email"] == "next1@edu.nextkbh.dk"
    teacher_frame = pd.DataFrame({
        "Initialer": ["hst", "matr"], "Lærernavn": ["Henrik", "Mathias"],
        "Fag": ["Dansk", "Historie"], "Lærer Teams-email": ["", "matr@nextkbh.dk"],
    })
    parsed_teachers, _ = beta.parse_teachers_beta(teacher_frame)
    emails = {item["id"]: item["teams_email"] for item in parsed_teachers}
    assert emails == {"hst": "hst@nextkbh.dk", "matr": "matr@nextkbh.dk"}


def test_generic_email_column_is_used_as_the_students_teams_id():
    student_frame = pd.DataFrame({
        "Elevnavn": ["Ane"], "Klasse": ["3q"], "Fag 1": ["Dansk A"],
        "Fag 2": ["Historie B"], "Email": ["NEXT1@EDU.NEXTKBH.DK"],
    })
    assert beta.parse_students_beta(student_frame)[0]["teams_email"] == "next1@edu.nextkbh.dk"

    class UploadedSchedule:
        name = "vejledningsplan.csv"

        @staticmethod
        def getvalue():
            return (
                "Elev,Klasse,Lærer 1,Fag 1,Lærer 2,Fag 2,Email\n"
                "Ane,3q,hst,Dansk A,matr,Historie B,NEXT1@EDU.NEXTKBH.DK\n"
            ).encode("utf-8")

    parsed = beta.parse_schedule_upload_beta(UploadedSchedule())
    assert parsed["students"][0]["teams_email"] == "next1@edu.nextkbh.dk"
    assert {teacher["teams_email"] for teacher in parsed["teachers"]} == {
        "hst@nextkbh.dk", "matr@nextkbh.dk",
    }

    schedule = beta.make_schedule_beta(
        parsed["students"], parsed["teachers"], parsed["solution"],
        clock(8, 15), clock(16, 15), 20, 0, 0, 0, True,
        lunch_mode="Fast tidspunkt for alle lærere", lunch_start_time=clock(12), lunch_minutes=30,
        search_attempts=5, teacher_blocks={}, floating_pauses=True,
    )
    teacher_html = beta.make_schedule_html_beta(schedule)
    assert teacher_html.count('class="action action-teams"') == 2
    assert "hst%40nextkbh.dk" in teacher_html
    assert "matr%40nextkbh.dk" in teacher_html
    assert "next1%40edu.nextkbh.dk" in teacher_html
    assert 'class="action action-disabled" title="Teams-link' not in teacher_html


def test_teams_action_stays_active_when_excel_table_is_on_a_later_sheet_after_title_rows():
    excel = io.BytesIO()
    plan = pd.DataFrame({
        "Elev": ["Ane"], "Email": ["next27888@edu.nextkbh.dk"], "Klasse": ["S 2024q"],
        "Lærer 1": ["Henrik Sterner (hst)"], "Fag 1": ["Dansk A"],
        "Lærer 2": ["Mathias Ramberg (matr)"], "Fag 2": ["Historie B"],
    })
    with pd.ExcelWriter(excel, engine="openpyxl") as writer:
        pd.DataFrame({"Information": ["Godkendt SOP-fordeling"]}).to_excel(writer, index=False, sheet_name="Læs mig")
        plan.to_excel(writer, index=False, sheet_name="Vejledningsplan", startrow=2)

    class UploadedExcel:
        name = "godkendt-fordeling.xlsx"

        @staticmethod
        def getvalue():
            return excel.getvalue()

    parsed = beta.parse_schedule_upload_beta(UploadedExcel())
    assert parsed["students"][0]["teams_email"] == "next27888@edu.nextkbh.dk"
    schedule, _ = make_beta_schedule(
        [("hst", "matr")], students_override=parsed["students"], teachers_override=parsed["teachers"],
    )
    teacher_html = beta.make_schedule_html_beta(schedule)
    assert teacher_html.count('class="action action-teams"') == 2
    assert re.search(r'<tr data-teacher="[^"]*hst[^"]*">.*class="action action-teams"', teacher_html)
    assert "hst%40nextkbh.dk" in teacher_html
    assert "matr%40nextkbh.dk" in teacher_html
    assert "next27888%40edu.nextkbh.dk" in teacher_html


def test_schedule_template_contains_teams_email_columns():
    workbook = load_workbook(io.BytesIO(beta.make_schedule_input_template_beta()), read_only=True)
    headers = [cell.value for cell in next(workbook["Tidsplan-input"].iter_rows())]
    assert "Email" in headers
    assert "Lærer 1 Teams-email" in headers
    assert "Lærer 2 Teams-email" in headers


def test_single_student_word_filename_is_exact_requirement():
    schedule, _ = make_beta_schedule([("a", "b")])
    documents = beta._word_documents(schedule)
    assert [item["filename"] for item in documents.values()] == ["SOP-Elev 1.docx"]


def test_unplaced_activity_has_neither_teams_link_nor_word_document():
    schedule, _ = make_beta_schedule([("a", "b")] * 4, end=clock(9, 0))
    unplaced_id = schedule["unplaced"].iloc[0]["Aktivitets-ID"]
    link, error = beta.teams_chat_link(schedule, unplaced_id)
    assert link == ""
    assert "bekræftet tidspunkt" in error
    assert unplaced_id not in beta._word_documents(schedule)


def test_teacher_html_escapes_student_content_and_action_attributes():
    dangerous = student(1, name='</td><script>alert("x")</script>', email="safe@edu.nextkbh.dk")
    teachers = [teacher("a", "A", "a@nextkbh.dk"), teacher("b", "B", "b@nextkbh.dk")]
    schedule, _ = make_beta_schedule([("a", "b")], students_override=[dangerous], teachers_override=teachers)
    teacher_html = beta.make_schedule_html_beta(schedule)
    assert "<script>alert" not in teacher_html
    assert "&lt;/td&gt;&lt;script&gt;" in teacher_html


def test_lock_history_supports_undo_redo_and_clears_redo_on_new_move():
    state = {}
    beta.push_lock_history(state, {})
    current = {"A-E001": "08:20"}
    assert beta.undo_lock_history(state, current) == {}
    assert beta.redo_lock_history(state, {}) == current
    state["beta_redo_stack"] = [{"old": "value"}]
    beta.push_lock_history(state, current)
    assert state["beta_redo_stack"] == []


def test_stale_plan_version_is_rejected(monkeypatch):
    fake_state = {}
    monkeypatch.setattr(beta.st, "session_state", fake_state)
    schedule, context = make_beta_schedule([("a", "b")] * 2)
    fake_state["v2_schedule"] = schedule
    fake_state["beta_context_marker"] = beta.core.stable_signature([
        [(item.get("id", ""), item.get("name", ""), item.get("className", "")) for item in context["students"]],
        context["solution"].get("assignments", []), context["start_time"], context["end_time"],
        context["student_minutes"], context["pause_count"], context["pause_minutes"],
        context["transition_minutes"], context["lunch_mode"], context["lunch_start_time"],
        context["lunch_minutes"], context["teacher_blocks"], context["floating_pauses"],
    ])
    with pytest.raises(ValueError, match="ændret siden"):
        beta.make_schedule_beta(**context, locked_activities={}, expected_version=999)


def test_invalid_settings_still_raise_clear_errors():
    students = [student(1)]
    teachers = [teacher("a", "A"), teacher("b", "B")]
    with pytest.raises(ValueError, match="Sluttidspunktet"):
        beta.make_schedule_beta(
            students, teachers, {"assignments": [["a", "b"]]}, clock(10), clock(9),
            20, 0, 0, 0, False,
        )


def test_beta_streamlit_app_starts_without_exceptions():
    rendered = AppTest.from_file("app_beta.py").run(timeout=45)
    assert not rendered.exception

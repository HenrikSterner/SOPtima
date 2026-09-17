from __future__ import annotations

import io
import random
import sys
import time
import types
import zipfile
from datetime import time as clock
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest

try:
    from docx import Document
except (ImportError, OSError):
    Document = None

import app


class Upload:
    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data

    def getvalue(self) -> bytes:
        return self._data


def student(
    name: str = "Elev 1",
    subjects: tuple[str, str] = ("Dansk A", "Historie B"),
    wishes: tuple[str, ...] = (),
    student_id: str = "E001",
) -> dict:
    return {
        "id": student_id,
        "name": name,
        "className": "3a",
        "subjects": list(subjects),
        "subjectsWithLevel": list(subjects),
        "wishes": list(wishes),
        "projectTitle": "Et projekt",
        "projectDescription": "En beskrivelse",
    }


def teacher(teacher_id: str, subjects: tuple[str, ...], name: str | None = None) -> dict:
    return {"id": teacher_id, "name": name or teacher_id.upper(), "subjects": list(subjects), "holds": []}


def optimise(students, teachers, capacities, **overrides):
    settings = {
        "K": 10,
        "double_limit": 0,
        "use_global_k": True,
        "allow_over_capacity": False,
        "lock_teacher_max": True,
        "prioritize_pairs": True,
        "prioritize_classes": True,
        "attempts": 20,
    }
    settings.update(overrides)
    return app.optimize(students, teachers, capacities, **settings)


def schedule_for_pairs(
    pairs: list[tuple[str, str]],
    *,
    group_pairs: bool = False,
    attempts: int = 20,
    avoid_teacher_gaps: bool = False,
):
    ids = sorted({teacher_id for pair in pairs for teacher_id in pair})
    teachers = [teacher(teacher_id, (f"Fag {teacher_id}",)) for teacher_id in ids]
    students = [
        student(
            name=f"Elev {index + 1}",
            subjects=(f"Fag {pair[0]}", f"Fag {pair[1]}"),
            student_id=f"E{index + 1:03d}",
        )
        for index, pair in enumerate(pairs)
    ]
    solution = {"assignments": [list(pair) for pair in pairs]}
    return app.make_schedule(
        students,
        teachers,
        solution,
        clock(8, 0),
        clock(16, 0),
        20,
        0,
        0,
        0,
        group_pairs,
        lunch_mode="Flydende for alle lærere",
        lunch_minutes=30,
        search_attempts=attempts,
        avoid_teacher_gaps=avoid_teacher_gaps,
    )


def guidance_round_count(schedule: dict) -> int:
    frame = schedule["students"]
    return frame[["Start", "Slut"]].drop_duplicates().shape[0]


def pair_gap_rounds(schedule: dict) -> int:
    frame = schedule["students"].sort_values(["Start", "Lærerpar"], kind="stable")
    rounds = {start: index for index, start in enumerate(sorted(frame["Start"].unique()))}
    gaps = 0
    for _, rows in frame.groupby("Lærerpar"):
        positions = sorted(rounds[start] for start in rows["Start"])
        if len(positions) > 1:
            gaps += positions[-1] - positions[0] + 1 - len(positions)
    return gaps


def assert_solution_invariants(students, teachers, capacities, solution):
    assessment = app.assignment_assessment(
        students,
        teachers,
        capacities,
        solution["assignments"],
        solution["K"],
        solution["double_limit"],
        solution["use_global_k"],
        solution.get("allow_over_capacity", False),
        solution.get("lock_teacher_max", False),
    )
    assert assessment["errors"] == []
    assert assessment["loads"] == solution["loads"]
    assert assessment["double_loads"] == solution["double_loads"]


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("  DANSK   A ", "dansk"),
        ("komm/it", "kommunikation og it"),
        ("Komm/info", "kommunikation og it"),
        ("Idéhistorie B", "idehistorie"),
        ("Idræt C", "idraet"),
        ("Erhvervsøkonomi A*", "erhvervsoekonomi"),
    ],
)
def test_subject_normalisation(raw, expected):
    assert app.canonical_subject(raw) == expected


def test_repair_text_repairs_double_encoded_danish():
    assert app.repair_text("LÃ¦rer") == "Lærer"
    assert app.repair_text(None) == ""
    assert app.normal_key("  AB  ") == "ab"


@pytest.mark.parametrize("raw, expected", [("12", 12), ("12.0", 12), ("12 elever", 12), ("", 0), (None, 0)])
def test_capacity_parser_accepts_clear_values(raw, expected):
    assert app.parse_capacity(raw) == expected


@pytest.mark.parametrize("raw", ["12,5", "12.5", "-1", "abc", "999"])
def test_capacity_parser_rejects_ambiguous_values(raw):
    with pytest.raises(ValueError):
        app.parse_capacity(raw)


def test_read_csv_detects_separator_utf8_and_cp1252():
    text = "Elevnavn;Klasse;Fag 1;Fag 2\nÅse Øster;3a;Dansk A;Idræt C\n"
    for encoding in ("utf-8-sig", "cp1252"):
        frame = app.read_uploaded_table(Upload("elever.csv", text.encode(encoding)))
        parsed = app.parse_students(frame)
        assert parsed[0]["name"] == "Åse Øster"
        assert parsed[0]["subjects"] == ["Dansk A", "Idræt C"]


def test_parse_students_accepts_alternative_headers_and_unknown_class():
    frame = pd.DataFrame([{"Navn": "Eva", "Subject 1": "Dansk", "Subject 2": "Historie", "Wish 1": "AB"}])
    parsed = app.parse_students(frame)
    assert parsed[0]["className"] == "Ukendt klasse/hold"
    assert parsed[0]["wishes"] == ["ab"]


def test_parse_students_reports_nonempty_row_without_name():
    frame = pd.DataFrame([{"Elevnavn": "", "Klasse": "3a", "Fag 1": "Dansk", "Fag 2": "Historie"}])
    with pytest.raises(ValueError, match="mangler elevnavn"):
        app.parse_students(frame)


def test_parse_students_preserves_one_subject_for_validation():
    parsed = app.parse_students(pd.DataFrame([{"Elevnavn": "Eva", "Klasse": "3a", "Fag 1": "Dansk", "Fag 2": ""}]))
    assert any("både Fag 1 og Fag 2" in error for error in app.data_readiness(parsed, []))


def test_parse_teachers_merges_rows_and_preserves_teacher_without_subject():
    frame = pd.DataFrame(
        [
            {"Initialer": "AB", "Lærernavn": "Anna", "Fag": "Dansk", "Max": "12"},
            {"Initialer": "ab", "Lærernavn": "Anna", "Fag": "Historie", "Max": "10"},
            {"Initialer": "CD", "Lærernavn": "Carl", "Fag": "", "Max": "8"},
        ]
    )
    teachers, capacities = app.parse_teachers(frame)
    assert len(teachers) == 2
    assert teachers[0]["subjects"] == ["Dansk", "Historie"]
    assert teachers[1]["subjects"] == []
    assert capacities == {"ab": 10, "cd": 8}
    assert any("mangler mindst ét fag" in error for error in app.data_readiness([], teachers))


def test_parse_teachers_rejects_data_row_without_initials():
    frame = pd.DataFrame([{"Initialer": "", "Lærernavn": "Anna", "Fag": "Dansk", "Max": 12}])
    with pytest.raises(ValueError, match="mangler initialer"):
        app.parse_teachers(frame)


def test_parse_teacher_workbook_uses_lower_capacity_from_second_sheet():
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        pd.DataFrame([{"Initialer": "AB", "Lærernavn": "Anna", "Fag": "Dansk", "Max": 12}]).to_excel(writer, index=False, sheet_name="Lærere")
        pd.DataFrame([{"Initialer": "ab", "Max": 8}]).to_excel(writer, index=False, sheet_name="Max")
    teachers, capacities = app.parse_teacher_upload(Upload("lærere.xlsx", stream.getvalue()))
    assert len(teachers) == 1
    assert capacities["ab"] == 8


def test_resolve_wishes_supports_id_name_and_label():
    teachers = [teacher("ab", ("Dansk",), "Anna Borg")]
    for wish in ("AB", "Anna Borg", "Anna Borg (ab)"):
        students = [student(wishes=(wish,))]
        app.resolve_wishes(students, teachers)
        assert students[0]["wishes"] == ["ab"]


def test_duplicate_teacher_names_remain_ambiguous():
    teachers = [teacher("ab", ("Dansk",), "Anna"), teacher("cd", ("Historie",), "Anna")]
    students = [student(wishes=("Anna",))]
    app.resolve_wishes(students, teachers)
    validation = app.input_validation(students, teachers)
    assert validation["unknown_wishes"] == []
    assert validation["ambiguous_wishes"][0]["Mulige initialer"] == "ab · cd"
    assert any("tvetydige" in error for error in app.data_readiness(students, teachers))


def test_data_readiness_reports_coverage_unknown_wish_and_duplicate_student_id():
    students = [student(wishes=("zz",)), student(name="Elev 2", student_id="E001")]
    teachers = [teacher("ab", ("Dansk",))]
    errors = app.data_readiness(students, teachers)
    assert any("findes ikke" in error for error in errors)
    assert any("ingen lærer" in error for error in errors)
    assert any("Dubletter" in error for error in errors)


def test_wish_status_counts_distinct_wishes_not_subject_slots():
    entry = student(wishes=("ab", "ab"))
    assert app.wish_status(entry, ["ab", "ab"]) == (1, 1)
    assert app.wish_status_label(entry, ["ab", "ab"]) == "1/1"
    assert app.wish_status(student(), ["ab", "cd"]) == (0, 0)


def test_minimal_optimisation_has_only_legal_solution():
    students = [student(wishes=("ab", "cd"))]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    solution = optimise(students, teachers, {"ab": 5, "cd": 5})
    assert solution["assignments"] == [["ab", "cd"]]
    assert solution["stats"]["all_wishes"] == 1
    assert_solution_invariants(students, teachers, {"ab": 5, "cd": 5}, solution)


def test_double_guidance_obeys_i_and_counts_load_once():
    students = [student()]
    teachers = [teacher("ab", ("Dansk", "Historie"))]
    blocked = optimise(students, teachers, {"ab": 5}, double_limit=0)
    assert blocked["stats"]["unassigned"] == 1
    allowed = optimise(students, teachers, {"ab": 5}, double_limit=1)
    assert allowed["assignments"] == [["ab", "ab"]]
    assert allowed["loads"]["ab"] == 1
    assert allowed["double_loads"]["ab"] == 1


def test_global_and_individual_capacity_modes():
    students = [student(name=f"Elev {index}", student_id=f"E{index:03d}") for index in range(1, 4)]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    locked = optimise(students, teachers, {"ab": 1, "cd": 1}, K=2, lock_teacher_max=True)
    assert locked["stats"]["unassigned"] == 2
    flexible = optimise(students, teachers, {"ab": 1, "cd": 1}, K=2, allow_over_capacity=True, lock_teacher_max=False)
    assert max(flexible["loads"].values()) <= 2
    assert flexible["stats"]["unassigned"] == 1


def test_optimisation_is_deterministic_and_depth_cannot_worsen_score():
    students = [student(name=f"Elev {index}", wishes=("ab", "cd"), student_id=f"E{index:03d}") for index in range(1, 7)]
    teachers = [teacher("ab", ("Dansk",)), teacher("ef", ("Dansk",)), teacher("cd", ("Historie",)), teacher("gh", ("Historie",))]
    capacities = {item["id"]: 6 for item in teachers}
    first = optimise(students, teachers, capacities, attempts=20)
    repeated = optimise(students, teachers, capacities, attempts=20)
    deeper = optimise(students, teachers, capacities, attempts=40)
    assert first["assignments"] == repeated["assignments"]
    assert deeper["stats"]["score"] <= first["stats"]["score"]
    assert_solution_invariants(students, teachers, capacities, deeper)


@pytest.mark.parametrize("seed", range(10))
def test_generated_distributions_respect_all_invariants(seed):
    rng = random.Random(seed)
    subjects = [f"Fag {index}" for index in range(5)]
    teachers = [
        teacher(f"t{subject_index}{candidate_index}", (subject,))
        for subject_index, subject in enumerate(subjects)
        for candidate_index in range(3)
    ]
    capacities = {item["id"]: 12 for item in teachers}
    students = []
    for index in range(30):
        selected = rng.sample(subjects, 2)
        possible_wishes = [item["id"] for item in teachers if item["subjects"][0] in selected]
        students.append(student(
            name=f"Elev {index}", subjects=tuple(selected), wishes=tuple(rng.sample(possible_wishes, rng.randint(0, 2))),
            student_id=f"E{index:03d}",
        ))
    solution = optimise(students, teachers, capacities, K=12, attempts=20)
    assert solution["stats"]["unassigned"] == 0
    assert_solution_invariants(students, teachers, capacities, solution)


def test_large_1000_student_200_teacher_distribution():
    rng = random.Random(20260914)
    subjects = [f"Fag {index}" for index in range(20)]
    teachers = [
        teacher(f"t{subject_index:02d}{candidate_index:02d}", (subject,))
        for subject_index, subject in enumerate(subjects)
        for candidate_index in range(10)
    ]
    capacities = {item["id"]: 20 for item in teachers}
    students = [
        student(name=f"Elev {index}", subjects=tuple(rng.sample(subjects, 2)), student_id=f"E{index:04d}")
        for index in range(1000)
    ]
    started = time.perf_counter()
    solution = optimise(students, teachers, capacities, K=20, attempts=20)
    assert time.perf_counter() - started < 30
    assert solution["stats"]["unassigned"] == 0
    assert_solution_invariants(students, teachers, capacities, solution)


def test_assignment_assessment_rejects_manual_k_and_i_breaches():
    students = [student(name="E1", student_id="E1"), student(name="E2", student_id="E2")]
    teachers = [teacher("ab", ("Dansk", "Historie"))]
    assessment = app.assignment_assessment(students, teachers, {"ab": 1}, [["ab", "ab"], ["ab", "ab"]], 1, 0, True, False, True)
    assert any("hårde grænse" in error for error in assessment["errors"])
    assert any("I-grænsen" in error for error in assessment["errors"])


def test_assignment_assessment_allows_visible_individual_overage_up_to_k():
    students = [student(name="E1", student_id="E1"), student(name="E2", student_id="E2")]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    assessment = app.assignment_assessment(students, teachers, {"ab": 1, "cd": 1}, [["ab", "cd"], ["ab", "cd"]], 2, 0, True, True, False)
    assert assessment["errors"] == []
    assert len(assessment["warnings"]) == 2


def test_parallel_conflict_chain_and_triangle_rounds():
    parallel = schedule_for_pairs([("ab", "cd"), ("ef", "gh"), ("ij", "kl"), ("mn", "op")])
    chain = schedule_for_pairs([("ab", "cd"), ("cd", "ef"), ("ef", "gh")])
    triangle = schedule_for_pairs([("ab", "cd"), ("cd", "ef"), ("ef", "ab")])
    assert guidance_round_count(parallel) == 1
    assert guidance_round_count(chain) == 2
    assert guidance_round_count(triangle) == 3


def test_schedule_never_double_books_teacher():
    schedule = schedule_for_pairs([("ab", "cd"), ("cd", "ef"), ("gh", "ab"), ("ij", "kl"), ("ef", "gh")], attempts=80)
    guidance = schedule["teachers"][schedule["teachers"]["Type"] == "Vejledning"]
    for _, rows in guidance.groupby("Initialer"):
        times = rows[["Start", "Slut"]].to_records(index=False).tolist()
        assert len(times) == len(set(times))


def test_schedule_depth_cannot_increase_round_count():
    pairs = [("ab", "cd"), ("cd", "ef"), ("ef", "gh"), ("gh", "ab"), ("ab", "ef")]
    shallow = schedule_for_pairs(pairs, attempts=10)
    deep = schedule_for_pairs(pairs, attempts=80)
    assert guidance_round_count(deep) <= guidance_round_count(shallow)


def test_grouping_pairs_preserves_conflict_rules_and_round_count():
    pairs = [("ab", "cd"), ("ef", "gh"), ("ab", "cd"), ("ij", "kl"), ("ef", "gh")]
    ungrouped = schedule_for_pairs(pairs, group_pairs=False, attempts=80)
    grouped = schedule_for_pairs(pairs, group_pairs=True, attempts=80)
    assert guidance_round_count(grouped) <= guidance_round_count(ungrouped)
    guidance = grouped["teachers"][grouped["teachers"]["Type"] == "Vejledning"]
    for _, rows in guidance.groupby("Initialer"):
        assert not rows.duplicated(subset=["Start", "Slut"]).any()


def test_grouping_pairs_stays_primary_when_teacher_gaps_are_also_optimised():
    pairs = [
        ("a", "b"), ("c", "d"), ("a", "b"), ("a", "e"),
        ("c", "d"), ("f", "g"), ("f", "g"),
    ]
    schedule = schedule_for_pairs(
        pairs,
        group_pairs=True,
        attempts=40,
        avoid_teacher_gaps=True,
    )

    assert pair_gap_rounds(schedule) == 0
    assert schedule["settings"]["Lærerparhuller (runder)"] == 0


def test_teacher_gap_optimisation_collects_rounds_without_changing_round_count():
    pairs = [("d", "h"), ("c", "f"), ("e", "i"), ("d", "g"), ("b", "e"), ("d", "g"), ("d", "f")]
    regular = schedule_for_pairs(pairs, attempts=30)
    optimised = schedule_for_pairs(pairs, attempts=30, avoid_teacher_gaps=True)

    assert optimised["settings"]["Lærerhuller (runder)"] < regular["settings"]["Lærerhuller (runder)"]
    assert guidance_round_count(optimised) == guidance_round_count(regular)
    guidance = optimised["teachers"][optimised["teachers"]["Type"] == "Vejledning"]
    for _, rows in guidance.groupby("Initialer"):
        assert not rows.duplicated(subset=["Start", "Slut"]).any()


def test_schedule_validates_missing_assignments_and_times():
    students = [student()]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    with pytest.raises(ValueError, match="uden to vejledere"):
        app.make_schedule(students, teachers, {"assignments": [["ab", None]]}, clock(9), clock(15), 20, 0, 0, 0, False)
    with pytest.raises(ValueError, match="efter starttidspunktet"):
        app.make_schedule(students, teachers, {"assignments": [["ab", "cd"]]}, clock(15), clock(9), 20, 0, 0, 0, False)


def test_schedule_reports_required_time_and_fixed_lunch_bounds():
    students = [student()]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    solution = {"assignments": [["ab", "cd"]]}
    with pytest.raises(ValueError, match="faste frokostpause"):
        app.make_schedule(students, teachers, solution, clock(9), clock(10), 20, 0, 0, 0, False, lunch_start_time=clock(12))
    with pytest.raises(ValueError, match="Der mangler derfor"):
        app.make_schedule(students, teachers, solution, clock(9), clock(9, 30), 20, 0, 0, 0, False, lunch_mode="Flydende for alle lærere", lunch_minutes=30)


def test_schedule_respects_individual_teacher_block_and_shows_it():
    students = [student()]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    schedule = app.make_schedule(
        students, teachers, {"assignments": [["ab", "cd"]]}, clock(8), clock(16), 20, 0, 0, 0, False,
        lunch_mode="Flydende for alle lærere", lunch_minutes=30,
        teacher_blocks={"ab": [(clock(8), clock(9))]},
    )
    assert schedule["students"].iloc[0]["Start"] == "09:00"
    blocked = schedule["teachers"][schedule["teachers"]["Type"] == "Spærret"]
    assert blocked[["Initialer", "Start", "Slut"]].values.tolist() == [["ab", "08:00", "09:00"]]
    assert schedule["settings"]["Lærerspærringer"] == 1


def test_teacher_blocks_merge_and_validate_incomplete_editor_rows():
    teachers = [teacher("ab", ("Dansk",))]
    blocks = app.normalise_teacher_blocks({"AB": [("08:00", "09:00"), ("08:30", "10:00")]}, teachers)
    assert blocks["ab"] == [(clock(8), clock(10))]
    frame = pd.DataFrame([{"Lærer": "AB", "Initialer": "ab", "Spærret fra": "08:00", "Spærret til": ""}])
    with pytest.raises(ValueError, match="både fra- og til-tid"):
        app.teacher_blocks_from_table(frame)


def test_teacher_block_time_choices_follow_the_selected_day():
    choices = app.schedule_time_choices(clock(8, 15), clock(8, 30))
    assert choices == [clock(8, 15), clock(8, 20), clock(8, 25), clock(8, 30)]
    assert app.schedule_time_choices(clock(16), clock(8)) == []


def test_teacher_blocks_explain_delay_in_floating_lunch_mode():
    students = [
        student(name="A", subjects=("Dansk", "Historie"), student_id="E001"),
        student(name="B", subjects=("Fysik", "Kemi"), student_id="E002"),
    ]
    teachers = [
        teacher("ab", ("Dansk",)), teacher("cd", ("Historie",)),
        teacher("ef", ("Fysik",)), teacher("gh", ("Kemi",)),
    ]
    solution = {"assignments": [["ab", "cd"], ["ef", "gh"]]}
    schedule = app.make_schedule(
        students, teachers, solution, clock(8), clock(12), 20, 0, 0, 0, False,
        lunch_mode="Flydende for alle lærere", lunch_minutes=30,
        teacher_blocks={"ab": [(clock(8), clock(9))]},
    )
    assert schedule["students"]["Start"].min() >= "09:00"
    with pytest.raises(ValueError, match="Lærerspærringer") as error:
        app.make_schedule(
            [students[0]], teachers[:2], {"assignments": [["ab", "cd"]]}, clock(8), clock(8, 30), 20, 0, 0, 0, False,
            lunch_mode="Flydende for alle lærere", lunch_minutes=30,
            teacher_blocks={"ab": [(clock(8), clock(9))]},
        )
    assert "faste frokosttid" not in str(error.value)
    assert "anden dag" in str(error.value)


def test_floating_lunch_is_present_even_for_one_round():
    schedule = schedule_for_pairs([("ab", "cd")])
    lunches = schedule["timeline"][schedule["timeline"]["Type"] == "Frokostpause"]
    assert len(lunches) == 1
    assert schedule["settings"]["Planlagt tidsforbrug (min.)"] == 50


def test_regular_pauses_are_in_teacher_plans_and_visual_gantt():
    students = [student(name="A", student_id="E001"), student(name="B", student_id="E002")]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    schedule = app.make_schedule(
        students, teachers, {"assignments": [["ab", "cd"], ["ab", "cd"]]}, clock(8), clock(12), 20, 1, 10, 0, False,
        lunch_mode="Flydende for alle lærere", lunch_minutes=30,
    )
    pauses = schedule["teachers"][schedule["teachers"]["Type"] == "Pause"]
    assert len(pauses) == 2
    assert set(pauses["Lærer"]) == {"ab", "cd"}
    assert ">Pause<" not in app.make_teacher_gantt_html(schedule, "ab")


def test_floating_pauses_are_shared_before_and_after_lunch_without_visual_blocks():
    students = [
        student(name="A", student_id="E001"),
        student(name="B", student_id="E002"),
        student(name="C", student_id="E003"),
    ]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",)), teacher("ef", ("Matematik",))]
    schedule = app.make_schedule(
        students, teachers, {"assignments": [["ab", "cd"], ["cd", "ef"], ["ef", "ab"]]}, clock(8), clock(12), 20, 2, 10, 0, False,
        lunch_mode="Flydende for alle lærere", lunch_minutes=30, floating_pauses=True,
    )
    pauses = schedule["teachers"][schedule["teachers"]["Type"] == "Pause"]
    assert len(pauses) == 6
    timeline_pauses = schedule["timeline"][schedule["timeline"]["Type"] == "Pause"]
    lunch = schedule["timeline"][schedule["timeline"]["Type"] == "Frokostpause"].iloc[0]
    assert len(timeline_pauses) == 2
    assert timeline_pauses.iloc[0]["Slut"] <= lunch["Start"]
    assert timeline_pauses.iloc[1]["Start"] >= lunch["Slut"]
    assert schedule["settings"]["Pauser fordeles om frokost"] == "Ja"
    assert ">Pause<" not in app.make_teacher_gantt_html(schedule, "ab")


def test_schedule_upload_finds_later_sheet_and_normalises_teacher_label():
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        pd.DataFrame({"Titel": ["Ikke tabellen"]}).to_excel(writer, index=False, sheet_name="Forside")
        pd.DataFrame(
            [
                ["Vejledningsliste", "", "", "", "", ""],
                ["Elev", "Klasse", "Lærer 1", "Fag 1", "Lærer 2", "Fag 2"],
                ["Eva", "3a", "Anna Borg (AB)", "Dansk", "Carl Dahl (CD)", "Historie"],
            ]
        ).to_excel(writer, index=False, header=False, sheet_name="Plan")
    parsed = app.parse_schedule_upload(Upload("plan.xlsx", stream.getvalue()))
    assert parsed["solution"]["assignments"] == [["ab", "cd"]]
    assert [item["name"] for item in parsed["teachers"]] == ["Anna Borg", "Carl Dahl"]


def test_schedule_upload_rejects_incomplete_row_atomically():
    frame = pd.DataFrame([{"Elev": "Eva", "Klasse": "3a", "Lærer 1": "AB", "Fag 1": "Dansk", "Lærer 2": "", "Fag 2": "Historie"}])
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine="openpyxl") as writer:
        frame.to_excel(writer, index=False)
    with pytest.raises(ValueError, match="mangler elevens to lærere eller fag"):
        app.parse_schedule_upload(Upload("plan.xlsx", stream.getvalue()))


def test_distribution_excel_has_selected_sheets_and_neutralises_formulas():
    students = [student(name="=HYPERLINK(\"bad\")")]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    capacities = {"ab": 2, "cd": 2}
    solution = optimise(students, teachers, capacities)
    data = app.make_export(students, teachers, solution, capacities, True, ["Fordeling", "Lærerbelastning"])
    workbook = load_workbook(io.BytesIO(data), data_only=False)
    assert workbook.sheetnames == ["Fordeling", "Lærerbelastning"]
    cell = workbook["Fordeling"]["B2"]
    assert cell.data_type != "f"
    assert str(cell.value).startswith("'=")
    assert workbook["Fordeling"].freeze_panes == "A2"


def test_input_excel_roundtrip_preserves_data():
    students = [student(wishes=("ab",))]
    teachers = [teacher("ab", ("Dansk",), "Anna"), teacher("cd", ("Historie",), "Carl")]
    capacities = {"ab": 3, "cd": 4}
    data = app.make_input_export(students, teachers, capacities)
    workbook = pd.read_excel(io.BytesIO(data), sheet_name=None, dtype=object)
    parsed_students = app.parse_students(workbook["Elevdata"])
    parsed_teachers, parsed_capacities = app.parse_teachers(workbook["Lærerdata"])
    app.resolve_wishes(parsed_students, parsed_teachers)
    assert parsed_students[0]["wishes"] == ["ab"]
    assert parsed_capacities == capacities


def test_word_zip_uses_unique_safe_names(monkeypatch):
    class FakeDocument:
        def __init__(self):
            self.lines = []

        def add_heading(self, text, _level):
            self.lines.append(text)

        def add_paragraph(self, text):
            self.lines.append(text)

        def save(self, output):
            output.write("\n".join(self.lines).encode("utf-8"))

    monkeypatch.setitem(sys.modules, "docx", types.SimpleNamespace(Document=FakeDocument))
    students = [student(name="Samme/Navn", student_id="E001"), student(name="Samme/Navn", student_id="E002")]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    solution = {"assignments": [["ab", "cd"], ["ab", "cd"]]}
    data = app.make_docx_zip(students, teachers, solution)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert len(archive.namelist()) == 2
        assert len(set(archive.namelist())) == 2
        assert all("/" not in name for name in archive.namelist())


@pytest.mark.skipif(Document is None, reason="python-docx/lxml kan ikke indlæses i det lokale, programkontrollerede Windows-miljø")
def test_real_word_document_is_valid_and_contains_student():
    students = [student(name="Samme/Navn", student_id="E001")]
    teachers = [teacher("ab", ("Dansk",)), teacher("cd", ("Historie",))]
    solution = {"assignments": [["ab", "cd"]]}
    data = app.make_docx_zip(students, teachers, solution)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        document = Document(io.BytesIO(archive.read(archive.namelist()[0])))
        assert "Samme/Navn" in "\n".join(paragraph.text for paragraph in document.paragraphs)


def test_schedule_exports_match_ui_frames_and_escape_html():
    schedule = schedule_for_pairs([("ab", "cd"), ("ef", "gh")])
    schedule["students"].loc[0, "Elev"] = "<script>alert(1)</script>"
    html = app.make_schedule_html(schedule)
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
    assert "<script>alert(1)</script>" not in html
    workbook = pd.read_excel(io.BytesIO(app.make_schedule_excel(schedule)), sheet_name=None)
    assert list(workbook) == ["Elevplan", "Lærerplan", "ab", "cd", "ef", "gh", "Tidslinje", "Lærerpar", "Indstillinger"]
    assert len(workbook["Elevplan"]) == len(schedule["students"])
    assert workbook["Lærerplan"]["Lærer"].tolist() == sorted(workbook["Lærerplan"]["Lærer"].tolist(), key=app.normal_key)
    assert set(workbook["cd"]["Lærer"]) == {"cd"}


def test_excel_teacher_sheet_names_are_safe_and_unique():
    used = {"lærerplan"}
    first = app.excel_sheet_name("Anne/Andersen [AB]", used)
    second = app.excel_sheet_name("Anne/Andersen [AB]", used)
    assert first == "Anne Andersen AB"
    assert second == "Anne Andersen AB (2)"


def test_student_lookup_html_and_teacher_gantt_include_search_and_co_teacher():
    schedule = schedule_for_pairs([("ab", "cd")])
    lookup = app.make_student_schedule_html(schedule)
    assert 'id="query"' in lookup
    assert "Elevnavn" in lookup
    assert "Navn eller klasse" not in lookup
    assert "3a" in lookup
    assert "Bemærkninger/ændringer fra lærerne" not in lookup
    assert app.compact_search_key("3.q") == "3q"
    assert app.compact_search_key("s 2024q") == "s2024q"
    assert "3q" in app.class_search_keys("S 2024q")
    assert "s2024q" in app.class_search_keys("S 2024q")
    assert "SOPtima – en algoritme udviklet af Henrik" in lookup
    assert "Elev 1" in lookup
    gantt = app.make_teacher_gantt_html(schedule, "ab")
    assert "Medvejleder: cd" in gantt
    teacher_html = app.make_schedule_html(schedule)
    assert "<h1>Vejledningsplan</h1>" in teacher_html
    assert "Fordelingen er foretaget ud fra SOPtima" in teacher_html
    assert "hst@nextkhb.dk" in teacher_html
    assert "teacher-link" in teacher_html
    assert "teacherQuery.value=link.dataset.teacherSelect" in teacher_html
    assert app.COPYRIGHT in teacher_html
    assert 'id="teacher-query"' in teacher_html
    assert 'id="teacher-agenda"' in teacher_html
    assert 'id="teacher-agenda-rows"' in teacher_html
    assert "agendaRows" in teacher_html
    assert app.COPYRIGHT in lookup


def test_demo_performance_and_invariants():
    students = app.parse_students(app.read_path_table(app.TEST_STUDENTS_FILE))
    teachers, capacities = app.parse_teachers(app.read_path_table(app.TEST_TEACHERS_FILE))
    app.resolve_wishes(students, teachers)
    started = time.perf_counter()
    solution = optimise(students, teachers, capacities, K=18, double_limit=0, allow_over_capacity=True, lock_teacher_max=False, attempts=120)
    duration = time.perf_counter() - started
    assert duration < 30
    assert_solution_invariants(students, teachers, capacities, solution)
    assert solution["stats"]["unassigned"] == 0
    schedule_started = time.perf_counter()
    schedule = app.make_schedule(
        students, teachers, solution, clock(9), clock(15), 20, 2, 15, 5, True,
        lunch_mode="Fast tidspunkt for alle lærere", lunch_start_time=clock(12), lunch_minutes=30,
        search_attempts=80,
    )
    assert time.perf_counter() - schedule_started < 15
    assert len(schedule["students"]) == len(students)
    guidance = schedule["teachers"][schedule["teachers"]["Type"] == "Vejledning"]
    for _, rows in guidance.groupby("Initialer"):
        assert not rows.duplicated(subset=["Start", "Slut"]).any()


def test_streamlit_app_smoke_and_professional_navigation():
    tested = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    assert len(tested.exception) == 0
    assert tested.sidebar.radio[0].value in {
        "1 · Elevdata", "2 · Lærerdata", "3 · Regler og max", "4 · Beregn", "5 · Resultat og eksport", "6 · Tidsplan"
    }
    assert any("350 elever" in item.value for item in tested.sidebar.markdown)
    assert any("fiktive" in item.value.casefold() for item in tested.sidebar.caption)


def test_all_six_streamlit_steps_render_without_exceptions():
    tested = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    expected_subheaders = {
        "1 · Elevdata": "Indlæs elevdata først",
        "2 · Lærerdata": "Indlæs lærerdata",
        "3 · Regler og max": "Fordelingsregler",
        "4 · Beregn": "Beregn fordeling",
        "5 · Resultat og eksport": "Resultat og eksport",
        "6 · Tidsplan": "1 · Vælg datagrundlag",
    }
    for step, expected in expected_subheaders.items():
        tested.sidebar.radio[0].set_value(step).run(timeout=30)
        assert len(tested.exception) == 0
        assert expected in [item.value for item in tested.subheader]


def test_rule_change_invalidates_calculated_solution_in_ui():
    tested = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    tested.sidebar.radio[0].set_value("4 · Beregn").run(timeout=30)
    calculate = next(button for button in tested.button if button.key == "v2_calculate")
    calculate.click().run(timeout=30)
    assert tested.session_state["solution"] is not None

    tested.sidebar.radio[0].set_value("3 · Regler og max").run(timeout=30)
    k_slider = next(slider for slider in tested.slider if slider.key == "v2_k")
    k_slider.set_value(17).run(timeout=30)
    assert tested.session_state["solution"] is None
    assert "ændret" in tested.session_state["v2_stale_notice"].casefold()


def test_demo_restore_requires_confirmation_in_ui():
    tested = AppTest.from_file(Path(__file__).resolve().parents[1] / "app.py", default_timeout=30).run()
    restore = next(button for button in tested.sidebar.button if button.key == "v2_restore_demo")
    restore.click().run(timeout=30)
    assert tested.session_state["v2_confirm_demo_restore"] is True
    assert any("erstatter" in warning.value.casefold() for warning in tested.sidebar.warning)
    cancel = next(button for button in tested.sidebar.button if button.key == "v2_restore_demo_cancel")
    cancel.click().run(timeout=30)
    assert tested.session_state["v2_confirm_demo_restore"] is False
    assert len(tested.session_state["students"]) == 350
